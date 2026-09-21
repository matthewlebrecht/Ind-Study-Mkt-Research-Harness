#!/usr/bin/env python3
"""
Item 19 (Matthew Lebrecht, 2026-09-15): composition, published coverage and the reconciler read an observation recorded
invalid AS invalid, and the counts they publish report total, valid and invalid as three separate measurements --
nothing netted out silently. The amended convention 41 wall: only those three read validity, only through
`core/validity.py::invalid_observation_ids` (the reconciler only inside `sync_observations`). Historical audit
artifacts are not rewritten. Throwaway copies of the workbook throughout.

    python core/tests/test_validity_consumers.py
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import fields
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from core import composition as C  # noqa: E402
from core import topics  # noqa: E402
from core import validity as V  # noqa: E402
from core.db import OBSERVATION_COLUMNS, MarketIntelDB, Observation, is_human_authored  # noqa: E402

PASS = FAIL = 0
LIVE = ROOT / "data" / "market_intel_db.xlsx"


def check(label, cond):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ok  {label}")
    else:
        FAIL += 1
        print(f"  FAIL {label}")


def copy_db() -> Path:
    tmp = Path(tempfile.mkdtemp()) / "test_copy.xlsx"
    shutil.copy(LIVE, tmp)
    return tmp


def obs_rows(db) -> dict:
    ws = db.wb["Observations"]
    return {ws.cell(r, 1).value: {c: ws.cell(r, i + 1).value for i, c in enumerate(OBSERVATION_COLUMNS)}
            for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}


def proposal(v: dict, **changes) -> Observation:
    names = {f.name for f in fields(Observation)}
    d = {k: x for k, x in v.items() if k in names}
    for k in ("publication_date", "retrieval_date", "evidence_excerpt", "reviewer_notes", "audit_verdict"):
        d[k] = "" if d.get(k) is None else str(d[k])
    d.update(changes)
    return Observation(**d)


def determine(wb, oids, status="invalidated_extraction_defect"):
    return V.append_determinations(wb, [{
        "observation_id": oid, "validity_status": status, "as_of_date": "2026-09-15", "determined_at": "2026-09-15",
        "determined_by": "test fixture", "basis": "test fixture: recorded invalid to exercise a consumer", "notes": ""}
        for oid in oids])


def run(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(ROOT / script), *args], cwd=ROOT, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")


def sheet_dicts(wb, name) -> list[dict]:
    it = wb[name].iter_rows(values_only=True)
    h = list(next(it))
    return [dict(zip(h, r)) for r in it if r and r[0] is not None]


def theme_of(topic: str):
    return topic if topic in topics.THEMES_BY_KEY else topics.theme_of_buyer_signal(topic)


def gap_csv(db_path: Path, *flags: str) -> tuple[list[dict], str]:
    out = ROOT / "harness_output" / "_test_validity_consumers_gap.csv"
    try:
        res = run("scripts/gap_report.py", "--db", str(db_path), "--csv", str(out.relative_to(ROOT)), *flags)
        with out.open(encoding="utf-8") as f:
            return list(csv.DictReader(f)), res.stdout
    finally:
        if out.exists():
            out.unlink()


def coverage_line(stdout: str, harness: str):
    for line in stdout.splitlines():
        parts = line.split()
        if parts and parts[0] == harness and len(parts) >= 9:
            return {"version": parts[1], "run": parts[2], "scoped": int(parts[3]), "covered": int(parts[4]),
                    "rate": parts[5], "cov_inv": int(parts[6]), "valid": parts[7]}
    return None


def main() -> int:
    print("1. composition reads an invalid observation as invalid, and says what it excluded")
    theme = min((t for t in topics.THEMES if t.buyer_detectable_since), key=lambda t: t.buyer_detectable_since)
    reach = {"H-T": {"realized_reach": "archival", "realized_reach_months": None,
                     "realized_reach_effective_from": "2020-01-01"}}

    def o(oid, cid, conf, state):
        return {"observation_id": oid, "company_id": cid, "harness_id": "H-T", "evidence_role": "buyer_acts",
                "topic": theme.key, "publication_date": "2026-09-08", "retrieval_date": "2026-09-10",
                "organizational_state": state, "source_grade": "B", "confidence_0_1": conf,
                "publication_state": "released"}
    base = dict(companies=["X1", "X2"], attempts=[], harness_reach=reach, instrument_class={}, themes=[theme],
                observations=[o("O1", "X1", 0.8, "unknown"), o("O2", "X1", 0.9, "active_transition"),
                              o("O3", "X2", 0.9, "active_transition")])
    day = dt.date(2026, 9, 15)
    marked = {"O2": "invalidated_extraction_defect", "O3": "invalidated_duplicate"}
    with_inv = C.derive(C.Inputs(**base, invalid=marked), derived_at=day)
    without = C.derive(C.Inputs(**base), derived_at=day)

    def at(rows, cid):
        return next(r for r in rows if r["company_id"] == cid and r["bucket_start"] <= "2026-09-10" <= r["bucket_end"])
    x1_plain, x1, x2 = at(without, "X1"), at(with_inv, "X1"), at(with_inv, "X2")
    check("control: with nothing recorded invalid, both X1 rows support the bucket (count 2, best confidence 0.9)",
          x1_plain["evidence_count"] == 2 and x1_plain["max_confidence"] == 0.9 and not x1_plain["notes"])
    check("the invalid row supports nothing: X1 rests on O1 alone, and its state and confidence come from O1",
          x1["status"] == "observed" and x1["supporting_observation_ids"] == "O1" and x1["evidence_count"] == 1
          and x1["max_confidence"] == 0.8 and x1["organizational_state"] == "unknown")
    check("the bucket's notes carry the three measurements and name the exclusion",
          str(x1["notes"]).startswith("evidence: 2 total, 1 valid, 1 invalid")
          and "O2 invalidated_extraction_defect" in str(x1["notes"]))
    check("a bucket only an invalid row would inform is not observed, and says why",
          x2["status"] != "observed" and not x2["supporting_observation_ids"]
          and str(x2["notes"]).startswith("evidence: 1 total, 0 valid, 1 invalid") and "O3 invalidated_duplicate" in str(x2["notes"]))
    ev = C.summarize(with_inv)["evidence"]
    check("summarize reports evidence citations as total, valid and invalid",
          ev == {"total": 3, "valid": 1, "invalid": 2, "buckets_with_invalid_excluded": 2})
    check("with nothing recorded invalid no row carries an evidence note", not [r for r in without if r["notes"]])
    wb = openpyxl.load_workbook(LIVE, read_only=True, data_only=True)
    live_invalid = V.invalid_observation_ids(wb)
    check(f"load_inputs reads the live determinations ({len(live_invalid)})",
          C.load_inputs(wb).invalid == live_invalid and len(live_invalid) >= 10)

    print("2. the reconciler counts a re-proposed invalid row as invalid and leaves it as recorded")
    db = MarketIntelDB(copy_db())
    rows = obs_rows(db)
    machine = sorted(k for k, v in rows.items() if not is_human_authored(v) and k not in live_invalid)
    m, n, h = machine[0], machine[1], "O00303"
    check("fixture: O00303 is human-authored and recorded invalid", is_human_authored(rows[h]) and h in live_invalid)
    determine(db.wb, [m])
    before_m, before_h = dict(rows[m]), dict(rows[h])
    rep = db.sync_observations([proposal(rows[m], observation_text="changed by test"),
                                proposal(rows[n], observation_text="changed by test"),
                                proposal(rows[h], observation_text="changed by test")])
    after = obs_rows(db)
    check("both invalid rows are counted invalid, with their statuses",
          sorted(i["observation_id"] for i in rep.invalid) == sorted([m, h])
          and all(i["validity_status"].startswith("invalidated") for i in rep.invalid))
    check("each invalid row is left exactly as recorded", after[m] == before_m and after[h] == before_h)
    check("the valid row is refreshed as before", [x.observation_id for x in rep.updated] == [n]
          and after[n]["observation_text"] == "changed by test")
    check("a human-authored invalid row is invalid, not a held conflict", not rep.conflicts)
    check("the run summary reports total, valid and invalid",
          rep.summary().startswith("3 proposed = 1 valid (") and rep.summary().endswith("+ 2 recorded invalid (left as recorded)"))
    check("only the valid write counts as written", rep.written == 1)

    print("3. published coverage: three measurements; the attempt-based rate does not move")
    live = run("scripts/published_coverage.py")
    obs = sheet_dicts(wb, "Observations")

    def split(state):
        ids = [x["observation_id"] for x in obs if x["publication_state"] == state]
        k = sum(1 for i in ids if i in live_invalid)
        return f"{len(ids)} {state} ({len(ids) - k} valid, {k} invalid)"
    check("the observation line reports released and quarantined as total, valid and invalid",
          f"Observations: {split('released')}, {split('quarantined')}" in live.stdout)
    fp = coverage_line(live.stdout, "H-FIRSTPARTY-01")
    check("the FIRSTPARTY line carries rate, cov-inv and valid rate", fp is not None)
    covered = sorted({a["company_id"] for a in sheet_dicts(wb, "Attempts")
                      if fp and a["run_id"] == fp["run"] and a["outcome"] == "covered"})
    target = next(c for c in covered if any(x["harness_id"] == "H-FIRSTPARTY-01" and x["company_id"] == c
                                            and x["publication_state"] == "released" for x in obs))
    ids = sorted(x["observation_id"] for x in obs if x["harness_id"] == "H-FIRSTPARTY-01" and x["company_id"] == target
                 and x["publication_state"] == "released" and x["observation_id"] not in live_invalid)
    path = copy_db()
    dbc = MarketIntelDB(path)
    determine(dbc.wb, ids)
    dbc.save()
    fp2 = coverage_line(run("scripts/published_coverage.py", "--db", str(path)).stdout, "H-FIRSTPARTY-01")
    check(f"invalidating every released FIRSTPARTY row of {target} leaves the rate unchanged",
          fp2 is not None and fp2["rate"] == fp["rate"] and fp2["covered"] == fp["covered"])
    check("... counts that covered attempt under cov-inv, and lowers the valid rate",
          fp2 is not None and fp2["cov_inv"] == fp["cov_inv"] + 1
          and float(fp2["valid"].rstrip("%")) < float(fp["valid"].rstrip("%")))

    print("4. gap report: valid, total and invalid-only companies on every theme")
    table, stdout = gap_csv(LIVE, "--include-low-grade")
    check("every theme: total = valid + invalid-only, for buyers and for providers",
          table and all(int(t["buyer_companies_total"]) == int(t["buyer_companies"]) + int(t["buyer_companies_invalid_only"])
                        and int(t["providers_messaging_total"]) == int(t["providers_messaging"]) + int(t["providers_invalid_only"])
                        for t in table))
    check("with low-grade rows counted, some theme has a company carried only by invalid rows (A073, O00529)",
          any(int(t["buyer_companies_invalid_only"]) > 0 for t in table))
    released = [x for x in obs if x["publication_state"] == "released"]
    k = sum(1 for x in released if x["observation_id"] in live_invalid)
    check("the released-rows line reports total, valid and invalid",
          f"released rows: {len(released)} total = {len(released) - k} valid + {k} recorded invalid" in stdout)
    counts: dict = {}
    for x in released:
        t = theme_of(str(x["topic"] or ""))
        if (t and x["evidence_role"] in ("buyer_acts", "buyer_articulates") and x["observation_id"] not in live_invalid
                and not str(x["evidence_excerpt"] or "").startswith(topics.LOW_GRADE_MARK)):
            counts.setdefault((t, x["company_id"]), []).append(x["observation_id"])
    (pair_theme, pair_company), (only_id,) = next((key, v) for key, v in sorted(counts.items()) if len(v) == 1)
    base_rows = {t["theme"]: t for t in gap_csv(LIVE)[0]}
    path = copy_db()
    dbg = MarketIntelDB(path)
    determine(dbg.wb, [only_id])
    dbg.save()
    moved = {t["theme"]: t for t in gap_csv(path)[0]}[pair_theme]
    was = base_rows[pair_theme]
    check(f"invalidating {only_id} ({pair_company}'s only full-grade {pair_theme} row) moves one company from valid to "
          f"invalid-only and leaves the total",
          int(moved["buyer_companies"]) == int(was["buyer_companies"]) - 1
          and int(moved["buyer_companies_invalid_only"]) == int(was["buyer_companies_invalid_only"]) + 1
          and moved["buyer_companies_total"] == was["buyer_companies_total"])

    print("5. the amended convention 41 wall")
    check("live repo: every read of validity in a gate module is permitted", V.wall_violations(ROOT) == [])
    src = {mod: (ROOT / mod).read_text(encoding="utf-8") for mod in V.GATE_MODULES}

    def breach(mod, text):
        return bool(V.wall_violations(ROOT, {mod: text}))
    check("core/audit.py importing validity is a breach", breach("core/audit.py", src["core/audit.py"] + "\nfrom core import validity\n"))
    check("core/attempts.py calling even the accessor is a breach",
          breach("core/attempts.py", src["core/attempts.py"] + "\nx = validity.invalid_observation_ids(wb)\n"))
    check("composition using anything but the accessor is a breach",
          breach("core/composition.py", src["core/composition.py"] + "\nrows = validity.current_rows([])\n"))
    check("from core.validity import ... is a breach",
          breach("core/composition.py", src["core/composition.py"] + "\nfrom core.validity import current_rows\n"))
    check("a permitted reader naming the table is a breach",
          breach("scripts/published_coverage.py", src["scripts/published_coverage.py"] + "\nS = 'Observation_Validity_History'\n"))
    check("an aliased import is a breach",
          breach("scripts/published_coverage.py", src["scripts/published_coverage.py"] + "\nfrom core import validity as v\n"))
    dsrc = src["core/db.py"]
    j = dsrc.index("\n", dsrc.index("    def apply_audit_verdict("))
    check("core/db.py calling the accessor outside sync_observations is a breach",
          breach("core/db.py", dsrc[:j + 1] + "        _x = validity.invalid_observation_ids(self.wb)\n" + dsrc[j + 1:]))
    check("... while the live core/db.py calls it inside sync_observations, permitted",
          "validity.invalid_observation_ids(" in dsrc and not breach("core/db.py", dsrc))
    check("the validator naming validity outside its fenced check is a breach",
          breach("scripts/validate_repo_db.py", src["scripts/validate_repo_db.py"] + "\n# core.validity\n"))

    print("6. historical audit artifacts are not rewritten")
    st = subprocess.run(["git", "status", "--porcelain", "--", "harness_output/audits"], cwd=ROOT,
                        capture_output=True, text=True).stdout
    check("no committed audit artifact is modified or deleted in the working tree",
          not [l for l in st.splitlines() if l[:2].strip() in ("M", "D", "R", "MM", "AM")])
    for script in ("scripts/published_coverage.py", "scripts/gap_report.py", "core/composition.py"):
        text = (ROOT / script).read_text(encoding="utf-8")
        check(f"{script} quotes no count recorded in an artifact",
              not re.search(r"population_size|precision_rate|random_control|sampled_n", text))

    print(f"\n{PASS} passed, {FAIL} failed")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
