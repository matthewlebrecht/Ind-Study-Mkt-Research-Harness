"""
Tests for the audit gate (docs/gates/gate_new_harness_output.md).

    python core/tests/test_audit_gate.py

Two things are under test and the second matters more.

The first is the rewritten never-overwrite rule. It used to key on the review_status
*value*; it now keys on `review_source` provenance, because the audit writes machine-side
verdicts that must stay overwritable while a human's review must not. Rewriting the rule
that protects the project's scarcest input is exactly the kind of change that quietly
breaks the case that was previously working, so the previously-working case is tested
explicitly alongside the new one.

The second is that check 9 can actually FAIL. A mechanical check that has only ever been
observed passing is not evidence of anything -- it is indistinguishable from a check that
returns True. So the negative cases are built deliberately: a published pair with no
artifact, an artifact missing required fields, an artifact whose own verdict is 'fail'.
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from core import audit  # noqa: E402
from core.db import (MarketIntelDB, Observation, OBSERVATION_COLUMNS,  # noqa: E402
                     is_human_authored)

PASSED = FAILED = 0


def check(cond, label):
    global PASSED, FAILED
    if cond:
        PASSED += 1
        print(f"  ok    {label}")
    else:
        FAILED += 1
        print(f"  FAIL  {label}")


def make_obs(topic, text, url="https://example.invalid/gate", **kw):
    base = dict(
        company_id="C0001", evidence_family="1_first_party_strategy_governance",
        evidence_role="buyer_articulates", topic=topic,
        organizational_state="active_transition", signal_strength="weak_clue",
        observation_text=text, evidence_excerpt="excerpt", source_url=url,
        publication_date="2026-08-01", retrieval_date="2026-08-31", source_grade="B",
        harness_id="H-GATETEST-01", harness_version="v1.0", confidence_0_1=0.5,
    )
    base.update(kw)
    return Observation(**base)


def find_row(db, obs_id):
    ws = db.wb["Observations"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == obs_id:
            return {h: ws.cell(r, i + 1).value for i, h in enumerate(headers)}
    return None


def set_cell(db, obs_id, col, value):
    ws = db.wb["Observations"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    i = headers.index(col) + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == obs_id:
            ws.cell(r, i).value = value


def artifact(harness_id="H-GATETEST-01", version="v1.0", supported=30, unsupported=0,
             wrong_entity=0, sampled=30, population=200, **kw):
    counts = {"supported": supported, "overgraded": 0, "unsupported": unsupported,
              "wrong_entity": wrong_entity}
    return audit.AuditArtifact(
        harness_id=harness_id, harness_version=version, run_id="HR-9999",
        auditor="test", population_size=population,
        strata=[audit.Stratum(stratum_id="single_common_word",
                              selection_rule="match rests on one dictionary word",
                              sampled_n=0)],
        random_control=audit.RandomControl(sampled_n=sampled, population_n=population,
                                           verdict_counts=counts),
        row_ids_sampled=[f"O{i:05d}" for i in range(sampled)], **kw).evaluate()


def main() -> int:
    print("1. the never-overwrite rule keys on provenance")
    check(is_human_authored({"review_source": "human", "review_status": "unreviewed"}),
          "review_source=human protects a row even with review_status unreviewed")
    check(not is_human_authored({"review_source": "machine",
                                 "review_status": "unreviewed"}),
          "a machine-authored unreviewed row is overwritable")
    check(not is_human_authored({}),
          "a row with neither field set is overwritable")
    # The previously-working case. Under provenance alone this row reads `machine` and
    # would be silently overwritten -- a data-loss path straight through convention 1.
    check(is_human_authored({"review_source": "machine", "review_status": "accepted"}),
          "a hand-set review_status still protects, even when review_source says machine")
    check(is_human_authored({"review_status": "corrected"}),
          "a reviewed status protects a row written before the gate migration")

    print("\n2. a machine verdict does not protect the row it sits on")
    check(not is_human_authored({"review_source": "machine", "review_status": "unreviewed",
                                 "audit_verdict": "supported"}),
          "audit_verdict alone leaves the row freely overwritable, as specified")

    print("\n3. fail-safe defaults")
    check(Observation.__dataclass_fields__["publication_state"].default == "quarantined",
          "a harness that declares nothing gets quarantined, not released")
    check(Observation.__dataclass_fields__["review_source"].default == "machine",
          "harness-written rows are machine-authored by default")
    check(OBSERVATION_COLUMNS[-3:] == ["publication_state", "audit_verdict",
                                       "review_source"],
          "column contract carries the three gate columns")

    print("\n4. sync against a throwaway copy of the workbook")
    src = Path(__file__).resolve().parents[2] / "data" / "market_intel_db.xlsx"
    tmp = Path(tempfile.mkdtemp()) / "copy.xlsx"
    shutil.copy2(src, tmp)
    db = MarketIntelDB(tmp)
    baseline = db.wb["Observations"].max_row

    rep = db.sync_observations([make_obs("gate_alpha", "original")])
    oid = rep.inserted[0].observation_id
    check(len(rep.inserted) == 1, "row inserted")
    check(find_row(db, oid)["publication_state"] == "quarantined",
          "new row lands quarantined")
    check(find_row(db, oid)["review_source"] == "machine", "new row is machine-authored")

    # An audit passes it, then the harness is fixed and the content changes.
    set_cell(db, oid, "audit_verdict", "supported")
    set_cell(db, oid, "publication_state", "released")
    rep = db.sync_observations([make_obs("gate_alpha", "REVISED after a harness fix")])
    row = find_row(db, oid)
    check(len(rep.updated) == 1, "machine row with a verdict is still overwritable")
    check(row["observation_text"] == "REVISED after a harness fix", "content refreshed")
    check(not row["audit_verdict"],
          "the stale verdict was cleared — it was reached against text that is gone")
    check(rep.verdicts_cleared == [oid], "clearing a verdict is reported, not silent")

    # Now a human owns it.
    set_cell(db, oid, "review_source", "human")
    rep = db.sync_observations([make_obs("gate_alpha", "SECOND REVISION")])
    check(len(rep.conflicts) == 1 and rep.written == 0,
          "review_source=human turns the change into a conflict")
    check(find_row(db, oid)["observation_text"] == "REVISED after a harness fix",
          "the human-owned row was left untouched")

    removed = db.delete_observations("H-GATETEST-01", keep_reviewed=False)
    check(db.wb["Observations"].max_row == baseline,
          f"test rows cleaned up ({removed} removed)")

    print("\n5. the stop rule")
    a = artifact()
    check(a.verdict == "pass" and not a.stop_rule_triggered,
          "a clean random control passes")
    check(abs(a.random_control.precision_rate - 1.0) < 1e-9, "precision rate is 30/30")

    a = artifact(supported=29, wrong_entity=1)
    check(a.verdict == "fail", "a single wrong_entity row fails, at any count")
    check("wrong_entity" in a.stop_rule_triggered[0], "and says why")

    a = artifact(supported=27, unsupported=3)   # 3/30 = 10.0%, exactly at threshold
    check(a.verdict == "pass", "an exclusion rate exactly at threshold does not trip it")
    a = artifact(supported=26, unsupported=4)   # 4/30 = 13.3%
    check(a.verdict == "fail", "an exclusion rate above threshold quarantines")

    a = artifact(supported=25)
    a.random_control.verdict_counts["overgraded"] = 5
    check(a.evaluate().verdict == "pass",
          "overgraded is a downgrade, not an exclusion, and does not trip the rule")

    print("\n6. artifact schema")
    good = artifact().to_dict()
    check(audit.validate_artifact(good) == [], "a complete artifact validates")
    check(isinstance(good["threshold_applied"], float),
          "threshold_applied is a literal number, not a config reference")
    check(good["random_control"]["population_n"] == 200,
          "the random control states its denominator explicitly")
    check(good["gate_doc_sha256"].startswith("sha256:"),
          "the gate document is pinned by content hash, not just path")
    for missing in ("run_id", "population_size", "row_ids_sampled", "threshold_applied"):
        d = dict(good)
        d.pop(missing)
        check(any(missing in p for p in audit.validate_artifact(d)),
              f"a missing {missing} is caught")
    d = json.loads(json.dumps(good))
    d["random_control"].pop("population_n")
    check(any("population_n" in p for p in audit.validate_artifact(d)),
          "a precision rate with no denominator is caught")

    print("\n7. reprocessing_required is derived, not authored")
    arts = {("H-X-01", "v1.0"): artifact("H-X-01", "v1.0", supported=29,
                                         wrong_entity=1).to_dict()}
    got = audit.derive_reprocessing_required("H-X-01", arts)
    check(got and got["reprocessing_required"], "a wrong_entity finding sets it")
    check("all prior versions" in got["blast_radius"],
          "default blast radius is the audited version and everything before it")

    arts[("H-X-01", "v1.1")] = artifact("H-X-01", "v1.1").to_dict()
    check(audit.derive_reprocessing_required("H-X-01", arts) is None,
          "a passing audit on a later version supersedes it")

    narrowed = artifact("H-Y-01", "v1.2", supported=29, wrong_entity=1)
    narrowed.defect_introduced_in = "v1.2"
    got = audit.derive_reprocessing_required("H-Y-01",
                                             {("H-Y-01", "v1.2"): narrowed.to_dict()})
    check("defect_introduced_in" in got["blast_radius"],
          "defect_introduced_in narrows the radius and says so")

    check(audit.derive_reprocessing_required("H-NOSUCH-01", {}) is None,
          "a harness with no audits requires nothing")

    print("\n8. check 9 can fail — the negative cases")
    import openpyxl
    from scripts.validate_repo_db import validate

    def run_check9(rows, artifacts_on_disk=None):
        """Stand up a workbook + audit dir and return check 9's failures."""
        d = Path(tempfile.mkdtemp())
        book = d / "db.xlsx"
        shutil.copy2(src, book)
        wb = openpyxl.load_workbook(book)
        hr = wb["Harness_Runs"]
        headers = [hr.cell(1, c).value for c in range(1, hr.max_column + 1)]
        for hid, ver, status in rows:
            vals = {"harness_run_id": f"HR-{8000 + hr.max_row}", "harness_id": hid,
                    "version": ver, "publication_status": status,
                    "records_excluded_by_audit": 0, "date_run": "2026-08-31"}
            hr.append([vals.get(h) for h in headers])
        wb.save(book)

        adir = d / "audits"
        adir.mkdir()
        # Seed the temp directory with the REAL artifacts first. The workbook this copies
        # carries whatever pairs are genuinely published and audited today, and pointing
        # the gate at an empty directory would report those as missing -- a failure about
        # the evidence base rather than about the case under test. The test has to isolate
        # its own variable, which is the synthetic pair appended above.
        real_dir = audit.AUDIT_DIR
        for art in (sorted(real_dir.glob("*__v*.json")) if real_dir.exists() else []):
            (adir / art.name).write_text(art.read_text(encoding="utf-8"), encoding="utf-8")
        audit.AUDIT_DIR = adir
        try:
            for data in (artifacts_on_disk or []):
                (adir / f"{data['harness_id']}__{data['harness_version']}.json").write_text(
                    json.dumps(data), encoding="utf-8")
            rep = validate(book)
        finally:
            audit.AUDIT_DIR = real_dir
        return [f for f in rep.failures]

    fails = run_check9([("H-NEWTHING-01", "v1.0", "published")])
    check(any("NO audit artifact" in f for f in fails),
          "a published version with no artifact FAILS the check")

    fails = run_check9([("H-NEWTHING-01", "v1.0", "quarantined")])
    check(not any("NO audit artifact" in f for f in fails),
          "the same version quarantined does not fail — quarantine is the pre-audit state")

    bad = artifact("H-NEWTHING-01", "v1.0").to_dict()
    bad.pop("row_ids_sampled")
    fails = run_check9([("H-NEWTHING-01", "v1.0", "published")], [bad])
    check(any("missing required fields" in f for f in fails),
          "an artifact missing a required field FAILS the check")

    failing = artifact("H-NEWTHING-01", "v1.0", supported=29, wrong_entity=1).to_dict()
    fails = run_check9([("H-NEWTHING-01", "v1.0", "published")], [failing])
    check(any("FAILING audit" in f for f in fails),
          "publishing over a failed audit FAILS the check")

    ok = artifact("H-NEWTHING-01", "v1.0").to_dict()
    fails = run_check9([("H-NEWTHING-01", "v1.0", "published")], [ok])
    # Scoped to check 9's own messages: the synthetic harness_id has no manifest, so
    # check 1 fails by design here and is not what this assertion is about.
    gate_msgs = ("audit artifact", "FAILING audit", "publication_status")
    check(not [f for f in fails if any(m in f for m in gate_msgs)],
          "a published version with a passing artifact raises no gate failure")

    fails = run_check9([("H-FMCSA-01", "v1.3", "published")])
    check(not any("NO audit artifact" in f for f in fails),
          "a grandfathered pair needs no artifact")

    # `superseded` -- the disposition added 2026-09-01 for a version that a later one
    # fully replaced. It must silence the artifact requirement (there is nothing live to
    # audit) WITHOUT becoming a way to publish unaudited output, so the claim it makes is
    # enforced: a superseded version still holding observations fails.
    fails = run_check9([("H-NEWTHING-01", "v1.0", "superseded")])
    check(not any("NO audit artifact" in f for f in fails),
          "a superseded version needs no artifact -- nothing of it is live")
    check(not any("expected 'published'" in f for f in fails),
          "superseded is an accepted publication_status, not a typo")

    # The negative. H-FMCSA-01 v1.3 genuinely holds 20 observations in the copied
    # workbook, so marking it superseded is a lie the check has to catch.
    fails = run_check9([("H-FMCSA-01", "v1.3", "superseded")])
    check(any("still holds" in f for f in fails),
          "a superseded version that still holds observations FAILS")

    # And it cannot be both. The same pair published on one run and superseded on
    # another would let a version quote a coverage number while claiming to be inert.
    check(any("one or the other" in f for f in fails),
          "superseded on one run and published on another FAILS")

    fails = run_check9([("H-NEWTHING-01", "v9.9", "wobbly")])
    check(any("publication_status" in f for f in fails),
          "an unknown publication_status still FAILS -- the vocabulary is closed")

    # ------------------------------------------------------------------
    print()
    print("9. an artifact is identified by its NAME, not its contents")
    # Found 2026-09-01 by reading load_artifacts()'s output, not by a test: the sampler's
    # `__strata.json` sidecars carry `harness_id` and `harness_version` at top level, so
    # globbing `*.json` and keying on those fields loaded sample-selection files as audit
    # verdicts. Two of the four entries returned were strata files. Nothing failed only
    # because both versions were quarantined; releasing either would have made check 9
    # report a valid, freshly-written artifact as malformed.
    #
    # Convention 40: the suite was green over a live defect, so the test lands with the
    # fix. It asserts the discrimination in BOTH directions -- a sidecar is not picked up,
    # and a real artifact next to one still is -- because a `load_artifacts` that returned
    # nothing at all would also pass the first half.
    d = Path(tempfile.mkdtemp())
    real_art = artifact("H-SIDECAR-01", "v1.0").to_dict()
    (d / "H-SIDECAR-01__v1.0.json").write_text(json.dumps(real_art), encoding="utf-8")
    (d / "H-SIDECAR-01__v1.0__strata.json").write_text(json.dumps({
        "harness_id": "H-SIDECAR-01", "harness_version": "v1.0", "run_id": "HR-9999",
        "population_size": 3, "seed": 20260831, "strata": [],
    }), encoding="utf-8")

    loaded = audit.load_artifacts(d)
    check(("H-SIDECAR-01", "v1.0") in loaded,
          "the real artifact is still found when a sidecar sits beside it")
    check(loaded.get(("H-SIDECAR-01", "v1.0"), {}).get("verdict") == "pass",
          "the entry returned is the ARTIFACT, not the sidecar that sorts after it")
    check(not audit.validate_artifact(loaded[("H-SIDECAR-01", "v1.0")]),
          "and it validates -- the pre-fix collision made this report as malformed")

    # A sidecar alone contributes no entry. Without this, deleting the artifact would
    # leave the strata file standing in for it and the gate would read as satisfied.
    d2 = Path(tempfile.mkdtemp())
    (d2 / "H-SIDECAR-01__v1.0__strata.json").write_text(json.dumps({
        "harness_id": "H-SIDECAR-01", "harness_version": "v1.0", "strata": [],
    }), encoding="utf-8")
    check(audit.load_artifacts(d2) == {},
          "a sidecar alone is not an artifact -- an unaudited version stays unaudited")

    # A file that claims an id and a version but is not named for either is surfaced
    # rather than skipped: a misnamed real artifact is a defect, and must not read as an
    # absent one.
    d3 = Path(tempfile.mkdtemp())
    (d3 / "backup-of-the-audit.json").write_text(json.dumps(real_art), encoding="utf-8")
    check(any("!misnamed" in str(k[1]) for k in audit.load_artifacts(d3)),
          "a misnamed artifact is reported, not silently ignored")

    print(f"\n{PASSED} checks passed, {FAILED} failed.")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
