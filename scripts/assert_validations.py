#!/usr/bin/env python3
"""
Reopen the saved workbook from disk and assert every declared dropdown validation is
present and correctly bound.

WHY THIS IS A SEPARATE SCRIPT AND NOT A FLAG ON THE MIGRATION
--------------------------------------------------------------
The migration knows what it *set*. That is not the question. All nine of this workbook's
dropdown validations silently vanished between 2026-08-22 and 2026-08-24 and nobody
noticed for four days, so no controlled vocabulary was enforced on anything written in
that window (convention 23). A process that reports "we set the validations" would have
printed a clean line every single day of that outage.

So this loads the file fresh -- not the in-memory workbook the writer just held -- and
checks three things per validation:

  1. a dataValidation element exists for that sheet/column
  2. its `sqref` is exactly the expected range (a binding that stops short silently
     accepts anything past its last bound row)
  3. its `formula1` points at the expected Lookups column *and* that column actually
     holds values (a range bound to an empty column is a dropdown with no options,
     which is indistinguishable from no dropdown at the point of data entry)

Check 3 is the one a "did we set it" test cannot do at all.

  python scripts/assert_validations.py            # exit 1 on any failure
  python scripts/assert_validations.py --quiet
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from scripts.migrate_schema import (VALIDATIONS, LOOKUPS_CEILING,  # noqa: E402
                                    NEW_VOCABULARIES, VOCAB_ADDITIONS)

DEFAULT_DB = ROOT / "data" / "market_intel_db.xlsx"


def column_of(sqref: str) -> set[str]:
    return {str(rng).split(":")[0].lstrip("$").rstrip("0123456789")
            for rng in str(sqref).split()}


def assert_validations(path: Path, quiet: bool = False) -> list[str]:
    """Return a list of failure strings. Empty means every binding is correct."""
    wb = openpyxl.load_workbook(path)
    failures: list[str] = []

    # Which Lookups columns actually hold values, read from the reopened file.
    lookups = wb["Lookups"]
    populated: dict[str, int] = {}
    for c in range(1, lookups.max_column + 1):
        letter = openpyxl.utils.get_column_letter(c)
        n = sum(1 for r in range(2, lookups.max_row + 1)
                if lookups.cell(r, c).value not in (None, ""))
        populated[letter] = n

    for sheet, col, src, ceiling in VALIDATIONS:
        want_sqref = f"{col}2:{col}{ceiling}"
        want_f1 = f"Lookups!${src}${2}:${src}${LOOKUPS_CEILING}"
        label = f"{sheet}!{col} -> Lookups!{src}"

        if sheet not in wb.sheetnames:
            failures.append(f"{label}: sheet {sheet} does not exist")
            continue
        ws = wb[sheet]

        match = None
        for dv in ws.data_validations.dataValidation:
            if column_of(dv.sqref) == {col}:
                match = dv
                break

        if match is None:
            failures.append(f"{label}: NO dataValidation bound to column {col}")
            continue
        if str(match.sqref) != want_sqref:
            failures.append(f"{label}: sqref is {str(match.sqref)!r}, expected "
                            f"{want_sqref!r}")
        if match.formula1 != want_f1:
            failures.append(f"{label}: formula1 is {match.formula1!r}, expected "
                            f"{want_f1!r}")
        # The VALUES, not just the binding. Until 2026-09-01 this script asserted that a
        # dropdown was bound to a Lookups column and that the column was non-empty, but
        # never that the column held what the repo declares -- so the declared vocabulary
        # was decorative and a value added to one side alone went unnoticed. That is the
        # same gap convention 23 was written about, one level down: the binding was
        # declared and checked, the vocabulary behind it was declared and assumed.
        #
        # The expectation is the union of the two declaration sites, because they are
        # applied by different migration steps: NEW_VOCABULARIES seeds a column that does
        # not exist yet and is skipped once it does, VOCAB_ADDITIONS appends missing
        # values to a column that already exists. A full migration runs both, so the
        # workbook should equal the union. Order is not asserted -- Excel offers values in
        # column order and additions legitimately land at the end.
        declared = list(dict.fromkeys(list(NEW_VOCABULARIES.get(src, (None, []))[1])
                                      + list(VOCAB_ADDITIONS.get(src, (None, []))[1])))
        if declared:
            have = [lookups.cell(r, openpyxl.utils.column_index_from_string(src)).value
                    for r in range(2, LOOKUPS_CEILING + 1)]
            have = [str(v) for v in have if v not in (None, "")]
            missing = [v for v in declared if v not in have]
            undeclared = [v for v in have if v not in declared]
            if missing:
                failures.append(f"{label}: Lookups column {src} is MISSING declared "
                                f"value(s) {missing} -- the repo declares a vocabulary "
                                f"the workbook cannot offer")
            if undeclared:
                failures.append(f"{label}: Lookups column {src} holds undeclared "
                                f"value(s) {undeclared} -- a value nothing in the repo "
                                f"knows about is enforced on writers")

        if populated.get(src, 0) == 0:
            failures.append(f"{label}: Lookups column {src} holds no values -- the "
                            f"dropdown would offer nothing")
        # A vocabulary that outgrew the Lookups source range loses its tail silently.
        if populated.get(src, 0) > LOOKUPS_CEILING - 1:
            failures.append(f"{label}: Lookups column {src} holds "
                            f"{populated[src]} values but the source range stops at row "
                            f"{LOOKUPS_CEILING}")
        elif not quiet and not failures:
            pass

    if not quiet:
        print(f"workbook (reopened from disk): {path}")
        print(f"declared validations: {len(VALIDATIONS)}")
        for sheet, col, src, ceiling in VALIDATIONS:
            bad = [f for f in failures if f.startswith(f"{sheet}!{col} ")]
            mark = "FAIL" if bad else "ok  "
            print(f"  {mark}  {sheet}!{col}2:{col}{ceiling} <- Lookups!{src} "
                  f"({populated.get(src, 0)} values)")
    return failures


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=str(DEFAULT_DB))
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    failures = assert_validations(Path(args.db), quiet=args.quiet)
    if failures:
        print()
        print(f"VALIDATION READ-BACK FAILED -- {len(failures)} problem(s):")
        for f in failures:
            print(f"  !!  {f}")
        return 1
    print("\nall declared validations present and correctly bound in the saved file")
    return 0


if __name__ == "__main__":
    sys.exit(main())
