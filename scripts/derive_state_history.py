#!/usr/bin/env python3
"""
Derive Company_State_History (the temporal schema's composed series) and append it.

    python scripts/derive_state_history.py              # dry run: derive, summarize, diff
    python scripts/derive_state_history.py --apply      # append as a new derivation
    python scripts/derive_state_history.py --csv out.csv

The derivation itself is core/composition.py; this script only loads the workbook, runs
it, reports what it would write, compares it with the latest derivation already on the
sheet, and -- on --apply -- appends the rows under a fresh derivation_id through
core/db.py::append_state_history (append-only; a derivation is never rewritten).

The dry run is the default and prints the full breakdown by status, reason, theme and
bucket, plus the row-level delta against the previous derivation, because an immutable
table should be read before it is written to.
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core import composition as comp  # noqa: E402
from core.db import MarketIntelDB, WORKBOOK_PATH  # noqa: E402
from core.workbook_backup import backup_workbook  # noqa: E402


def previous_rows(wb, derivation_id: str) -> dict:
    ws = wb["Company_State_History"]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    out = {}
    for r in it:
        if not r or r[0] is None:
            continue
        d = dict(zip(headers, r))
        if str(d.get("derivation_id")) == derivation_id:
            out[(d["company_id"], d["theme_id"], d["bucket_id"])] = d
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=str(WORKBOOK_PATH))
    ap.add_argument("--apply", action="store_true", help="append the derivation")
    ap.add_argument("--force", action="store_true",
                    help="append even if identical to the latest derivation")
    ap.add_argument("--csv", help="write the derived rows to this CSV as well")
    ap.add_argument("--as-of", help="derive as of this date (YYYY-MM-DD); default today")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.db, read_only=True, data_only=True)
    if "Company_State_History" not in wb.sheetnames:
        print("ABORT: Company_State_History does not exist; run scripts/migrate_schema.py --apply")
        return 1
    inputs = comp.load_inputs(wb)
    derived_at = _dt.date.fromisoformat(args.as_of) if args.as_of else _dt.date.today()

    # provisional id for the dry run; the writer allocates the real one on --apply
    ws = wb["Company_State_History"]
    existing = sorted({str(r[1]) for r in ws.iter_rows(min_row=2, values_only=True)
                       if r and r[1]})
    next_id = f"DR-{(int(existing[-1][3:]) + 1) if existing else 1:04d}"
    rows = comp.derive(inputs, derived_at=derived_at, derivation_id=next_id)

    s = comp.summarize(rows)
    print(f"{next_id} derived as of {derived_at}  ({comp.DERIVATION_VERSION}, "
          f"{comp.BUCKET_GRAIN} buckets)")
    print(f"  companies {len(inputs.companies)} x themes {len(inputs.themes)} x buckets "
          f"{len(s['by_bucket'])} = {s['rows']} rows")
    print("  reach in force (weakest realized reach of each harness's primary sources):")
    for hid, r in sorted(inputs.harness_reach.items()):
        print(f"    {hid:22s} {str(r['realized_reach']):13s} months={r['realized_reach_months']} "
              f"from={r['realized_reach_effective_from']}")
    print("  by status:", dict(s["by_status"]))
    print("  by reason:", dict(s["by_reason"]))
    ev = s["evidence"]
    print(f"  evidence cited (observation x bucket): {ev['total']} total, {ev['valid']} valid, {ev['invalid']} invalid "
          f"-- invalid excluded in {ev['buckets_with_invalid_excluded']} bucket(s); {len(inputs.invalid)} "
          f"observation(s) recorded invalid")
    print("  by bucket:")
    for b, c in sorted(s["by_bucket"].items()):
        print(f"    {b}: {dict(c)}")
    print("  by theme:")
    for t, c in s["by_theme"].items():
        print(f"    {t:32s} {dict(c)}")

    if existing:
        prev = previous_rows(wb, existing[-1])
        # 2026-09-15: the evidence a row rests on is compared too, so a change in what supports a bucket (an
        # observation recorded invalid, say) is seen move even when the status does not.
        changed, keys = 0, ("status", "reason", "organizational_state",
                            "min_retrospective_reach", "definition_hash")
        evidence_only = 0
        for r in rows:
            p = prev.get((r["company_id"], r["theme_id"], r["bucket_id"]))
            if p is None or any(str(p.get(k)) != str(r.get(k)) for k in keys):
                changed += 1
            elif any(str(p.get(k) or "") != str(r.get(k) or "") for k in ("supporting_observation_ids", "evidence_count")):
                changed += 1
                evidence_only += 1
        new_buckets = {r["bucket_id"] for r in rows} - {k[2] for k in prev}
        print(f"  vs {existing[-1]}: {changed} row(s) differ or are new ({evidence_only} differ only in supporting "
              f"evidence); new buckets {sorted(new_buckets) or 'none'}")
        if not changed and not args.force and args.apply:
            print("  identical to the latest derivation -- not appended (use --force)")
            return 0
    else:
        print("  no prior derivation on the sheet")

    if args.csv:
        with open(args.csv, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=comp.STATE_HISTORY_COLUMNS)
            w.writeheader()
            w.writerows(rows)
        print(f"  wrote {args.csv}")

    if not args.apply:
        print("\n  DRY RUN -- nothing written. Re-run with --apply to append.")
        return 0

    backup, pruned = backup_workbook(args.db)
    print(f"  backup {backup.name}" + (f"; pruned {[p.name for p in pruned]}" if pruned else ""))
    db = MarketIntelDB(args.db)
    ids = db.append_state_history(rows)
    db.save()
    print(f"\nappended {len(ids)} rows as {next_id} ({ids[0]} .. {ids[-1]})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
