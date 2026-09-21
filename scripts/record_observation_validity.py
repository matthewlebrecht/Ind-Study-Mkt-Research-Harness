#!/usr/bin/env python3
"""
Restore hard-deleted observations exactly and record their validity -- one named batch at a time.

Matthew Lebrecht's standing rule (2026-09-15), made retroactive the same day: no observation is ever hard-deleted.
Rows deleted before the rule come back through core/db.py::restore_observation (exact values, own id, original sheet
position) and their invalidity is recorded in Observation_Validity_History (core/validity.py).

BATCHES
  o00303-2026-09-15       O00303 (applied 2026-09-15, commit 0b1628f; kept as the record of what ran)
  front-line-2026-09-20  the three rows citing the Front Line Power Construction press release:
                          O00530 (its first determination) plus O00350 and O00531, whose
                          extraction-defect status is superseded by the true reason, wrong entity.
  hr0083-overstated-2026-09-17
                          the 8 rows of HR-0083's 97 whose recorded status names the wrong cause: 6
                          are wrong-entity rows (the article never names the company) and 2 are
                          McCarthy rows that aged out of the five-year window. Corrected by
                          SUPERSESSION; the rows stay invalid.
  retroactive-2026-09-15  the nine other deleted observations whose claim has no live row: the six page-furniture
                          rows deleted with O00303, O00604 (unsupported sidebar teaser), O00612 (duplicate of O00611)
                          and O00564 (a quarantined row no longer produced)

The 23 other retired ids are claims still live under a new id (pre-convention-43 renumbering): the evidence was never
lost, and restoring the old rows would give one claim two live ids, so they are not in any batch.

Guards, all before anything is written: each snapshot must equal the row as committed at its source commit, and
(for rows removed in an identifiable commit) be absent from the next one; the id must be retired under the same
natural key (restore_observation re-checks); after the batch every OTHER Observations row must be unchanged.

    python scripts/record_observation_validity.py --batch retroactive-2026-09-15            # dry run on a copy
    python scripts/record_observation_validity.py --batch retroactive-2026-09-15 --apply    # write (all or nothing)
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core import validity  # noqa: E402
from core.db import OBSERVATION_COLUMNS, MarketIntelDB  # noqa: E402
from core.workbook_backup import backup_workbook  # noqa: E402

DELETIONS = "harness_output/audits/DELETIONS_2026-09-15_firstparty_page_furniture.json"
WORDS = "Oh yeah agreed they're all unsupported and should be removed"
SHEET_REF = "REVIEW_QUEUE_2026-09-15_firstparty_page_furniture.md"
RULE = ("Restored exactly under its own id and recorded here instead of staying deleted, under Matthew Lebrecht's "
        "standing rule of 2026-09-15 that no observation is ever hard-deleted, made retroactive the same day.")


def furniture_basis(oid: str, root: Path) -> str:
    ev = json.loads((root / DELETIONS).read_text(encoding="utf-8"))["evidence"][oid]
    lines = "; ".join(f"page line {h['line']} ({h['section_above']}): {h['text'][:90]}" for h in ev["furniture_lines"])
    a0, a1 = ev["article_span"]
    return (f"Theme matched page furniture: every matched-term line on the cached page is outside the article "
            f"(article lines {a0}-{a1}) -- {lines}. Reviewer verdict unsupported, Matthew Lebrecht 2026-09-15: "
            f"\"{WORDS}\". Sheet {SHEET_REF}.")


def _furniture(oid: str) -> dict:
    return {"observation_id": oid, "source": "deletions", "source_commit": "b21a24c", "removed_in": "9641128",
            "validity_status": "invalidated_extraction_defect", "as_of_date": "2026-09-15",
            "determined_by": "Matthew Lebrecht", "basis": None,
            "notes": f"Hard-deleted in commit 9641128 with O00303 (convention 32, delete_by_reviewer_verdict). {RULE}"}


def _cookie(oid: str, url: str, lines: str, banner: int, copyright_line: int, extra: str) -> dict:
    """Item 7 (Matthew Lebrecht, 2026-09-15): a Caddell announcement whose single matched term is in the cookie banner."""
    page = "harness_output/H-FIRSTPARTY-01/raw/pages/2026-08-31/p_caddell.com_" + url.rstrip("/").rsplit("/", 1)[-1] + ".txt"
    return {"observation_id": oid, "source": "live",
            "expect": {"company_id": "A076", "harness_id": "H-FIRSTPARTY-01", "topic": "data_analytics_ai", "source_url": url},
            "validity_status": "invalidated_extraction_defect", "as_of_date": "2026-09-15",
            "determined_by": "Matthew Lebrecht",
            "basis": (f"Theme matched page furniture: the row's single matched term, 'analytics', occurs only in the site's "
                      f"GDPR Cookie Consent banner -- page lines {lines} ('cookielawinfo-checkbox-analytics' and the "
                      f"'Analytics' cookie category), all after the banner begins at line {banner} and after the page's "
                      f"copyright line ({copyright_line}); it occurs nowhere in the announcement (title at line 60). "
                      f"Cached page {page}, lines as rendered by harnesses/h_execid_01/extract.py::html_lines. Matthew "
                      f"Lebrecht, 2026-09-15 (item 7): the cookie-banner rows O00539, O00542 and O00544 are removed, "
                      f"i.e. recorded invalid."),
            "notes": f"{extra}Determination on a row that exists; the row is unchanged (convention 45)."}


# ---- hr0083-overstated-2026-09-17 -------------------------------------------------------
# HR-0083 recorded every row H-FIRSTPARTY-01 v1.3 stopped producing as invalidated_extraction_defect,
# on Matthew Lebrecht's instruction (2026-09-15, item 22). For 89 of the 97 that IS the reason. These
# eight were dropped for other reasons and their matched terms are STILL in the article body
# (core/tests/test_firstparty_body.py section 9 pins the set), so the recorded status names the wrong
# cause -- in two OPPOSITE directions. Diagnosis:
# docs/diagnostics/hr0083_overstated_statuses_2026-09-17.md.
# Every count below was read from the cached page in core/tests/fixtures/firstparty_pages/ on
# 2026-09-17, and case is stated where it matters: the near-miss names differ only in case and spacing.
_HR0083_SHEET = "docs/diagnostics/hr0083_overstated_statuses_2026-09-17.md"
_HR0083_NOTE = ("Supersedes the determination HR-0083 wrote. The row is unchanged and stays invalid "
                "(convention 45; a determination is permanent and none reinstates) -- only the recorded "
                "REASON is corrected. Matthew Lebrecht, 2026-09-17.")


def _wrong_entity(oid, cid, topic, url, article, evidence):
    return {"observation_id": oid, "source": "live",
            "expect": {"company_id": cid, "harness_id": "H-FIRSTPARTY-01", "topic": topic,
                       "source_url": url},
            "validity_status": "invalidated_wrong_entity", "as_of_date": "2026-09-17",
            "determined_by": "Matthew Lebrecht",
            "basis": (f"Wrong entity, not an extraction defect: the article is about {article}, and "
                      f"{evidence} H-FIRSTPARTY-01 v1.3's body-scoped identity test refused the row for "
                      f"this reason. HR-0083 recorded it as invalidated_extraction_defect, which says the "
                      f"claim was read off the wrong part of the RIGHT page, when the page is about "
                      f"another company altogether. Counts read from the cached page "
                      f"(core/tests/fixtures/firstparty_pages/) on 2026-09-17. Sheet {_HR0083_SHEET}."),
            "notes": _HR0083_NOTE}


def _aged_out(oid, cid, topic, url):
    return {"observation_id": oid, "source": "live",
            "expect": {"company_id": cid, "harness_id": "H-FIRSTPARTY-01", "topic": topic,
                       "source_url": url},
            "validity_status": "invalidated_not_reproduced", "as_of_date": "2026-09-17",
            "determined_by": "Matthew Lebrecht",
            "basis": ("Not an extraction defect: the company is right, the article is about it (the page "
                      "names 'McCarthy Holdings' 6 times and 'McCarthy' 13), and the matched terms are in "
                      "the article body. The row was dropped because the article AGED OUT of the five-year "
                      "currency window (core/windows.py) between the v1.3 dry run and the run: published "
                      "2021-09-09, it passed 1,825 days on 2026-09-08 and was 1,832 days old when HR-0083 "
                      "ran on 2026-09-15. Nothing about the claim was ever defective. Recorded "
                      "invalidated_not_reproduced -- the accurate statement available in a vocabulary that "
                      "has no status for evidence which aged out; Matthew Lebrecht chose this over adding "
                      f"one (2026-09-17). Sheet {_HR0083_SHEET}."),
            "notes": _HR0083_NOTE}


BATCHES = {
    "front-line-2026-09-20": {
        "record": "harness_output/audits/VALIDITY_2026-09-20_front_line.json",
        "rows": [
            {"observation_id": "O00530", "source": "live",
             "expect": {"company_id": "A073", "harness_id": "H-FIRSTPARTY-01",
                        "topic": "digital_transformation_process", "source_url": 'https://www.prnewswire.com/news-releases/front-line-power-construction-announces-strategic-investment-from-ariel-alternatives-302691595.html'},
             "validity_status": "invalidated_wrong_entity", "as_of_date": "2026-09-20",
             "determined_by": "Matthew Lebrecht",
             "basis": 'Identity: the press release (PR Newswire, 2026-02-18, \'Front Line Power Construction Announces Strategic Investment from Ariel Alternatives\') is about FRONT LINE POWER CONSTRUCTION, a Houston utility-infrastructure contractor, not A073 Power Construction, a general contractor in Chicago (powerconstruction.net). Counts read from the cached page (core/tests/fixtures/firstparty_pages/) on 2026-09-20: \'Front Line Power Construction\' 6 occurrences; \'Power Construction\' NOT preceded by \'Front Line\' 0; \'Houston\' 3 and \'Texas\' 2; \'Chicago\' 0, \'Illinois\' 0 and \'powerconstruction.net\' 0. The company matched only as a SUBSTRING of the other firm\'s name -- convention 11, a single matching fragment is not an identity. Reviewer verdict, Matthew Lebrecht 2026-09-20: "This is a wrong identity case, seems like it should be Front Line Power Construction." Decision sheet harness_output/audits/DECISION_O00530_front_line_identity.md.',
             "notes": "Determination on a row that exists; the row is unchanged (convention 45). A `supported` audit verdict applied to this row from a blanket approval on 2026-09-20 was RETRACTED the same day (scripts/retract_misattributed_verdict.py) because the question is identity, not strength; this determination is the reviewer's direct call on that question."},
            {"observation_id": "O00350", "source": "live",
             "expect": {"company_id": "A073", "harness_id": "H-FIRSTPARTY-01",
                        "topic": "data_analytics_ai", "source_url": 'https://www.prnewswire.com/news-releases/front-line-power-construction-announces-strategic-investment-from-ariel-alternatives-302691595.html'},
             "validity_status": "invalidated_wrong_entity", "as_of_date": "2026-09-20",
             "determined_by": "Matthew Lebrecht",
             "basis": 'Identity: the press release (PR Newswire, 2026-02-18, \'Front Line Power Construction Announces Strategic Investment from Ariel Alternatives\') is about FRONT LINE POWER CONSTRUCTION, a Houston utility-infrastructure contractor, not A073 Power Construction, a general contractor in Chicago (powerconstruction.net). Counts read from the cached page (core/tests/fixtures/firstparty_pages/) on 2026-09-20: \'Front Line Power Construction\' 6 occurrences; \'Power Construction\' NOT preceded by \'Front Line\' 0; \'Houston\' 3 and \'Texas\' 2; \'Chicago\' 0, \'Illinois\' 0 and \'powerconstruction.net\' 0. The company matched only as a SUBSTRING of the other firm\'s name -- convention 11, a single matching fragment is not an identity. Reviewer verdict, Matthew Lebrecht 2026-09-20: "This is a wrong identity case, seems like it should be Front Line Power Construction." Decision sheet harness_output/audits/DECISION_O00530_front_line_identity.md.',
             "notes": 'Supersedes the invalidated_extraction_defect this row received from the HR-0083 run on 2026-09-15. The row stays invalid and unchanged -- only the recorded REASON moves, from extraction to identity, which is the same correction applied to eight other HR-0083 rows on 2026-09-17. It cites the same press release as O00530 and fails for the same reason.'},
            {"observation_id": "O00531", "source": "live",
             "expect": {"company_id": "A073", "harness_id": "H-FIRSTPARTY-01",
                        "topic": "workforce_enablement", "source_url": 'https://www.prnewswire.com/news-releases/front-line-power-construction-announces-strategic-investment-from-ariel-alternatives-302691595.html'},
             "validity_status": "invalidated_wrong_entity", "as_of_date": "2026-09-20",
             "determined_by": "Matthew Lebrecht",
             "basis": 'Identity: the press release (PR Newswire, 2026-02-18, \'Front Line Power Construction Announces Strategic Investment from Ariel Alternatives\') is about FRONT LINE POWER CONSTRUCTION, a Houston utility-infrastructure contractor, not A073 Power Construction, a general contractor in Chicago (powerconstruction.net). Counts read from the cached page (core/tests/fixtures/firstparty_pages/) on 2026-09-20: \'Front Line Power Construction\' 6 occurrences; \'Power Construction\' NOT preceded by \'Front Line\' 0; \'Houston\' 3 and \'Texas\' 2; \'Chicago\' 0, \'Illinois\' 0 and \'powerconstruction.net\' 0. The company matched only as a SUBSTRING of the other firm\'s name -- convention 11, a single matching fragment is not an identity. Reviewer verdict, Matthew Lebrecht 2026-09-20: "This is a wrong identity case, seems like it should be Front Line Power Construction." Decision sheet harness_output/audits/DECISION_O00530_front_line_identity.md.',
             "notes": 'Supersedes the invalidated_extraction_defect this row received from the HR-0083 run on 2026-09-15. The row stays invalid and unchanged -- only the recorded REASON moves, from extraction to identity, which is the same correction applied to eight other HR-0083 rows on 2026-09-17. It cites the same press release as O00530 and fails for the same reason.'},
        ],
    },
    "hr0083-overstated-2026-09-17": {
        "record": "harness_output/audits/VALIDITY_2026-09-17_hr0083_overstated.json",
        "rows": [
            _wrong_entity("O00297", "A007", "data_analytics_ai", 'https://www.prnewswire.com/news-releases/lio-raises-30m-series-a-to-bring-agentic-ai-to-enterprise-procurement-302705236.html',
                          'a Series A funding round for a startup called Lio', "the page never names the company in any form: 'National Product Sales' 0 times, 'NPS' 0, 'Product Sales' 0, and 'National' only twice case-insensitively, in unrelated text."),
            _wrong_entity("O00298", "A007", "digital_transformation_process", 'https://www.prnewswire.com/news-releases/lio-raises-30m-series-a-to-bring-agentic-ai-to-enterprise-procurement-302705236.html',
                          'a Series A funding round for a startup called Lio', "the page never names the company in any form: 'National Product Sales' 0 times, 'NPS' 0, 'Product Sales' 0, and 'National' only twice case-insensitively, in unrelated text."),
            _wrong_entity("O00431", "A007", "workforce_enablement", 'https://www.prnewswire.com/news-releases/lio-raises-30m-series-a-to-bring-agentic-ai-to-enterprise-procurement-302705236.html',
                          'a Series A funding round for a startup called Lio', "the page never names the company in any form: 'National Product Sales' 0 times, 'NPS' 0, 'Product Sales' 0, and 'National' only twice case-insensitively, in unrelated text."),
            _wrong_entity("O00354", "A083", "data_analytics_ai", 'https://www.prweb.com/releases/jdm-technology-group-expands-construction-software-offering-acquiring-penta-technologies-and-struxi-302080723.html',
                          'JDM Technology Group acquiring Penta Technologies and STRUXI', "the page never names the company: 'PENTA Building Group' 0 times, 'PENTA Building' 0, 'Penta Building' 0. It carries 'Penta' 18 times, every one of them Penta Technologies -- a construction-software vendor being acquired, not The PENTA Building Group."),
            _wrong_entity("O00355", "A083", "workforce_enablement", 'https://www.prweb.com/releases/jdm-technology-group-expands-construction-software-offering-acquiring-penta-technologies-and-struxi-302080723.html',
                          'JDM Technology Group acquiring Penta Technologies and STRUXI', "the page never names the company: 'PENTA Building Group' 0 times, 'PENTA Building' 0, 'Penta Building' 0. It carries 'Penta' 18 times, every one of them Penta Technologies -- a construction-software vendor being acquired, not The PENTA Building Group."),
            _wrong_entity("O00359", "A090", "data_analytics_ai", 'https://www.businesswire.com/news/home/20260811241297/en/ESS-Tech-Inc.-Announces-Second-Quarter-2026-Financial-Results',
                          "ESS Tech, Inc.'s second-quarter 2026 financial results", "the page never names the company: 'ESS Companies' 0 times. It carries 'ESS' 72 times case-sensitively, every one of them ESS Tech, Inc., an energy-storage manufacturer."),
            _aged_out("O00324", "A038", "data_analytics_ai", 'https://www.prnewswire.com/news-releases/leading-us-commercial-construction-company-mccarthy-holdings-inc-selects-gep-smart-software-to-transform-and-unify-procurement-301372567.html'),
            _aged_out("O00325", "A038", "erp_core_systems", 'https://www.prnewswire.com/news-releases/leading-us-commercial-construction-company-mccarthy-holdings-inc-selects-gep-smart-software-to-transform-and-unify-procurement-301372567.html'),
        ],
    },
    "o00303-2026-09-15": {
        "record": "harness_output/audits/VALIDITY_2026-09-15_O00303.json",
        "rows": [_furniture("O00303") | {"notes": (
            "Hard-deleted in commit 9641128 under the rule then in force (convention 32, delete_by_reviewer_verdict). "
            "Restored exactly under its own id and recorded here instead, under Matthew Lebrecht's standing rule of "
            "2026-09-15 that no observation is ever hard-deleted. Deletion record DELETIONS_2026-09-15_firstparty_"
            "page_furniture.json; restoration record VALIDITY_2026-09-15_O00303.json.")}],
    },
    "retroactive-2026-09-15": {
        "record": "harness_output/audits/VALIDITY_2026-09-15_retroactive.json",
        "rows": [_furniture(o) for o in ("O00302", "O00339", "O00441", "O00484", "O00491", "O00529")] + [
            {"observation_id": "O00604", "source": "git", "source_commit": "f563dcd", "removed_in": "15f7c18",
             "validity_status": "invalidated_extraction_defect", "as_of_date": "2026-09-06",
             "determined_by": "Matthew Lebrecht",
             "basis": ("Audit verdict unsupported, Matthew Lebrecht 2026-09-06: the quoted span is the Construction Dive "
                       "sidebar teaser for a separate Novo Construction Q&A, swept into the last answer of a Walbridge "
                       "Q&A and attributed to Randy Abdallah -- the defect H-TRADEPRESS-01 v1.7's end-of-article cut "
                       "fixed."),
             "notes": f"Deleted on audit in commit 15f7c18 (session 15). {RULE}"},
            {"observation_id": "O00612", "source": "git", "source_commit": "769d6e2", "removed_in": "bf617a9",
             "validity_status": "invalidated_duplicate", "as_of_date": "2026-09-06",
             "determined_by": "Matthew Lebrecht",
             "basis": ("Audit 2026-09-06, Matthew Lebrecht: duplicate of O00611 -- the same Estes Express Lines breach "
                       "(breach 09/26/2023, reported 12/31/2023, same organisation string) surfaced by the second "
                       "California query term 'ESTES'. O00611 stands; H-BREACHPORTAL-01 v1.1 deduplicates California "
                       "hits."),
             "notes": f"Deleted on audit in commit bf617a9 (session 16 follow-up). {RULE}"},
            {"observation_id": "O00564", "source": "git", "source_commit": "f0411b1", "removed_in": "64f5b96",
             "validity_status": "invalidated_not_reproduced", "as_of_date": "2026-09-03",
             "determined_by": "H-EXECVOICE-01 (machine: not reproduced; no reviewer verdict)",
             "basis": ("H-EXECVOICE-01 v1.4 row, quarantined and never released, with no human verdict; last present at "
                       "f0411b1 (2026-09-03). The next committed run (64f5b96, session 10) did not propose the claim, and "
                       "delete-and-rewrite removed it. Its claim is live under no other id."),
             "notes": f"Removed by delete-and-rewrite in commit 64f5b96. {RULE}"},
        ],
    },
    "o00349-wrong-entity-2026-09-15": {
        "record": "harness_output/audits/VALIDITY_2026-09-15_O00349.json",
        "rows": [{
            "observation_id": "O00349", "source": "live",
            "expect": {"company_id": "A073", "harness_id": "H-FIRSTPARTY-01", "topic": "data_analytics_ai",
                       "source_url": "https://www.constructiondive.com/news/bechtel-kiewit-japan-ohio-power-generation-ai/815768/"},
            "validity_status": "invalidated_wrong_entity", "as_of_date": "2026-09-15",
            "determined_by": "Matthew Lebrecht",
            "basis": ("Identity: the article (Construction Dive, 2026-03-26, 'Bechtel, Kiewit tapped for Japan-backed $33B "
                      "Ohio power generation project') is about Bechtel and Kiewit. 'Power Construction' appears nowhere "
                      "on the cached page (0 occurrences in the raw HTML, case-insensitive); the company's name reduces to "
                      "two dictionary words that run through a construction article about a power project, which does "
                      "not establish identity (convention 31). The theme match itself is in the article. Reviewer "
                      "verdict wrong_entity, Matthew Lebrecht 2026-09-15: \"O00349: wrong_entity confirmed.\" Sheet "
                      "REVIEW_QUEUE_2026-09-15_firstparty_page_furniture.md (Separate: O00349)."),
            "notes": ("Determination on a row that exists; the row is unchanged (convention 45). O00529 cites the same "
                      "article and is already recorded invalidated_extraction_defect.")}],
    },
    "caddell-cookie-banner-2026-09-15": {
        "record": "harness_output/audits/VALIDITY_2026-09-15_caddell_cookie_banner.json",
        "rows": [_cookie(oid, url, lines, banner, copyright_line, notes) for oid, url, lines, banner, copyright_line, notes in (
            ("O00539", "https://caddell.com/caddell-construction-announces-key-c-suite-and-executive-promotions/",
             "125, 127, 149, 150", 112, 80,
             "Human-reviewed row: judged `supported` in the H-FIRSTPARTY-01 v1.2 random control (audit 2026-09-03, 29 of "
             "30 supported). That artifact is a historical record and is not rewritten; this later determination records "
             "what the cached page shows. "),
            ("O00542", "https://caddell.com/caddell-construction-announces-strategic-executive-hires-and-promotions/",
             "142, 144, 166, 167", 129, 97, ""),
            ("O00544", "https://caddell.com/caddell-construction-announces-three-new-presidents/",
             "134, 136, 158, 159", 121, 89, ""))],
    },
}


def _live_rows(db: MarketIntelDB) -> dict:
    it = db.wb["Observations"].iter_rows(values_only=True)
    h = list(next(it))
    return {r[0]: dict(zip(h, r)) for r in it if r and r[0]}


def committed_rows(commit: str) -> dict:
    blob = subprocess.run(["git", "show", f"{commit}:data/market_intel_db.xlsx"], cwd=ROOT, capture_output=True,
                          check=True).stdout
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    it = wb["Observations"].iter_rows(values_only=True)
    h = list(next(it))
    return {r[0]: dict(zip(h, r)) for r in it if r[0]}


def others_digest(wb, exclude: set[str]) -> str:
    h = hashlib.sha256()
    for row in wb["Observations"].iter_rows(values_only=True):
        if row[0] not in exclude:
            h.update(repr(row).encode("utf-8"))
    return h.hexdigest()


def run(db: MarketIntelDB, batch: str, root: Path = ROOT, today: str | None = None) -> dict:
    """Restore and record `batch` on `db` (in memory). Refuses before writing anything if a guard fails."""
    from core.db import today as _today
    today = today or _today()
    spec = BATCHES[batch]
    deleted_snaps = {s["observation_id"]: s for s in
                     json.loads((root / DELETIONS).read_text(encoding="utf-8"))["snapshots"]}
    cache: dict[str, dict] = {}

    def at(commit):
        if commit not in cache:
            cache[commit] = committed_rows(commit)
        return cache[commit]

    live_now = _live_rows(db)
    problems, snaps = [], {}
    for b in spec["rows"]:
        oid = b["observation_id"]
        if b["source"] == "live":
            # A determination on a row that exists: nothing to restore; the row must be the one the verdict names.
            row = live_now.get(oid)
            if row is None:
                problems.append(f"{oid}: no live Observations row")
            elif any(str(row.get(k) or "") != str(v) for k, v in b["expect"].items()):
                problems.append(f"{oid}: the live row does not match {b['expect']}")
            continue
        committed = at(b["source_commit"]).get(oid)
        if committed is None:
            problems.append(f"{oid}: not in the workbook at {b['source_commit']}")
            continue
        if b["source"] == "deletions":
            snap = deleted_snaps.get(oid)
            if snap is None or any(snap[c] != committed[c] for c in OBSERVATION_COLUMNS):
                problems.append(f"{oid}: deletion-record snapshot missing or differs from {b['source_commit']}")
                continue
        else:
            snap = {c: committed[c] for c in OBSERVATION_COLUMNS}
        if b.get("removed_in") and oid in at(b["removed_in"]):
            problems.append(f"{oid}: still present at {b['removed_in']}, so that is not where it was removed")
        snaps[oid] = snap
    if problems:
        raise SystemExit("REFUSED, nothing written: " + "; ".join(problems))
    ids = {b["observation_id"] for b in spec["rows"] if b["source"] != "live"}   # a live row must not change at all
    before = others_digest(db.wb, ids)
    restored = []
    for b in sorted((b for b in spec["rows"] if b["source"] != "live"), key=lambda b: b["observation_id"]):
        live = sorted(validity.live_observation_ids(db.wb))
        after = next((x for x in live if x > b["observation_id"]), None)
        restored.append(db.restore_observation(snaps[b["observation_id"]], reason=RULE, before_id=after))
    dets = []
    for b in spec["rows"]:
        basis = b["basis"] or furniture_basis(b["observation_id"], root)
        dets.append({"observation_id": b["observation_id"], "validity_status": b["validity_status"],
                     "as_of_date": b["as_of_date"], "determined_at": today, "determined_by": b["determined_by"],
                     "basis": basis, "notes": b["notes"]})
    report = validity.append_determinations(db.wb, dets)
    if others_digest(db.wb, ids) != before:
        raise SystemExit("REFUSED: an Observations row outside the batch changed")
    live = validity.live_observation_ids(db.wb)
    reg = validity.registry(db.wb)
    cur = validity.current_rows(validity.rows_from_sheet(db.wb[validity.SHEET]))
    citations = []
    for d in validity.derivation_citations(db.wb):
        states = validity.citation_states(d["cited"], live, reg, cur)
        if any(s != "valid" for s in states.values()):
            citations.append({k: d[k] for k in ("state_id", "derivation_id", "bucket_id", "company_id", "theme_key",
                                                "status")} | {"cited": states})
    return {"batch": batch, "restored": restored, "determinations": report, "citations": citations,
            "hard_deletions_after": validity.hard_deletions(db.wb)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--batch", required=True, choices=sorted(BATCHES))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    path = ROOT / "data" / "market_intel_db.xlsx"
    if args.apply:
        db = MarketIntelDB(path)
    else:
        tmp = Path(tempfile.mkdtemp()) / "dry.xlsx"
        shutil.copy(path, tmp)
        db = MarketIntelDB(tmp)
    res = run(db, args.batch)
    for r in res["restored"]:
        print(f"restored {r['observation_id']} at sheet row {r['sheet_row']}")
    print("determinations:", res["determinations"])
    print("hard deletions remaining:", res["hard_deletions_after"])
    print(f"Company_State_History rows citing anything not valid: {len(res['citations'])}")
    for c in res["citations"]:
        print(f"  {c['derivation_id']} {c['bucket_id']} {c['company_id']} {c['theme_key']} {c['status']}: {c['cited']}")
    if args.apply:
        backup_workbook(path)
        db.save()
        spec = BATCHES[args.batch]
        (ROOT / spec["record"]).write_text(json.dumps({
            "record": "restoration and validity determination (convention 45: no observation is hard-deleted)",
            "rule": "Matthew Lebrecht, 2026-09-15: no observation, from any harness, is ever hard-deleted again (retroactive)",
            "rows": spec["rows"], **res}, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
        print(f"saved; record {spec['record']}")
    else:
        print("DRY RUN (temporary copy) -- re-run with --apply to write.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
