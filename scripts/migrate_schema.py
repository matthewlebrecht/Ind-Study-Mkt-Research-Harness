#!/usr/bin/env python3
"""
Apply the Attempts audit schema to the workbook.

Implements `docs/archive/build_handoff_2026-08-27.md` steps 1-7, which in turn implement
`docs/schema/attempts_schema_spec.md` (Rev 5). Step 8 stays in `extend_validation.py`.

Idempotent by construction: every step checks for its own output before doing anything, so
re-running is a no-op. That matters because this runs against the system of record and the
handoff's verification step asks for a save/reload round-trip.

  python scripts/migrate_schema.py                # report only
  python scripts/migrate_schema.py --apply

WHY THIS ALSO CREATES VALIDATIONS RATHER THAN ONLY EXTENDING THEM
------------------------------------------------------------------
The handoff assumes the workbook's 9 dropdown validations exist and only need rebinding
past row 500. They do not exist. They are present in every archived copy up to
`backup-pre-v12` (2026-08-22 21:45) and absent from `backup-pre-v13` (2026-08-24 22:45)
onward, including the copy that was live when this ran.

An openpyxl load/save round-trip was ruled out as the cause -- these are standard
`dataValidation` elements, not the `x14` extension kind openpyxl warns about, and a
round-trip preserves all 9. Whatever removed them, the effect is that controlled
vocabularies have not actually been enforced on any row written since 2026-08-22, which is
the same silent-loss failure `CLAUDE.md` warns about for the 500-row binding, arrived at by
a different route.

So validations are declared here, in full, and created if missing. `VALIDATIONS` below is
the declarative expected set; `scripts/validate_repo_db.py` checks reality against it, so
a future disappearance is a test failure instead of a silence.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.db import is_human_authored  # noqa: E402
from core.topics import RETIRED_THEME_KEYS  # noqa: E402
from core.workbook_backup import backup_workbook  # noqa: E402

DEFAULT_DB = ROOT / "data" / "market_intel_db.xlsx"

# --------------------------------------------------------------------------------------
# Step 2 -- six new controlled vocabularies, appended to Lookups as columns J-O.
# Additive-only: retire a value by marking it inactive, never by deleting it (spec §5).
# --------------------------------------------------------------------------------------
NEW_VOCABULARIES = {
    # `qualification_status` predates this migration -- the column and its four values
    # were already in the workbook, so add_vocabularies() skips it (the header matches)
    # and this entry changes nothing for an existing book. It is declared anyway because
    # until 2026-09-01 the repo declared only the ONE value it added later
    # (`provider_benchmark`, in VOCAB_ADDITIONS) and none of the four it was added to,
    # which meant a fresh rebuild would have produced a column holding a single value.
    # Found by the vocabulary read-back added to assert_validations.py the same day.
    "F": ("qualification_status", [
        "qualified", "excluded", "pending_review", "soft_gate_exception",
    ]),
    # `industry` likewise predates the migration (three hand-set pilot values); declared
    # here so the read-back in assert_validations.py knows them, with session 15's six
    # verticals appended through VOCAB_ADDITIONS.
    "H": ("industry", ["manufacturing", "supply_chain_operations", "logistics"]),
    "J": ("attempt_outcome", [
        "covered", "absent_confirmed", "partial", "not_covered",
    ]),
    "K": ("failure_stage", [
        "discovery", "fetch", "entity_resolution", "extraction",
        "classification", "temporal", "governance",
    ]),
    # 17 values, in stage order, from spec §4's table. The handoff says "16 values" --
    # that count predates `suppressed_redundant`, added by Rev 4 §11 Q5. The table is
    # authoritative over the count; see MIGRATION_NOTES.md.
    "L": ("failure_category", [
        "no_api_available", "source_not_found",
        "access_blocked", "js_rendered_unreachable", "source_unavailable",
        "entity_no_candidate", "entity_below_threshold", "entity_ambiguous_multiple",
        "parse_failure", "content_unstructured",
        "false_positive_rejected", "classification_ambiguous",
        "source_drift_detected", "stale_beyond_threshold",
        "suppressed_by_cap", "write_conflict_post_review", "suppressed_redundant",
    ]),
    "M": ("fix_class", ["transient", "code_change", "source_limitation"]),
    "N": ("attempt_scope", ["scoped", "incidental"]),
    "O": ("harness_source_role", ["primary", "supporting", "resolution_only"]),
    # ---- audit gate, 2026-08-31 (session 3 task 1) ----
    # `publication_state` is deliberately NOT a `review_status` value. review_status is
    # the human's verdict on a claim; publication_state is the machine's statement about
    # whether the row may contribute to a coverage number. Overloading one column with
    # both would make "may a re-run overwrite this" undecidable, which is the exact
    # question the gate has to answer.
    "P": ("publication_state", ["quarantined", "released"]),
    "Q": ("audit_verdict", ["supported", "overgraded", "unsupported", "wrong_entity"]),
    "R": ("review_source", ["human", "machine"]),
    "S": ("publication_status", ["quarantined", "published"]),
    # ---- taxonomy dimension 7, signal_taxonomy.md §21 (2026-09-02) ----
    # Scored PER SOURCE, not per harness -- H-PROCUREMENT-01 alone spans two values --
    # which is why it lands on Harness_Sources rather than on a harness or a manifest.
    "V": ("retrospective_reach", ["current_only", "bounded", "archival"]),
    # ---- Company_State_History, session 13 (2026-09-05) ----
    # `status` is what a bucket says; `reason` is why. `null` carries the two split
    # not_instrumented reasons from the week-3 package §4 -- `theme_not_detectable`
    # (resolves by waiting) and `no_reach` (resolves by acquiring access) -- plus
    # `no_absence_license`, the §26 "formally untested" case: reach covered the bucket
    # but every covering instrument is IC1/IC2 or presence-only, so silence says nothing.
    "Y": ("observation_id_status", ["live", "retired"]),
    # ---- coherence framework, build handoff 2026-09-08 ----
    # `entity_type` partitions the reference table. `generative_force` rows are
    # reference-only and are never tagging options (handoff §1 note); `failure_family`
    # values are derived downstream at synthesis and are never written to
    # Observation_Coherence_Tags (§2). The tag table's dimension column is therefore
    # constrained to COH-D* by validate_repo_db check 11, not by this vocabulary.
    "Z": ("coherence_entity_type", ["dimension", "failure_family", "generative_force"]),
    "AA": ("coherence_valence", ["failure_candidate", "counterevidence",
                                 "successful_coherence_candidate", "not_applicable",
                                 "unknown_insufficient_evidence"]),
    "AB": ("coherence_pilot_cohort", ["over_firing", "comparison"]),
    # ---- SEC reporting status, Harness Advisor design (session 17 wrap-up, 2026-09-13) ----
    # Eight values, fixed by the design. `entity_unresolved` is IDENTITY doubt (a filer was
    # found and cannot be confirmed to be the company; no status is asserted);
    # `status_uncertain` is STATUS doubt (the filer is confirmed, its filings do not settle
    # which status applies). Definitions and the writer's enforcement: core/sec_status.py.
    "AC": ("sec_reporting_status", ["active_reporter", "form_d_only", "withdrawn_registration",
                                    "never_registered", "insider_only", "deregistered",
                                    "entity_unresolved", "status_uncertain"]),
    # ---- evidence directionality, build handoff 2026-09-15 (Signal Advisor, finalized) ----
    # Which side of a transaction the company sits on in the evidence. Optional and sparse:
    # an untagged observation is untagged, not not_applicable. core/directionality.py.
    "AD": ("evidence_directionality", ["buyer_side", "seller_side", "mixed", "not_applicable"]),
    # ---- role-classification reviews, Matthew Lebrecht 2026-09-15 (convention 44) ----
    # `correct` confirms the role under review; any other value names the role the reviewer
    # found instead, or that it could not be placed. core/role_review.py.
    "AE": ("role_review_verdict", ["correct", "buyer_acts", "buyer_articulates",
                                   "provider_market_responds", "other", "cant_tell"]),
    # ---- observation validity, Matthew Lebrecht 2026-09-15 (standing rule: no hard deletion) ----
    # Additive-only: other invalidation reasons are appended as values. core/validity.py.
    "AF": ("observation_validity_status", ["invalidated_extraction_defect"]),
    "W": ("state_status", ["observed", "absent", "null"]),
    "X": ("state_reason", ["observed", "absence_licensed_IC3", "absence_licensed_IC4",
                           "no_absence_license", "no_reach", "theme_not_detectable"]),
    # ---- taxonomy patch 2026-08-31 rev 2, §4.6 ----
    # `instrument_class` is NOT `signal_class` and must never be collapsed into it.
    # signal_class (locked 2026-08-27) says whether a signal type yields evidence or is a
    # prerequisite enabling other harnesses -- it is what lets H-EXECID-01 exist with a
    # null evidence family. instrument_class says whether the instrument could observe
    # the thing at all, and governs what an ABSENCE licenses (§4.6.1): IC1 absence
    # licenses nothing, IC4 absence is moderate evidence. Rev 1 of the patch tried to
    # merge them, which is why the distinction is written down twice.
    "T": ("instrument_class", ["IC1", "IC2", "IC3", "IC4"]),
    "U": ("signal_type_status", ["active", "routing_only", "inactive"]),
}

# Values appended to vocabularies that already exist. Additive-only, same rule as the new
# ones: adding a value is cheap, renaming or removing one breaks version-over-version
# trend queries.
#
# `provider_benchmark` exists because the seller benchmark needs entities to hang
# Observations on, and Observations key to `company_id`. Providers therefore live in
# Companies under a `P0xx` prefix. They are NOT marked `excluded` -- that value means "was
# assessed against the buyer gate and failed it", and overloading it would corrupt any
# query counting excluded buyers. This keeps one entity table, so the Observations foreign
# key and the referential-integrity check stay simple, while leaving sellers trivially
# separable in every query.
VOCAB_ADDITIONS = {
    # ---- observation validity statuses, 2026-09-15 (convention 45, Matthew: retirement becomes invalidation) ----
    "AF": ("observation_validity_status", ["invalidated_not_reproduced", "invalidated_duplicate",
                                          "invalidated_wrong_entity"]),
    "F": ("qualification_status", ["provider_benchmark"]),
    # ---- session 15 (2026-09-06): `industry` gains H-TRADEPRESS-01's outlet verticals so
    # Companies.industry_primary, populated from archived NAICS codes by
    # scripts/enrich_industry.py, can drive that harness's outlet map for the whole
    # universe. Additive (convention 22); the three original values stay.
    "H": ("industry", ["construction", "trucking", "food", "medical", "energy", "consumer"]),
    # Closes coordination item (5) from the 2026-08-27 list. `suppressed_redundant` did
    # not distinguish "the natural key already holds this" (deterministic, cheap, always
    # right) from "a classifier judged this instance to add nothing new" (a judgment call
    # that can be wrong, and that the audit gate's suppression stratum has to be able to
    # sample). Additive-only per convention 22: the old value stays and is marked
    # superseded in the docs rather than deleted -- 0 rows carry it today, but a query
    # over historical Attempts must not break the day one does.
    "L": ("failure_category", ["suppressed_redundant_key", "suppressed_redundant_judged",
                               # session 10 (C6): generic-tier-only text, nothing written
                               "suppressed_generic_only"]),
    # `superseded`, 2026-09-01 (session 5 task 1). `quarantined` was holding two states
    # that route to different actions: "awaiting an audit" and "will never be audited,
    # because a later version re-ran the same scope and holds its output". The first is a
    # backlog item; the second is a decision, and conflating them means a version sits in
    # the audit queue forever with nothing anyone can do about it.
    #
    # It is deliberately NOT grandfathering. Grandfathering exempts builds that predate
    # the gate; H-TRADEPRESS-01 v1.1 is post-gate and auditable in principle, and reaching
    # for the exemption because its review sheet was clobbered is the way around a failing
    # audit that the gate text rules out. `superseded` asserts something narrower and
    # CHECKABLE: nothing of this version is live. validate_repo_db check 9 enforces that
    # literally -- a superseded version still holding observations fails, as does one
    # marked superseded on one run and published on another.
    #
    # Same shape as convention 39's suppressed_redundant split directly above, and as
    # tonight's access_blocked split: one flag holding several states that need different
    # actions, so the distinction is recorded rather than inferred.
    "S": ("publication_status", ["superseded"]),
    # ---- §22.1, 2026-09-02: `status` had no vocabulary for a type that is REAL but
    # produces nothing, which is why the Week 2 carried types were hard to resolve.
    #
    # `inactive` is NOT in §22.1's five, but convention 22 is additive-only and a value
    # is retired by marking it unused, never by deleting it -- a query over historical
    # Signal_Types rows must not break the day one carries it. So the column holds six
    # values and §22.1's five are the live set. Flagged rather than silently reconciled.
    #
    # The hard rule attached to this vocabulary is a ROLLUP filter, not a write filter:
    # attempts against an `access_bounded` or `out_of_theme` type still land in Attempts
    # and are excluded at the rollup by join. Implementing it as a write filter would
    # silently understate the coverage denominator -- the same reasoning that kept
    # `records_quarantined` off Harness_Runs.
    "U": ("signal_type_status", ["out_of_theme", "access_bounded", "reference"]),
}

# Lookups source ranges bind to row 50, not 20. `failure_category` already holds 17 of the
# 19 slots a `$X$2:$X$20` range would give it, and the vocabulary is explicitly
# additive-only -- a value past the source range silently never appears in any dropdown.
# Same failure as the 500-row binding, on a different axis (handoff §2).
LOOKUPS_CEILING = 50

# --------------------------------------------------------------------------------------
# Steps 3-6 -- new sheets. Column order is the schema contract, as with Observations.
# --------------------------------------------------------------------------------------
NEW_SHEETS = {
    # ---- coherence framework, build handoff 2026-09-08 (additive; Signal Advisor ruling
    # 2026-09-07/08, no Matthew sign-off required) ----
    #
    # Three tables, deliberately separate from Observations. The framework is explicitly
    # falsifiable and expected to churn, so a tag lives in its own table keyed BY
    # observation_id rather than as a column on the row -- an observation's identity
    # (convention 43) must not move when the framework does.
    "Coherence_Framework_Taxonomy": [
        "id", "framework_version", "label", "definition", "parent_group",
        "superseded_by", "effective_date", "entity_type",
    ],
    "Observation_Coherence_Tags": [
        "observation_id", "coherence_dimension_candidate", "coherence_valence",
        "framework_version", "tagged_at", "tagging_run_id",
    ],
    "Coherence_Pilot_Runs": [
        "pilot_id", "hypothesis", "company_id", "cohort", "systems_integration_count",
        "matched_to_company_id", "included_at", "framework_version",
    ],
    # Addendum 2026-09-08: the family-to-dimension mapping, normalized rather than a
    # delimited column on the taxonomy -- it is many-to-many, a delimited cell breaks the
    # join §2's synthesis needs, and it versions independently of the family's own row.
    # PK is (family_id, dimension_id, framework_version); no superseded_by, because a
    # mapping change is new rows at a new framework_version.
    "Coherence_Family_Dimensions": [
        "family_id", "dimension_id", "framework_version", "role_note",
    ],
    # Build handoff 2026-09-15 (Signal Advisor, finalized). The STANDING PATTERN (convention
    # 44): any optional, non-blocking classification of an observation lives in its own linked
    # table keyed BY observation_id, never as a column on Observations, whether or not its
    # vocabulary is expected to churn. Same convention 41 wall as Observation_Coherence_Tags
    # (validate_repo_db check 13). Writer: core/directionality.py::append_tags.
    "Observation_Directionality_Tags": [
        "observation_id", "evidence_directionality", "tagged_at", "tagging_run_id", "notes",
    ],
    # Matthew Lebrecht, 2026-09-15 (option A of the buyer_articulates role-review report).
    # Role-classification verdicts, the directionality table's shape (convention 44): keyed BY
    # observation_id, never written to Observations, so a role verdict re-clears neither identity
    # nor extraction and leaves review_source / review_status alone. Each row carries its
    # numeric sample basis and names the committed verdicts artifact it must agree with.
    # Same convention 41 wall (validate_repo_db check 14). Writer: core/role_review.py.
    "Observation_Role_Reviews": [
        "observation_id", "role_at_review", "role_verdict", "review_scope", "reviewer",
        "review_source", "reviewed_at", "review_run_id", "sample_design", "sample_n",
        "population_n", "stratum", "stratum_sampled_n", "stratum_population_n",
        "reviewer_words", "artifact", "notes",
    ],
    # Matthew Lebrecht, 2026-09-15: no observation is ever hard-deleted again. An invalid observation keeps its row
    # and id; each determination is a new row here. Append-only like SEC_Reporting_Status_History: nothing is
    # updated except the forward pointer `superseded_by`. Convention 41 wall (validate_repo_db check 15). Writer:
    # core/validity.py::append_determinations.
    "Observation_Validity_History": [
        "id", "observation_id", "validity_status", "as_of_date", "determined_at", "determined_by", "basis",
        "superseded_by", "notes",
    ],
    # Session 17 wrap-up (2026-09-13), Harness Advisor design, final. A DEDICATED table, not
    # an extension of Company_State_History. Append-only: every determination is a new row;
    # the only field ever set after insert is `superseded_by`, the forward pointer from a
    # historical row to the row that replaced it. `as_of_date` is what the evidence
    # establishes; `determined_at` is the load date. There is no Companies.public_private
    # column and there never will be: "is this company currently public" is DERIVED as the
    # latest non-superseded row per company with sec_reporting_status = active_reporter
    # (core/sec_status.py), and validate_repo_db check 12 fails on any such column.
    # Writer: core/db.py::append_sec_status.
    # Tab `SEC_Reporting_Status_History`: Excel caps sheet names at 31 characters; the design name is 36.
    "SEC_Reporting_Status_History": [
        "id", "company_id", "sec_reporting_status", "source_filing_type", "source_reference",
        "as_of_date", "determined_at", "determined_by", "superseded_by", "notes",
    ],
    # Session 14 (2026-09-06), convention 43: the observation-id registry. One row per id
    # ever assigned; an id belongs to its natural key forever and is reused if the same
    # claim is re-proposed after a delete-and-rewrite, never handed to a different claim.
    # current_id (Matthew Lebrecht, 2026-09-15, item 20) is appended by add_registry_lineage_column, not declared
    # here: add_sheets aborts on an existing sheet whose header differs, so a declared column could never be added.
    "Observation_Ids": [
        "observation_id", "natural_key", "company_id", "harness_id", "topic", "source_url",
        "first_assigned", "status", "retired_at", "retired_note",
    ],
    # Session 13 (2026-09-05). Append-only and immutable: one row per (derivation,
    # company, theme, bucket); every gating value stamped, nothing joined live. The
    # writer is core/db.py::append_state_history; the derivation is core/composition.py.
    "Company_State_History": [
        "state_id", "derivation_id", "derivation_version", "derived_at", "bucket_grain",
        "bucket_id", "bucket_start", "bucket_end", "company_id", "theme_id", "theme_key",
        "display_label", "definition_hash", "buyer_detectable_since", "status", "reason",
        "organizational_state", "min_retrospective_reach", "covering_instruments",
        "supporting_observation_ids", "evidence_count", "lowest_source_grade",
        "max_confidence", "staleness_years", "notes",
    ],
    "Attempts": [
        "attempt_id", "run_id", "harness_id", "harness_version", "company_id",
        "attempted_signal", "evidence_family", "outcome", "failure_stage",
        "failure_category", "failure_detail", "fix_class", "records_written",
        "records_reconciled", "output_sheet", "scope", "candidates_evaluated",
        "candidates_discarded", "attempt_timestamp", "source_url_attempted",
    ],
    "Signal_Types": [
        "signal_type_id", "signal_type_name", "signal_class", "evidence_family", "status",
    ],
    "Harness_Sources": [
        "harness_id", "harness_version", "source_id", "role", "notes",
    ],
    "Company_Executives": [
        "executive_id", "company_id", "full_name", "name_variants", "title",
        "title_normalized", "role_relevance", "source_url", "source_grade",
        "publication_date", "retrieval_date", "harness_id", "harness_version",
        "status", "superseded_by", "confidence_0_1", "review_status",
    ],
}

# Step 4 -- seed only what exists today. The taxonomy is not settled and
# `Attempts.attempted_signal` stays free text (no dropdown) until it is: binding a
# dropdown to an unfinished vocabulary would reject valid values from next week's
# harnesses. Signal Advisor owns the contents.
SIGNAL_TYPE_SEED = [
    ("ST-0001", "carrier_registry_status", "evidence",
     "10_logistics_supply_network", "active"),
    ("ST-0002", "job_posting", "evidence", "3_workforce_org_exhaust", "active"),
    ("ST-0003", "executive_identification", "prerequisite", None, "active"),
    ("ST-0004", "seller_service_page", "evidence", "seller_discourse", "active"),
    ("ST-0005", "osha_inspection", "evidence", "9_industrial_safety_environmental",
     "active"),
    ("ST-0006", "epa_echo_compliance", "evidence", "9_industrial_safety_environmental",
     "active"),
    ("ST-0007", "removed_page", "evidence",
     "18_historical_change_disappearing_evidence", "active"),
]

# --------------------------------------------------------------------------------------
# Audit gate, 2026-08-31 -- new columns on two existing sheets.
#
# Observations gains three. `publication_state` and `audit_verdict` are machine-authored
# and freely overwritable; `review_source` is the provenance flag that the rewritten
# never-overwrite rule keys on. The old rule keyed on the review_status *value*, which
# made a machine-set status indistinguishable from a human-set one -- fine while only
# humans ever set it, wrong the moment an audit starts writing verdicts.
#
# Harness_Runs gains two. There is deliberately no `records_quarantined` counter:
# exclusion is a filter at the rollup layer, never a mutated count. A quarantined run
# contributes neither numerator nor denominator, and its Attempts rows are excluded by
# join on run_id rather than by a duplicated status column that could drift out of sync.
# --------------------------------------------------------------------------------------
AUDIT_GATE_COLUMNS = {
    "Observations": ["publication_state", "audit_verdict", "review_source"],
    "Harness_Runs": ["publication_status", "records_excluded_by_audit"],
    # Taxonomy dimension only. Deliberately NOT added to Observations: it is a property
    # of the instrument, not of any individual claim, so storing it per row would be
    # 315 copies of a fact that belongs in one registry cell.
    "Signal_Types": ["instrument_class"],
}

# --------------------------------------------------------------------------------------
# Retrospective reach, signal_taxonomy.md §21 (2026-09-02)
# --------------------------------------------------------------------------------------
# Reach splits into NOMINAL (what the source could speak to) and REALIZED (what we can
# actually reach). §21.1: a paywalled trade archive is `archival` nominal and
# `current_only` realized, because the archive exists and we cannot enter it.
#
# The invariant that makes the split worth storing: **inference composes on
# `realized_reach` only.** `nominal_reach` is never an input to a negative inference --
# it feeds an acquisition-backlog report. Composing on nominal overclaims by exactly the
# size of the access gap.
#
# `realized_reach_effective_from` exists because realized reach CHANGES when an access
# gap closes, and a derived historical row must record the value it composed against
# rather than joining this table live. A live join looks correct in every test and
# silently rewrites history the first time we gain access to something.
REACH_COLUMNS = {
    "Harness_Sources": [
        "nominal_reach", "nominal_reach_months",
        "realized_reach", "realized_reach_months", "realized_reach_effective_from",
    ],
}

# Reach VALUES, per source (session 13, 2026-09-05). Provenance per row: taxonomy §21.2,
# the Harness Advisor reach-research ping of 2026-09-02, or a Code classification made
# while populating (marked "unverified" -- confirm before relying on it for a finding).
# Written only into EMPTY cells; a value someone has set by hand is never overwritten.
# (nominal_reach, nominal_months, realized_reach, realized_months, provenance note)
REACH_VALUES = {
    "SRC-0001": ("archival", None, "archival", None, "§21.2 dated newsroom archive (High)"),
    "SRC-0002": ("archival", None, "archival", None, "§21.2 PR-wire archive (High)"),
    "SRC-0003": ("archival", None, "current_only", None, "Bizjournals archive is paywalled: archival nominal, current_only realized (§21.1 pattern; Code 2026-09-05, unverified)"),
    "SRC-0009": ("current_only", None, "current_only", None, "live job listings, undated history (Code 2026-09-05)"),
    "SRC-0010": ("current_only", None, "current_only", None, "live job listings (Code 2026-09-05)"),
    "SRC-0013": ("current_only", None, "current_only", None, "aggregate rating is a present-state snapshot (Code 2026-09-05); source robots-refused, see access, not reach"),
    "SRC-0017": ("current_only", None, "current_only", None, "job postings naming systems: live listings (Code 2026-09-05)"),
    "SRC-0018": ("current_only", None, "current_only", None, "§21.2 vendor case studies: undated, silently revised (Medium); H-VENDOR-01, session 14"),
    "SRC-0023": ("archival", None, "archival", None, "ping 2026-09-02: Establishment Search since 1972 (High); harness queries from 2016 by scope, citation detail public ~5y post-closure"),
    "SRC-0024": ("archival", None, "bounded", 120, "ping 2026-09-02: reclassified bounded->archival with a ~10-year practical window on facility-level detail (High); realized recorded as bounded/120 for that window"),
    "SRC-0025": ("bounded", 24, "bounded", 24, "ping 2026-09-02: SMS 24-month rolling window confirmed (High)"),
    "SRC-0027": ("archival", None, "archival", None, "CPSC recall database is a dated permanent record (Code 2026-09-05, unverified)"),
    "SRC-0029": ("archival", None, "current_only", None, "§21.2 trade outlets: archival nominal, current_only realized (High)"),
    "SRC-0033": ("archival", None, "current_only", None, "dated uploads exist; reached through the live search index only (Code 2026-09-05, unverified)"),
    "SRC-0034": ("archival", None, "current_only", None, "dated episodes exist; reached through the live search index only (Code 2026-09-05, unverified)"),
    "SRC-0037": ("archival", None, "bounded", 0, "§21.1 ST-REMOVEDPAGE is bounded by OUR crawl window (self-referential): reach begins at the first capture and grows only with retention; months 0 at population"),
    "SRC-0038": ("current_only", None, "current_only", None, "provider service pages, undated and silently revised (§21.2 vendor-page reasoning; Code 2026-09-05)"),
    "SRC-0039": ("current_only", None, "current_only", None, "leadership pages show present roster only (Code 2026-09-05)"),
    "SRC-0040": ("current_only", None, "current_only", None, "search index: present state of the web (Code 2026-09-05)"),
    "SRC-0041": ("current_only", None, "current_only", None, "aggregate rating snapshot (Code 2026-09-05)"),
    "SRC-0042": ("archival", None, "current_only", None, "§21.2 trade outlets: archival nominal, current_only realized (High)"),
    "SRC-0043": ("archival", None, "current_only", None, "§21.2 trade outlets: archival nominal, current_only realized (High)"),
    "SRC-0044": ("archival", None, "archival", None, "openFDA recall/MAUDE: dated permanent records (Code 2026-09-05, unverified)"),
    "SRC-0045": ("archival", None, "archival", None, "openFDA enforcement: dated permanent records (Code 2026-09-05, unverified)"),
    "SRC-0046": ("archival", None, "archival", None, "NHTSA recalls/complaints: dated permanent records (Code 2026-09-05, unverified)"),
    "SRC-0047": ("archival", None, "archival", None, "§21.2 CourtListener dockets (High)"),
    "SRC-0048": ("archival", None, "archival", None, "§21.2 NLRB (High)"),
    "SRC-0049": ("current_only", None, "current_only", None, "§21.2 own-domain careers pages: current_only (High)"),
    "SRC-0050": ("archival", None, "archival", None, "ping 2026-09-02: CA portal is a permanent, indexable public record (High); list since 2012"),
    "SRC-0004": ("archival", None, "archival", None, "§21.2 SEC EDGAR / 8-K archival (High); submissions records reach back past the 2023-12-18 Item 1.05 rule for every scoped filer (Code 2026-09-13, probed live)"),
    "SRC-0011": ("archival", None, "archival", None, "state WARN lists (Code 2026-09-13, probed live): TX yearly listings 2020+, UT table 2009+; CA structured only for the current fiscal year, so CA realized reach is bounded to 2026-07-01 -- stated in every CA row"),
    "SRC-0021": ("archival", None, "archival", None, "Socrata permit datasets for Chicago, Seattle, Austin carry issue dates back to the mid-2000s (Code 2026-09-13, probed live); county-dependent elsewhere (§21.2)"),
    "SRC-0035": ("archival", None, "archival", None, "§21.2 Granicus/Legistar minutes archival (Medium); Seattle Legistar matters dated back to at least 2015 (Code 2026-09-13, probed live)"),
    "SRC-0052": ("current_only", None, "current_only", None, "own-domain served homepage: present state only (Code 2026-09-13)"),
    "SRC-0005": ("archival", None, "archival", None, "§21.2 USASpending awards: archival (High); API search floor 2007-10-01, harness reads a five-year window (2026-09-13)"),
    "SRC-0020": ("archival", None, "archival", None, "§21.2 USASpending awards: archival (High); API search floor 2007-10-01, harness reads a five-year window (2026-09-13)"),
    "SRC-0051": ("archival", None, "archival", None, "WA AG list is a permanent public record (Code 2026-09-06, probed live: 38 pages back to 2015; unverified beyond the probe)"),
}

# Step 7 -- Harness_Runs rollups, all computed over `scope = scoped` rows only.
# `coverage_rate` counts absent_confirmed as coverage; that is the point of the value.
HARNESS_RUN_ROLLUPS = [
    "attempts_total", "attempts_covered", "attempts_absent_confirmed",
    "attempts_partial", "attempts_not_covered", "coverage_rate",
]

# --------------------------------------------------------------------------------------
# The complete expected validation set: (sheet, column, lookups_column, ceiling).
# Observations/Companies restore the 9 that went missing; Attempts and Company_Executives
# are new. `Attempts` binds to 250,000 because it is append-only -- every re-run adds a
# full set of rows rather than reconciling (spec §12).
# --------------------------------------------------------------------------------------
VALIDATIONS = [
    ("Companies", "C", "H", 20_000),    # industry_primary
    ("Companies", "D", "H", 20_000),    # industry_secondary
    ("Companies", "K", "F", 20_000),    # qualification_status
    ("Observations", "C", "G", 20_000),  # evidence_family
    ("Observations", "D", "A", 20_000),  # evidence_role
    ("Observations", "F", "E", 20_000),  # organizational_state
    ("Observations", "G", "D", 20_000),  # signal_strength
    ("Observations", "M", "B", 20_000),  # source_grade
    ("Observations", "P", "C", 20_000),  # review_status
    ("Attempts", "G", "G", 250_000),    # evidence_family
    ("Attempts", "H", "J", 250_000),    # outcome
    ("Attempts", "I", "K", 250_000),    # failure_stage
    ("Attempts", "J", "L", 250_000),    # failure_category
    ("Attempts", "L", "M", 250_000),    # fix_class
    ("Attempts", "P", "N", 250_000),    # scope
    ("Harness_Sources", "D", "O", 20_000),        # role
    ("Company_Executives", "I", "B", 20_000),     # source_grade
    ("Company_Executives", "Q", "C", 20_000),     # review_status
    # ---- audit gate, 2026-08-31 ----
    ("Observations", "S", "P", 20_000),           # publication_state
    ("Observations", "T", "Q", 20_000),           # audit_verdict
    ("Observations", "U", "R", 20_000),           # review_source
    ("Harness_Runs", "R", "S", 20_000),           # publication_status
    # ---- taxonomy patch rev 2, 2026-08-31 ----
    ("Signal_Types", "E", "U", 20_000),           # status
    ("Signal_Types", "F", "T", 20_000),           # instrument_class
    ("Harness_Sources", "F", "V", 20_000),        # nominal_reach
    ("Harness_Sources", "H", "V", 20_000),        # realized_reach
    # ---- Company_State_History, session 13 (2026-09-05); append-only, so 250,000 ----
    ("Company_State_History", "O", "W", 250_000),  # status
    ("Company_State_History", "P", "X", 250_000),  # reason
    ("Company_State_History", "Q", "E", 250_000),  # organizational_state
    ("Company_State_History", "R", "V", 250_000),  # min_retrospective_reach
    # ---- Observation_Ids, session 14 ----
    ("Observation_Ids", "H", "Y", 250_000),        # status
    # ---- coherence framework, 2026-09-08 ----
    ("Coherence_Framework_Taxonomy", "H", "Z", 20_000),   # entity_type
    ("Observation_Coherence_Tags", "C", "AA", 250_000),   # coherence_valence
    ("Coherence_Pilot_Runs", "D", "AB", 20_000),          # cohort
    # ---- SEC reporting status history, session 17 wrap-up (2026-09-13) ----
    ("SEC_Reporting_Status_History", "C", "AC", 20_000),  # sec_reporting_status
    # ---- evidence directionality, build handoff 2026-09-15; a tag table, so 250,000 ----
    ("Observation_Directionality_Tags", "B", "AD", 250_000),  # evidence_directionality
    # ---- role-classification reviews, 2026-09-15; a review table keyed by observation, 250,000 ----
    ("Observation_Role_Reviews", "C", "AE", 250_000),  # role_verdict
    # ---- observation validity history, 2026-09-15; append-only, so 250,000 ----
    ("Observation_Validity_History", "C", "AF", 250_000),  # validity_status
]


class Migration:
    def __init__(self, path: Path):
        self.path = path
        self.wb = openpyxl.load_workbook(path)
        self.actions: list[str] = []
        self.skipped: list[str] = []

    def note(self, msg: str) -> None:
        self.actions.append(msg)

    def skip(self, msg: str) -> None:
        self.skipped.append(msg)

    # ---------- step 1 ----------
    def add_source_id(self) -> None:
        ws = self.wb["Source_Families"]
        if ws.cell(1, 1).value == "source_id":
            self.skip("Source_Families.source_id already present")
            return
        ws.insert_cols(1)
        ws.cell(1, 1).value = "source_id"
        n = 0
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value in (None, ""):
                continue
            n += 1
            ws.cell(r, 1).value = "SRC-%04d" % n
        self.note(f"Source_Families: added source_id as column A, populated SRC-0001..SRC-{n:04d}")

    # ---------- step 2 ----------
    def add_vocabularies(self) -> None:
        ws = self.wb["Lookups"]
        for col, (name, values) in NEW_VOCABULARIES.items():
            idx = openpyxl.utils.column_index_from_string(col)
            if ws.cell(1, idx).value == name:
                self.skip(f"Lookups!{col} ({name}) already present")
                continue
            if ws.cell(1, idx).value not in (None, ""):
                raise SystemExit(
                    f"ABORT: Lookups!{col}1 holds {ws.cell(1, idx).value!r}, expected empty "
                    f"or {name!r}. Refusing to overwrite an existing vocabulary."
                )
            ws.cell(1, idx).value = name
            for i, v in enumerate(values, start=2):
                ws.cell(i, idx).value = v
            self.note(f"Lookups: added {name} at column {col} ({len(values)} values)")

    def extend_vocabularies(self) -> None:
        ws = self.wb["Lookups"]
        for col, (name, values) in VOCAB_ADDITIONS.items():
            idx = openpyxl.utils.column_index_from_string(col)
            if ws.cell(1, idx).value != name:
                raise SystemExit(
                    f"ABORT: expected Lookups!{col}1 to be {name!r}, found "
                    f"{ws.cell(1, idx).value!r}"
                )
            present = {ws.cell(r, idx).value for r in range(2, ws.max_row + 1)}
            row = max((r for r in range(2, ws.max_row + 1)
                       if ws.cell(r, idx).value not in (None, "")), default=1)
            added = []
            for v in values:
                if v in present:
                    continue
                row += 1
                ws.cell(row, idx).value = v
                added.append(v)
            if added:
                self.note(f"Lookups.{name}: appended {added}")
            else:
                self.skip(f"Lookups.{name} already has {values}")

    # ---------- steps 3-6 ----------
    def add_sheets(self) -> None:
        for name, headers in NEW_SHEETS.items():
            if name in self.wb.sheetnames:
                ws = self.wb[name]
                actual = [ws.cell(1, i + 1).value for i in range(len(headers))]
                if actual != headers:
                    raise SystemExit(
                        f"ABORT: sheet {name} exists with different headers.\n"
                        f"  expected: {headers}\n  actual:   {actual}"
                    )
                self.skip(f"sheet {name} already present ({len(headers)} columns)")
                continue
            ws = self.wb.create_sheet(name)
            ws.append(headers)
            for i in range(1, len(headers) + 1):
                ws.cell(1, i).font = openpyxl.styles.Font(bold=True)
            ws.freeze_panes = "A2"
            self.note(f"created sheet {name} ({len(headers)} columns)")

    def seed_signal_types(self) -> None:
        ws = self.wb["Signal_Types"]
        existing = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        added = 0
        for row in SIGNAL_TYPE_SEED:
            if row[0] in existing:
                continue
            ws.append(list(row))
            added += 1
        if added:
            self.note(f"Signal_Types: seeded {added} registry row(s) "
                      f"({', '.join(r[1] for r in SIGNAL_TYPE_SEED)})")
        else:
            self.skip("Signal_Types already seeded")

    # ---------- step 7 ----------
    def add_run_rollups(self) -> None:
        ws = self.wb["Harness_Runs"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        added = []
        for col in HARNESS_RUN_ROLLUPS:
            if col in headers:
                continue
            headers.append(col)
            ws.cell(1, len(headers)).value = col
            ws.cell(1, len(headers)).font = openpyxl.styles.Font(bold=True)
            added.append(col)
        if added:
            self.note(f"Harness_Runs: appended rollup columns {added}")
        else:
            self.skip("Harness_Runs rollups already present")

    # ---------- audit gate (2026-08-31) ----------
    def add_audit_gate_columns(self) -> None:
        for sheet, cols in AUDIT_GATE_COLUMNS.items():
            ws = self.wb[sheet]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            added = []
            for col in cols:
                if col in headers:
                    continue
                headers.append(col)
                ws.cell(1, len(headers)).value = col
                ws.cell(1, len(headers)).font = openpyxl.styles.Font(bold=True)
                added.append(col)
            if added:
                self.note(f"{sheet}: appended audit-gate columns {added}")
            else:
                self.skip(f"{sheet} audit-gate columns already present")

    # ---------- retrospective reach (2026-09-02) ----------
    def add_reach_columns(self) -> None:
        for sheet, cols in REACH_COLUMNS.items():
            ws = self.wb[sheet]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            added = []
            for col in cols:
                if col in headers:
                    continue
                headers.append(col)
                ws.cell(1, len(headers)).value = col
                ws.cell(1, len(headers)).font = openpyxl.styles.Font(bold=True)
                added.append(col)
            if added:
                self.note(f"{sheet}: appended reach columns {added}")
            else:
                self.skip(f"{sheet} reach columns already present")

    def add_company_columns(self) -> None:
        """Session 15: `industry_source` beside `industry_primary`, on the same pattern as
        `employee_count_source` / `revenue_source` -- a value nobody can trace is a guess."""
        ws = self.wb["Companies"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        added = []
        # hq_city (2026-09-15, Matthew): the city beside hq_state, for the same-HQ-city identity
        # standard; populated by scripts/populate_hq_city.py, never guessed.
        for col in ("industry_source", "hq_city"):
            if col in headers:
                continue
            headers.append(col)
            ws.cell(1, len(headers)).value = col
            ws.cell(1, len(headers)).font = openpyxl.styles.Font(bold=True)
            added.append(col)
        if added:
            self.note(f"Companies: appended {added}")
        else:
            self.skip("Companies.industry_source and hq_city already present")

    def add_registry_lineage_column(self) -> None:
        """Item 20 (Matthew Lebrecht, 2026-09-15): `current_id` on Observation_Ids, after `retired_note` -- the live id
        a retired id's claim now carries. Filled by scripts/record_id_lineage.py, blank on live ids; check 15."""
        if "Observation_Ids" not in self.wb.sheetnames:
            return
        ws = self.wb["Observation_Ids"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if "current_id" in headers:
            self.skip("Observation_Ids.current_id already present")
            return
        col = headers.index("retired_note") + 2
        if ws.cell(1, col).value is not None:
            raise SystemExit(f"Observation_Ids column {col} is not empty ({ws.cell(1, col).value!r}); not adding current_id")
        ws.cell(1, col).value = "current_id"
        ws.cell(1, col).font = openpyxl.styles.Font(bold=True)
        self.note("Observation_Ids: appended current_id")

    def register_observation_ids(self) -> None:
        """Every live Observations row gets a registry entry (idempotent)."""
        from core.db import OBSERVATION_COLUMNS, natural_key_of
        ws = self.wb["Observation_Ids"]
        known = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
        obs = self.wb["Observations"]
        headers = [obs.cell(1, c).value for c in range(1, len(OBSERVATION_COLUMNS) + 1)]
        added = 0
        for r in range(2, obs.max_row + 1):
            oid = obs.cell(r, 1).value
            if not oid or oid in known:
                continue
            values = {h: obs.cell(r, i + 1).value for i, h in enumerate(headers)}
            key = natural_key_of(values)
            ws.append([oid, key, values.get("company_id"), values.get("harness_id"),
                       values.get("topic"), values.get("source_url"),
                       str(values.get("retrieval_date") or "")[:10] or None,
                       "live", None, None])
            added += 1
        if added:
            self.note(f"Observation_Ids: registered {added} live observation id(s)")
        else:
            self.skip("Observation_Ids already holds every live observation id")

    def populate_reach(self) -> None:
        """Write REACH_VALUES into EMPTY reach cells on Harness_Sources.

        `realized_reach_effective_from` is the harness's first run date from Harness_Runs:
        our realized reach through a source began when we first read it. A cell that
        already holds a value is left alone and reported, so a hand-set correction is
        never clobbered by a re-run of this migration.
        """
        ws = self.wb["Harness_Sources"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        hr = self.wb["Harness_Runs"]
        hh = [hr.cell(1, c).value for c in range(1, hr.max_column + 1)]
        first_run: dict[str, str] = {}
        for r in range(2, hr.max_row + 1):
            hid = hr.cell(r, hh.index("harness_id") + 1).value
            d = hr.cell(r, hh.index("date_run") + 1).value
            if hid and d:
                d = str(d)[:10]
                first_run[str(hid)] = min(first_run.get(str(hid), "9999"), d)
        written, kept, unknown = 0, 0, set()
        for r in range(2, ws.max_row + 1):
            sid = ws.cell(r, col["source_id"]).value
            if not sid:
                continue
            spec = REACH_VALUES.get(str(sid))
            if spec is None:
                unknown.add(str(sid))
                continue
            nom, nom_m, real, real_m, note = spec
            hid = str(ws.cell(r, col["harness_id"]).value)
            values = {"nominal_reach": nom, "nominal_reach_months": nom_m,
                      "realized_reach": real, "realized_reach_months": real_m,
                      "realized_reach_effective_from": first_run.get(hid)}
            touched = False
            for k, v in values.items():
                cell = ws.cell(r, col[k])
                if cell.value not in (None, ""):
                    kept += 1
                    continue
                if v is not None:
                    cell.value = v
                    touched = True
            if touched:
                written += 1
                ncell = ws.cell(r, col["notes"])
                stamp = f"reach: {note}"
                if stamp not in str(ncell.value or ""):
                    ncell.value = f"{ncell.value} | {stamp}" if ncell.value else stamp
        if written:
            self.note(f"Harness_Sources: reach populated on {written} row(s); "
                      f"{kept} pre-set cell(s) left untouched")
        else:
            self.skip(f"Harness_Sources reach already populated ({kept} cells kept)")
        if unknown:
            self.note(f"Harness_Sources: no REACH_VALUES entry for {sorted(unknown)} -- "
                      f"left blank, composition treats them as no reach")

    def backfill_audit_gate(self) -> None:
        """Fill the new columns for rows that predate the gate. Blanks only, never
        overwrite -- this runs against the system of record and has to stay idempotent.

        Everything already committed is grandfathered: it is `released` because it is
        already reported, and marking it otherwise would retroactively unpublish the
        Week 1-2 evidence base rather than gating what comes next. `review_source` is
        backfilled from review_status because that is exactly the information the old
        rule was inferring implicitly -- the 27 rows a human touched become `human`, the
        288 the harness wrote become `machine`.
        """
        ws = self.wb["Observations"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        i_status = headers.index("review_status") + 1
        i_pub = headers.index("publication_state") + 1
        i_src = headers.index("review_source") + 1
        pub_n = human_n = machine_n = 0
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                continue
            if ws.cell(r, i_pub).value in (None, ""):
                ws.cell(r, i_pub).value = "released"
                pub_n += 1
            if ws.cell(r, i_src).value in (None, ""):
                status = str(ws.cell(r, i_status).value or "").strip().lower()
                if status in ("accepted", "corrected", "rejected"):
                    ws.cell(r, i_src).value = "human"
                    human_n += 1
                else:
                    ws.cell(r, i_src).value = "machine"
                    machine_n += 1
        if pub_n or human_n or machine_n:
            self.note(f"Observations backfill: {pub_n} row(s) -> publication_state="
                      f"released; review_source set on {human_n} human + "
                      f"{machine_n} machine row(s)")
        else:
            self.skip("Observations audit-gate backfill already applied")

        ws = self.wb["Harness_Runs"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        i_pubs = headers.index("publication_status") + 1
        i_excl = headers.index("records_excluded_by_audit") + 1
        n = 0
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                continue
            if ws.cell(r, i_pubs).value in (None, ""):
                ws.cell(r, i_pubs).value = "published"
                ws.cell(r, i_excl).value = 0
                n += 1
        if n:
            self.note(f"Harness_Runs backfill: {n} pre-gate run(s) -> "
                      f"publication_status=published, records_excluded_by_audit=0")
        else:
            self.skip("Harness_Runs audit-gate backfill already applied")

    # ---- taxonomy patch rev 2: instrument_class for the pre-existing registry ----
    #
    # Assigned from §4.6's table. Two of these are judgment calls and are flagged as open
    # items for Signal Advisor rather than presented as settled:
    #
    #   removed_page (H-WAYBACK-01) -> IC3. The observable is the *act of removal*, which
    #     is revealed behaviour -- a byproduct of operating, not addressed to an audience.
    #     The archive is the instrument that makes it visible, not the thing being scored.
    #     The counter-argument is IC4: a company removes a page precisely to un-publish it,
    #     so the archive discloses over its preference.
    #   executive_identification (H-EXECID-01) -> IC1. It is a prerequisite, not evidence,
    #     but the class still does real work: a leadership page is curated, so the 47
    #     companies with no identified executive license NO inference that they have no
    #     relevant executives. That is exactly why H-EXECVOICE-01's gap is an H-EXECID-01
    #     coverage gap and not a finding about the companies.
    #
    # Blank is not the safe default here. A blank instrument_class reads downstream as
    # "unbiased", which is the one thing no instrument in this project is.
    INSTRUMENT_CLASSES = {
        "carrier_registry_status": "IC3",        # FMCSA -- named in §4.6 as IC3
        "job_posting": "IC3",                    # named in §4.6 as IC3
        "executive_identification": "IC1",       # curated leadership page
        "seller_service_page": "IC1",            # seller discourse benchmark, named IC1
        "osha_inspection": "IC4",                # named in §4.6 as IC4
        "epa_echo_compliance": "IC4",            # named in §4.6 as IC4
        "removed_page": "IC3",                   # judgment call, see above
        "executive_public_statement": "IC2",     # third-party coverage, reporter selects
        "first_party_announcement": "IC1",       # company announcements, named IC1
        "employee_review_aggregate": "IC4",      # employee reviews, named IC4
    }

    def backfill_instrument_class(self) -> None:
        ws = self.wb["Signal_Types"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if "instrument_class" not in headers:
            self.skip("Signal_Types.instrument_class not present yet")
            return
        i_name = headers.index("signal_type_name") + 1
        i_ic = headers.index("instrument_class") + 1
        filled, unknown = [], []
        for r in range(2, ws.max_row + 1):
            name = str(ws.cell(r, i_name).value or "").strip()
            if not name or ws.cell(r, i_ic).value not in (None, ""):
                continue
            ic = self.INSTRUMENT_CLASSES.get(name)
            if ic is None:
                unknown.append(name)
                continue
            ws.cell(r, i_ic).value = ic
            filled.append(f"{name}={ic}")
        if filled:
            self.note(f"Signal_Types.instrument_class backfilled: {', '.join(filled)}")
        else:
            self.skip("Signal_Types.instrument_class already populated")
        if unknown:
            # Never guessed. A signal type with no assigned class is reported, because a
            # silently blank class reads as "unbiased" everywhere downstream.
            self.note(f"  [!] no instrument_class assigned for: {', '.join(unknown)} "
                      f"-- assign in register_sources.py or INSTRUMENT_CLASSES")

    # ---------- validations ----------
    # ---------- theme-key retirement (core/topics.py::RETIRED_THEME_KEYS) ----------
    def rename_topic_keys(self) -> None:
        """Rewrite `Observations.topic` for every theme key retired in `core/topics.py`.

        Idempotent: a row already carrying the new key is not a match, so a second run
        finds nothing. The map is declared in ONE place (RETIRED_THEME_KEYS) and read
        here, by validate_repo_db.py check 8 and by theme_regression.py, so the three
        cannot disagree about what was renamed.

        A human-reviewed row carrying a retired key ABORTS the migration rather than
        being rewritten or skipped. Convention 35: provenance decides what is
        untouchable. Skipping it silently would leave a row the validator then fails on;
        rewriting it would move reviewed evidence under a person's verdict. Either is a
        decision for Matthew, so the migration stops and names the row.

        First use, 2026-09-02: `cybersecurity_ot` -> `cybersecurity` on O00225, O00230,
        O00234, O00255 (H-SELLERCONTENT-01 v1.2, P004/P005/P006/P012, all machine and
        unreviewed). theme_id THEME-08 unchanged; observation_text unchanged -- it
        records what the harness wrote at retrieval time, and the source did not change.
        """
        ws = self.wb["Observations"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        i_topic = headers.index("topic") + 1
        renamed: Counter = Counter()
        ids: list[str] = []
        refused: list[str] = []
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                continue
            topic = str(ws.cell(r, i_topic).value or "").strip()
            if topic not in RETIRED_THEME_KEYS:
                continue
            values = {h: ws.cell(r, i + 1).value for i, h in enumerate(headers) if h}
            oid = str(values.get("observation_id"))
            if is_human_authored(values):
                refused.append(f"{oid} ({topic})")
                continue
            ws.cell(r, i_topic).value = RETIRED_THEME_KEYS[topic]
            renamed[(topic, RETIRED_THEME_KEYS[topic])] += 1
            ids.append(oid)
        if refused:
            raise SystemExit(
                "ABORT: human-reviewed Observations carry a retired theme key and the "
                f"migration will not rewrite reviewed evidence: {', '.join(refused)}. "
                "Decide the row's disposition first (convention 35).")
        if not renamed:
            self.skip("no Observations row carries a retired theme key")
            return
        for (old, new), n in sorted(renamed.items()):
            self.note(f"Observations.topic: {n} row(s) {old} -> {new} "
                      f"[{', '.join(ids)}]")

    # ---------- observation_text carrying a retired theme LABEL ----------
    #
    # `observation_text` is machine prose generated from the theme label at retrieval time.
    # When a label changes, the text on rows written under the old label no longer matches
    # what a re-run would write, and `sync_observations` would then see a changed
    # fingerprint on an otherwise identical claim. Declared here as (harness, old phrase,
    # new phrase); only machine-authored rows are touched, and the step is idempotent
    # because the old phrase is gone after one pass.
    #
    # 2026-09-02 (session 7 follow-up, item 1): THEME-08's label narrowed from
    # "cybersecurity and OT/IT convergence" to "cybersecurity" in the split, and the
    # theme flipped buyer_detectable on 2026-08-31. H-SELLERCONTENT-01's template appends
    # a "NOTE: no buyer-side harness can currently detect this theme ..." sentence only
    # while buyer_detectable is False, so a faithful regeneration drops that sentence too.
    TEXT_RELABELS = [
        ("H-SELLERCONTENT-01",
         "markets capability in cybersecurity and OT/IT convergence (",
         "markets capability in cybersecurity ("),
    ]
    SELLERCONTENT_NOTE = (" NOTE: no buyer-side harness can currently detect this theme, so an "
                          "absence of matching buyer evidence is an instrumentation gap rather "
                          "than a demonstrated divergence.")

    def relabel_observation_text(self) -> None:
        """Regenerate `observation_text` on machine rows written under a retired label."""
        from core.topics import THEMES_BY_KEY
        ws = self.wb["Observations"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        i_text = headers.index("observation_text") + 1
        i_topic = headers.index("topic") + 1
        i_harness = headers.index("harness_id") + 1
        done: list[str] = []
        refused: list[str] = []
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                continue
            text = str(ws.cell(r, i_text).value or "")
            harness = str(ws.cell(r, i_harness).value or "")
            for h, old, new in self.TEXT_RELABELS:
                if harness != h or old not in text:
                    continue
                values = {hd: ws.cell(r, i + 1).value for i, hd in enumerate(headers) if hd}
                oid = str(values.get("observation_id"))
                if is_human_authored(values):
                    refused.append(oid)
                    continue
                text = text.replace(old, new)
                theme = THEMES_BY_KEY.get(str(ws.cell(r, i_topic).value or ""))
                if h == "H-SELLERCONTENT-01" and theme is not None and theme.buyer_detectable:
                    text = text.replace(self.SELLERCONTENT_NOTE, "")
                ws.cell(r, i_text).value = text
                done.append(oid)
        if refused:
            raise SystemExit("ABORT: human-reviewed rows carry a retired theme label in "
                             f"observation_text and will not be rewritten: {', '.join(refused)}")
        if done:
            self.note(f"Observations.observation_text regenerated under the current label on "
                      f"{len(done)} machine row(s) [{', '.join(done)}]")
        else:
            self.skip("no machine Observations row carries a retired theme label in its text")

    def apply_validations(self) -> None:
        created, rebound = [], []
        by_sheet: dict[str, list] = {}
        for sheet, col, src, ceiling in VALIDATIONS:
            by_sheet.setdefault(sheet, []).append((col, src, ceiling))

        for sheet, specs in by_sheet.items():
            ws = self.wb[sheet]
            for col, src, ceiling in specs:
                want_sqref = f"{col}2:{col}{ceiling}"
                want_f1 = f"Lookups!${src}$2:${src}${LOOKUPS_CEILING}"
                match = None
                for dv in ws.data_validations.dataValidation:
                    cols = {str(rng).split(":")[0].lstrip("$").rstrip("0123456789")
                            for rng in str(dv.sqref).split()}
                    if cols == {col}:
                        match = dv
                        break
                if match is None:
                    dv = DataValidation(type="list", formula1=want_f1, allow_blank=True,
                                        showErrorMessage=True)
                    dv.error = f"Value must come from Lookups column {src}."
                    dv.errorTitle = "Not in controlled vocabulary"
                    ws.add_data_validation(dv)
                    dv.sqref = want_sqref
                    created.append(f"{sheet}!{col}")
                elif str(match.sqref) != want_sqref or match.formula1 != want_f1:
                    match.sqref = openpyxl.worksheet.cell_range.MultiCellRange(want_sqref)
                    match.formula1 = want_f1
                    rebound.append(f"{sheet}!{col}")
        if created:
            self.note(f"validations CREATED (were missing): {', '.join(created)}")
        if rebound:
            self.note(f"validations rebound: {', '.join(rebound)}")
        if not created and not rebound:
            self.skip("all validations already correct")

    def run(self) -> None:
        self.add_source_id()
        self.add_vocabularies()
        self.extend_vocabularies()
        self.add_sheets()
        self.seed_signal_types()
        self.add_run_rollups()
        self.add_audit_gate_columns()
        self.add_reach_columns()
        self.populate_reach()
        self.add_company_columns()
        self.add_registry_lineage_column()
        self.register_observation_ids()
        self.backfill_audit_gate()
        self.backfill_instrument_class()
        self.rename_topic_keys()
        self.relabel_observation_text()
        self.apply_validations()

    def save(self) -> tuple[Path, list[Path]]:
        backup, pruned = backup_workbook(self.path)
        self.wb.save(self.path)
        return backup, pruned


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=str(DEFAULT_DB))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    path = Path(args.db)
    if not path.exists():
        raise SystemExit(f"ABORT: no workbook at {path}")

    m = Migration(path)
    m.run()

    print(f"workbook: {path}")
    print()
    if m.skipped:
        print("already in place:")
        for s in m.skipped:
            print(f"  ... {s}")
        print()
    if not m.actions:
        print("nothing to do -- schema is current")
        return 0
    print("changes:")
    for a in m.actions:
        print(f"  +   {a}")

    if not args.apply:
        print(f"\n{len(m.actions)} change(s) pending -- re-run with --apply")
        return 0

    backup, pruned = m.save()
    print(f"\napplied {len(m.actions)} change(s); backup at {backup.name}")
    if pruned:
        print(f"pruned {len(pruned)} older backup(s): {', '.join(p.name for p in pruned)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
