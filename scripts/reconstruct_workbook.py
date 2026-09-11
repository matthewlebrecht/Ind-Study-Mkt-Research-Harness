#!/usr/bin/env python3
"""
Rebuild `data/market_intel_db.xlsx` from the last known-good base plus the Anvil-100 import.

WHY THIS EXISTS
---------------
On 2026-08-26 the Anvil-100 CRM export was imported into a *pre-cleanup* copy of the
workbook. `repo_structure_spec.md` P2 names the result: "pilot IDs shifted by one,
HR-0001 naming the wrong harness". That file is preserved at
`data/archive/market_intel_db.STALE-2026-08-26.xlsx`.

The damage is asymmetric, which is what makes an automated fix safe:

  * The stale file has the ONLY copy of the Anvil-100 companies (A001-A100). Those live
    in their own ID namespace, so the off-by-one never touched them.
  * The stale file LOST the 27 reviewed FMCSA observations, the correct C0001-C0008
    pilot IDs, and the HR-0001..HR-0004 run history — all of which survive intact in
    `market_intel_db.backup-pre-jobpost-231246.xlsx`.

So: take the pilot side from the good base, take the Anvil side from the stale file, and
nothing has to be reconciled by hand.

WHY `backup-pre-jobpost-231246.xlsx` IS THE CORRECT BASE
--------------------------------------------------------
Verified against CLAUDE.md's authoritative account, not against file mtime (the handoff
warns explicitly that recency is a weak tiebreaker here — the stale copy is NEWER):

  * C0001 Midmark / C0002 Mack / C0003 Duke / C0004 Western Express / C0005 Venture /
    C0006 Kenco / C0007 PLS (pending_review) / C0008 CT Logistics (excluded)
  * 27 observations, review split 18 accepted / 9 corrected / 0 rejected
  * HR-0001..HR-0004 = H-FMCSA-01 v1.0, v1.1, v1.2, v1.3
  * Findings header-only

Independent corroboration from outside the workbook: the H-JOBPOST-01 response cache
names its files by company_id — `workday_C0006_*` (Kenco), `icims_C0007_*` (PLS),
`paylocity_C0005_*` (Venture). That is the base's numbering, not the stale file's. The
harness code and its cached evidence agree with this base, so adopting it keeps code,
cache, and database consistent.

WHAT THIS SCRIPT DOES NOT DO
----------------------------
It does not restore H-JOBPOST-01's 8 observations or its HR-0005 row. The base predates
that run by design. Those are restored by replaying the harness against its cache, which
is a harness operation, not a spreadsheet operation:

    python -m harnesses.h_jobpost_01.harness --commit --offline

Run this script first, then that.

Usage:
  python scripts/reconstruct_workbook.py                 # report only
  python scripts/reconstruct_workbook.py --apply
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.workbook_backup import backup_workbook  # noqa: E402

BASE = ROOT / "data/archive/market_intel_db.backup-pre-jobpost-231246.xlsx"
STALE = ROOT / "data/archive/market_intel_db.STALE-2026-08-26.xlsx"
TARGET = ROOT / "data/market_intel_db.xlsx"

# Columns the Anvil import added to Companies, in the order the stale file has them.
ANVIL_COLUMNS = ["website", "anvil_rank", "anvil_selection_score", "anvil_score_notes"]

# CLAUDE.md's authoritative pilot list. The base must match this exactly or we abort:
# this is the assertion that distinguishes the good base from the stale copy.
EXPECTED_PILOTS = [
    ("C0001", "Midmark Corporation", "qualified"),
    ("C0002", "Mack Group", "qualified"),
    ("C0003", "Duke Manufacturing Co.", "qualified"),
    ("C0004", "Western Express", "qualified"),
    ("C0005", "Venture Logistics", "qualified"),
    ("C0006", "Kenco Group", "qualified"),
    ("C0007", "PLS Logistics", "pending_review"),
    ("C0008", "CT Logistics", "excluded"),
]

EXPECTED_REVIEW_SPLIT = {"accepted": 18, "corrected": 9}
EXPECTED_OBSERVATIONS = 27
EXPECTED_RUNS = ["HR-0001", "HR-0002", "HR-0003", "HR-0004"]


class AbortBuild(SystemExit):
    """Raised loudly rather than proceeding against a workbook we cannot vouch for."""


def _rows(ws, col=1):
    """Row indices that have a value in `col`, skipping the header."""
    return [r for r in range(2, ws.max_row + 1) if ws.cell(r, col).value not in (None, "")]


def verify_base(wb) -> list[str]:
    """Assert the base really is the post-cleanup, post-v1.3, pre-jobpost state."""
    problems = []

    comp = wb["Companies"]
    actual = [
        (comp.cell(r, 1).value, comp.cell(r, 2).value, comp.cell(r, 11).value)
        for r in _rows(comp)
    ]
    if actual != EXPECTED_PILOTS:
        problems.append(
            "Companies does not match CLAUDE.md's authoritative pilot list.\n"
            f"    expected: {EXPECTED_PILOTS}\n    actual:   {actual}"
        )

    obs = wb["Observations"]
    obs_rows = _rows(obs)
    if len(obs_rows) != EXPECTED_OBSERVATIONS:
        problems.append(f"expected {EXPECTED_OBSERVATIONS} observations, found {len(obs_rows)}")

    headers = [obs.cell(1, c).value for c in range(1, obs.max_column + 1)]
    rs_col = headers.index("review_status") + 1
    split: dict[str, int] = {}
    for r in obs_rows:
        v = str(obs.cell(r, rs_col).value or "").strip()
        split[v] = split.get(v, 0) + 1
    if split != EXPECTED_REVIEW_SPLIT:
        problems.append(
            f"review_status split is {split}, expected {EXPECTED_REVIEW_SPLIT} "
            "(the 2026-08-24 manual review must be present and intact)"
        )

    runs = wb["Harness_Runs"]
    run_ids = [runs.cell(r, 1).value for r in _rows(runs)]
    if run_ids != EXPECTED_RUNS:
        problems.append(f"Harness_Runs is {run_ids}, expected {EXPECTED_RUNS}")

    if _rows(wb["Findings"]):
        problems.append("Findings has data rows; the good base is header-only")

    return problems


def verify_stale(wb) -> tuple[list[str], list[int]]:
    """Locate the Anvil rows in the stale file and confirm they are what we think."""
    problems = []
    comp = wb["Companies"]
    headers = [comp.cell(1, c).value for c in range(1, comp.max_column + 1)]

    for col in ANVIL_COLUMNS:
        if col not in headers:
            problems.append(f"stale file is missing expected Anvil column {col!r}")

    anvil_rows = [r for r in _rows(comp) if str(comp.cell(r, 1).value).startswith("A")]
    if len(anvil_rows) != 100:
        problems.append(f"expected 100 Anvil companies in the stale file, found {len(anvil_rows)}")

    return problems, anvil_rows


def reconstruct(apply: bool) -> int:
    for path in (BASE, STALE):
        if not path.exists():
            raise AbortBuild(f"ABORT: required input missing: {path}")

    base_wb = openpyxl.load_workbook(BASE)
    stale_wb = openpyxl.load_workbook(STALE)

    print(f"base  : {BASE.name}")
    print(f"anvil : {STALE.name}")
    print()

    problems = verify_base(base_wb)
    stale_problems, anvil_rows = verify_stale(stale_wb)
    problems += stale_problems

    if problems:
        print("ABORT — inputs failed verification. Nothing was written.\n")
        for p in problems:
            print(f"  !! {p}")
        raise AbortBuild(1)

    print("verification passed:")
    print(f"  base   8 pilots C0001-C0008 match CLAUDE.md, {EXPECTED_OBSERVATIONS} observations "
          f"({EXPECTED_REVIEW_SPLIT['accepted']} accepted / "
          f"{EXPECTED_REVIEW_SPLIT['corrected']} corrected), runs {EXPECTED_RUNS}")
    print(f"  stale  {len(anvil_rows)} Anvil companies A001-A100 with columns {ANVIL_COLUMNS}")
    print()

    # ---- graft the Anvil columns onto the base's Companies sheet ----
    comp = base_wb["Companies"]
    base_headers = [comp.cell(1, c).value for c in range(1, comp.max_column + 1)]
    stale_comp = stale_wb["Companies"]
    stale_headers = [stale_comp.cell(1, c).value for c in range(1, stale_comp.max_column + 1)]

    added_cols = []
    for col in ANVIL_COLUMNS:
        if col not in base_headers:
            base_headers.append(col)
            comp.cell(1, len(base_headers)).value = col
            added_cols.append(col)

    # ---- copy the Anvil rows across, mapping by header name, not position ----
    # Position-mapping is exactly how the original import went wrong. Every value is
    # placed by the name of its column in the source sheet.
    first_new_row = comp.max_row + 1
    copied = 0
    for r in anvil_rows:
        target_row = comp.max_row + 1
        for src_idx, name in enumerate(stale_headers, start=1):
            if not name or name not in base_headers:
                continue
            dst_idx = base_headers.index(name) + 1
            comp.cell(target_row, dst_idx).value = stale_comp.cell(r, src_idx).value
        copied += 1

    print(f"  + {len(added_cols)} column(s) appended to Companies: {added_cols}")
    print(f"  + {copied} Anvil rows copied into rows {first_new_row}-{comp.max_row} "
          f"(mapped by column name)")

    # ---- post-conditions ----
    ids = [comp.cell(r, 1).value for r in _rows(comp)]
    if len(ids) != 108:
        raise AbortBuild(f"ABORT: expected 108 companies after merge, got {len(ids)}")
    if len(set(ids)) != len(ids):
        raise AbortBuild("ABORT: duplicate company_id after merge")
    if ids[:8] != [p[0] for p in EXPECTED_PILOTS]:
        raise AbortBuild("ABORT: pilot IDs moved during merge")
    print(f"  = {len(ids)} companies, all IDs unique, pilots still C0001-C0008")

    if not apply:
        print(f"\nreport only — re-run with --apply to write {TARGET.relative_to(ROOT)}")
        return 0

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    if TARGET.exists():
        backup, pruned = backup_workbook(TARGET)
        print(f"\nexisting target backed up to {backup.name}")
        if pruned:
            print(f"pruned {len(pruned)} older backup(s): {', '.join(p.name for p in pruned)}")
    base_wb.save(TARGET)
    print(f"\nwrote {TARGET.relative_to(ROOT)}")
    print("NEXT: restore H-JOBPOST-01's 8 observations by replaying it against its cache:")
    print("      python -m harnesses.h_jobpost_01.harness --commit --offline")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="write the workbook (default: report only)")
    args = ap.parse_args()
    return reconstruct(args.apply)


if __name__ == "__main__":
    sys.exit(main())
