#!/usr/bin/env python3
"""
Backfill the Observation_Ids registry with ids that once existed and were deleted.

    python scripts/backfill_observation_ids.py            # report
    python scripts/backfill_observation_ids.py --apply    # write

The registry (convention 43) must know every id EVER assigned, or a fresh allocation could
hand a retired id to a different claim. `scripts/migrate_schema.py` registers the ids that
are live today; this script walks the workbook's git history for the ids that are not --
every delete-and-rewrite before 2026-09-06 renumbered them -- and registers each as
`retired` with the natural key it carried when it was last seen.

Where one natural key carried several ids over time (the same claim renumbered on each
replay), every id is registered against that key. A future re-proposal of the key reuses
the most recently assigned of them (`core/db.py::_assign_observation_id`); the older ones
stay retired forever.
"""

from __future__ import annotations

import argparse
import io
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.db import (MarketIntelDB, OBSERVATION_COLUMNS, OBSERVATION_ID_COLUMNS,  # noqa: E402
                     natural_key_of, today)
from core.workbook_backup import backup_workbook  # noqa: E402

WB = "data/market_intel_db.xlsx"


def history() -> list[tuple[str, str]]:
    out = subprocess.run(["git", "log", "--format=%h %ad", "--date=short", "--", WB],
                         capture_output=True, text=True, cwd=ROOT, check=True)
    return [tuple(line.split()) for line in out.stdout.splitlines() if line.strip()]


def rows_at(commit: str) -> dict[str, dict]:
    blob = subprocess.run(["git", "show", f"{commit}:{WB}"], capture_output=True,
                          cwd=ROOT, check=True).stdout
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if "Observations" not in wb.sheetnames:
        return {}
    it = wb["Observations"].iter_rows(values_only=True)
    headers = list(next(it))
    out = {}
    for r in it:
        if r and r[0]:
            d = dict(zip(headers, r))
            out[str(d["observation_id"])] = d
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    db = MarketIntelDB()
    if "Observation_Ids" not in db.wb.sheetnames:
        print("ABORT: Observation_Ids missing; run scripts/migrate_schema.py --apply first")
        return 1
    reg = db._registry()
    known = set(reg["by_id"])

    # newest commit first; the first time we see a retired id is its LAST known key
    seen: dict[str, tuple[str, dict, str, str]] = {}
    commits = history()
    print(f"walking {len(commits)} commit(s) of {WB}")
    for sha, date in commits:
        for oid, d in rows_at(sha).items():
            if oid in known or oid in seen:
                continue
            seen[oid] = (natural_key_of(d), d, sha, date)

    if not seen:
        print("nothing to backfill: every id in history is already registered")
        return 0
    print(f"{len(seen)} retired id(s) found in history and absent from the registry:")
    for oid in sorted(seen):
        key, d, sha, date = seen[oid]
        print(f"  {oid}  last seen {date} ({sha})  {d.get('harness_id')} {d.get('harness_version')}  "
              f"{str(d.get('company_id'))}/{str(d.get('topic'))[:30]}")
    if not args.apply:
        print("\nDRY RUN -- nothing written. Re-run with --apply.")
        return 0

    b, pruned = backup_workbook(WB)
    print(f"backup {b.name}" + (f"; pruned {[p.name for p in pruned]}" if pruned else ""))
    ws = reg["ws"]
    for oid in sorted(seen):
        key, d, sha, date = seen[oid]
        ws.append([oid, key, str(d.get("company_id") or ""), str(d.get("harness_id") or ""),
                   str(d.get("topic") or ""), str(d.get("source_url") or ""),
                   str(d.get("retrieval_date") or date)[:10], "retired", today(),
                   f"backfilled from git history; last seen at {sha} ({date}); "
                   f"renumbered or removed before the registry existed"[:200]])
    db.save()
    print(f"registered {len(seen)} retired id(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
