#!/usr/bin/env python3
"""
Load `Coherence_Framework_Taxonomy` from the framework workbook.

    python scripts/load_coherence_taxonomy.py            # report only
    python scripts/load_coherence_taxonomy.py --apply

Build handoff 2026-09-08 §1: "insert all 46 rows from the attached workbook as-is at
framework_version = v0.1, effective_date = 2026-09-08. No transformation needed -- copy
verbatim." This script does exactly that and nothing else. It reads the source's own
canonical `Coherence_Framework_Taxonomy` sheet, whose eight columns are precisely the
target's, and copies the cells across unchanged.

WHAT IT CHECKS BEFORE WRITING, and refuses on
---------------------------------------------
Verbatim is a promise about transformation, not a licence to load anything. The source is
reference data that every future tag keys to, so the load refuses unless:

  * the source's header matches the target's eight columns exactly;
  * there are 46 rows, 22 `dimension` / 19 `failure_family` / 5 `generative_force`;
  * every id is unique, and unique against what is already in the sheet;
  * every row carries a label and a definition;
  * `parent_group` is one of I-V on failure families and blank on everything else;
  * `framework_version` is uniform;
  * no id collides with the live workbook's id namespace (handoff §5).

Re-running is a no-op: rows already present at the same `framework_version` are skipped, so
this can be run again after a partial load without duplicating. A row whose id exists at the
same version but whose CONTENT differs is refused rather than overwritten -- the framework
revises by `superseded_by` pointing forward, never by rewriting a live row (§1).

THREE COLUMNS THE TARGET SCHEMA CANNOT HOLD
-------------------------------------------
The source workbook also carries per-entity detail sheets with content the agreed eight
columns have nowhere to put: `key_question` (22 dimensions), and `dimensions_implicated`
and `economic_consequence` (19 failure families). They are NOT loaded, because the handoff
fixes the target's columns. The source workbook is committed to `docs/schema/` so none of it
is lost, and `dimensions_implicated` is flagged in the spec: it is the family-to-dimension
mapping that §2's synthesis-time join will need, and it currently lives only in that file.
"""

from __future__ import annotations

import argparse
import collections
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.db import MarketIntelDB  # noqa: E402
from core.workbook_backup import backup_workbook  # noqa: E402

SHEET = "Coherence_Framework_Taxonomy"
COLUMNS = ["id", "framework_version", "label", "definition", "parent_group",
           "superseded_by", "effective_date", "entity_type"]
EXPECTED = {"dimension": 22, "failure_family": 19, "generative_force": 5}
PARENT_GROUPS = {"I", "II", "III", "IV", "V"}
DEFAULT_SOURCE = ROOT / "docs" / "schema" / "Coherence_Framework_Taxonomy_v0.1.xlsx"


def read_source(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"ABORT: {path.name} has no {SHEET!r} sheet (has {wb.sheetnames})")
    it = wb[SHEET].iter_rows(values_only=True)
    header = [h for h in next(it)]
    if header != COLUMNS:
        raise SystemExit(f"ABORT: source header does not match the agreed schema.\n"
                         f"  expected {COLUMNS}\n  actual   {header}")
    return [dict(zip(header, r)) for r in it if r and r[0]]


def verify(rows: list[dict]) -> list[str]:
    problems = []
    kinds = collections.Counter(str(r["entity_type"]) for r in rows)
    if len(rows) != sum(EXPECTED.values()):
        problems.append(f"expected {sum(EXPECTED.values())} rows, found {len(rows)}")
    for kind, n in EXPECTED.items():
        if kinds.get(kind, 0) != n:
            problems.append(f"expected {n} {kind} row(s), found {kinds.get(kind, 0)}")
    unknown = set(kinds) - set(EXPECTED)
    if unknown:
        problems.append(f"unknown entity_type value(s): {sorted(unknown)}")

    ids = [str(r["id"]).strip() for r in rows]
    dupes = [i for i, n in collections.Counter(ids).items() if n > 1]
    if dupes:
        problems.append(f"duplicate id(s) in source: {sorted(dupes)}")
    malformed = [i for i in ids if not re.fullmatch(r"COH-[DFG]\d{2}", i)]
    if malformed:
        problems.append(f"malformed id(s): {sorted(malformed)}")

    versions = {str(r["framework_version"]) for r in rows}
    if len(versions) != 1:
        problems.append(f"framework_version is not uniform: {sorted(versions)}")

    for r in rows:
        rid = str(r["id"]).strip()
        if not str(r.get("label") or "").strip():
            problems.append(f"{rid}: no label")
        if not str(r.get("definition") or "").strip():
            problems.append(f"{rid}: no definition")
        pg = str(r.get("parent_group") or "").strip()
        if str(r["entity_type"]) == "failure_family":
            if pg not in PARENT_GROUPS:
                problems.append(f"{rid}: failure_family parent_group {pg!r} not in I-V")
        elif pg:
            problems.append(f"{rid}: {r['entity_type']} must have a blank parent_group, "
                            f"found {pg!r}")
    return problems


def namespace_collisions(db: MarketIntelDB, ids: set[str]) -> list[str]:
    """Handoff §5: opaque ids should be globally unique. Scan every other sheet."""
    hits = []
    for name in db.wb.sheetnames:
        if name == SHEET:
            continue
        ws = db.wb[name]
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip() in ids:
                    hits.append(f"{name}: {v.strip()}")
    return sorted(set(hits))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(DEFAULT_SOURCE))
    ap.add_argument("--db", default=str(ROOT / "data" / "market_intel_db.xlsx"))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    src = Path(args.source)
    if not src.exists():
        raise SystemExit(f"ABORT: no source workbook at {src}")
    rows = read_source(src)
    problems = verify(rows)
    kinds = collections.Counter(str(r["entity_type"]) for r in rows)
    print(f"source {src.name}: {len(rows)} row(s) {dict(kinds)}")
    if problems:
        print("\nREFUSING TO LOAD:")
        for p in problems:
            print(f"  !!  {p}")
        return 1

    db = MarketIntelDB(Path(args.db))
    ws = db.wb[SHEET]
    header = [ws.cell(1, c).value for c in range(1, len(COLUMNS) + 1)]
    if header != COLUMNS:
        raise SystemExit(f"ABORT: target header drift.\n  expected {COLUMNS}\n  actual   {header}")

    existing = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            existing[str(ws.cell(r, 1).value).strip()] = {
                c: ws.cell(r, i + 1).value for i, c in enumerate(COLUMNS)}

    collisions = namespace_collisions(db, {str(r["id"]).strip() for r in rows})
    if collisions:
        print("\nREFUSING TO LOAD -- id collision with the live namespace (handoff §5, "
              "escalate rather than resolve):")
        for c in collisions[:10]:
            print(f"  !!  {c}")
        return 1
    print("id namespace: no collision with any other sheet")

    fresh, skipped, conflicts = [], [], []
    for r in rows:
        rid = str(r["id"]).strip()
        prior = existing.get(rid)
        if prior is None:
            fresh.append(r)
        elif all(str(prior.get(c) or "") == str(r.get(c) or "") for c in COLUMNS):
            skipped.append(rid)
        else:
            differing = [c for c in COLUMNS if str(prior.get(c) or "") != str(r.get(c) or "")]
            conflicts.append(f"{rid} differs in {differing}")

    if conflicts:
        print("\nREFUSING TO LOAD -- id(s) already present with different content. The "
              "framework revises by superseded_by pointing forward, never by rewriting a "
              "live row (§1):")
        for c in conflicts[:10]:
            print(f"  !!  {c}")
        return 1

    print(f"  {len(fresh)} to load, {len(skipped)} already present and identical")
    for r in fresh[:3]:
        print(f"    {r['id']}  {r['entity_type']:<17} {str(r['label'])[:52]}")
    if len(fresh) > 3:
        print(f"    ... and {len(fresh) - 3} more")
    if not fresh:
        print("\n  nothing to do")
        return 0
    if not args.apply:
        print("\n  DRY RUN -- nothing written. Re-run with --apply.")
        return 0

    b, _ = backup_workbook(Path(args.db))
    for r in fresh:
        ws.append([r[c] for c in COLUMNS])
    db.save()
    print(f"\n  loaded {len(fresh)} row(s); backup {b.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
