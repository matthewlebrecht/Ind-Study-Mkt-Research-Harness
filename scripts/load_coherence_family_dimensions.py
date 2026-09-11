#!/usr/bin/env python3
"""
Load `Coherence_Family_Dimensions` -- the failure-family to dimension mapping.

    python scripts/load_coherence_family_dimensions.py            # report only
    python scripts/load_coherence_family_dimensions.py --apply

Build handoff addendum 2026-09-08, closing the open item in spec §13.1. The mapping is
many-to-many, so it is a normalized table rather than a delimited column on
`Coherence_Framework_Taxonomy`; §2's synthesis-time join joins on this.

TWO SOURCES, CROSS-CHECKED, BECAUSE THE ADDENDUM ASKED FOR A DIFF
-----------------------------------------------------------------
The addendum reproduces the mapping in prose "so the mapping doesn't require re-parsing
prose out of the archived file", and its escalation note says that a row-count mismatch
means "either this transcription or the archived source has an error worth finding, not a
validator to relax". So this loader does not trust either one alone:

  * `dimensions_implicated` is parsed from the archived source workbook, which is the
    authoritative artifact and is committed to the repo; and
  * `TRANSCRIPTION` below holds the addendum's table verbatim.

They are diffed family by family before anything is written, and a disagreement REFUSES.
Checked 2026-09-08: all 18 listed families agree exactly, set for set, and the three
qualifier annotations agree. There is no data discrepancy anywhere in this mapping.

THE COUNT IN THE ADDENDUM IS ARITHMETICALLY WRONG, AND THE DATA IS NOT
----------------------------------------------------------------------
The addendum states "18 families x their listed dimensions = 82 rows, plus COH-F19's
expansion below = 100 rows total". Both figures are miscalculations:

  * its own table lists 105 dimension references across F01-F18, not 82 -- and the archived
    source lists the same 105, independently; and
  * even taking 82 as given, 82 + 22 = 104, not 100.

The verified total is 105 + 22 = 127. `EXPECTED_ROWS` is 127 for that reason and not to make
a check pass: the escalation note forbids relaxing the validator to fit a suspect source, and
the diff establishes that the source is not suspect -- the prose arithmetic is. Every one of
the 127 rows traces to a dimension named in both the addendum's table and the archived file.

COH-F19
-------
The source's `dimensions_implicated` for `COH-F19` reads "All - this is the cross-cutting
case". Per the addendum this expands literally to all 22 dimensions, each its own row with
`role_note = "cross-cutting / special status"`, so that every family has a real queryable
dimension set and synthesis need not special-case one family id.
"""

from __future__ import annotations

import argparse
import collections
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.db import MarketIntelDB  # noqa: E402
from core.workbook_backup import backup_workbook  # noqa: E402

SHEET = "Coherence_Family_Dimensions"
COLUMNS = ["family_id", "dimension_id", "framework_version", "role_note"]
TAXONOMY = "Coherence_Framework_Taxonomy"
SOURCE = ROOT / "docs" / "schema" / "Coherence_Framework_Taxonomy_v0.1.xlsx"
SOURCE_SHEET = "Failure_Families_F01-F19"
FRAMEWORK_VERSION = "v0.1"
CROSS_CUTTING = "cross-cutting / special status"
EXPECTED_ROWS = 127          # 105 listed references + COH-F19 expanded to all 22

# The addendum's table, verbatim, held only to be diffed against the archived source.
TRANSCRIPTION = {
    "COH-F01": "D02, D04, D05, D06, D07, D21",
    "COH-F02": "D02, D03, D06, D14, D21",
    "COH-F03": "D04, D05, D07, D08, D10, D22",
    "COH-F04": "D06(incentive), D07, D10, D11, D16, D21",
    "COH-F05": "D01, D06, D07, D08, D09, D20",
    "COH-F06": "D01, D06, D07, D09, D16, D17",
    "COH-F07": "D11, D12, D14, D20(sequencing), D21",
    "COH-F08": "D02, D05, D10, D13, D14, D16",
    "COH-F09": "D08(decision), D10, D11, D12, D20, D22",
    "COH-F10": "D01, D02, D03, D14, D17, D20",
    "COH-F11": "D01, D02, D05, D06, D14",
    "COH-F12": "D01, D05, D09, D10, D13, D16",
    "COH-F13": "D01, D02, D09, D16, D17, D19",
    "COH-F14": "D06, D07, D10, D15, D17, D18, D22",
    "COH-F15": "D07, D08, D13, D14, D15, D22",
    "COH-F16": "D01, D02, D09, D19, D20",
    "COH-F17": "D11, D12, D13, D19, D20, D22",
    "COH-F18": "D06, D08, D10, D11, D14, D21",
    "COH-F19": "All -- cross-cutting",
}
_REF = re.compile(r"(D\d{2})\s*(?:\(([^)]*)\))?")


def parse(cell: str) -> list[tuple[str, str | None]]:
    """[(dimension_id, role_note or None)] from a `dimensions_implicated` cell."""
    return [(f"COH-{d}", (q.strip() or None) if q else None)
            for d, q in _REF.findall(str(cell or ""))]


def read_source(path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SOURCE_SHEET not in wb.sheetnames:
        raise SystemExit(f"ABORT: {path.name} has no {SOURCE_SHEET!r} sheet")
    it = wb[SOURCE_SHEET].iter_rows(values_only=True)
    head = list(next(it))
    if "dimensions_implicated" not in head:
        raise SystemExit(f"ABORT: {SOURCE_SHEET} has no dimensions_implicated column")
    return {str(r["id"]): str(r["dimensions_implicated"])
            for r in (dict(zip(head, x)) for x in it if x and x[0])}


def build(src: dict[str, str]) -> tuple[list[dict], list[str]]:
    """(rows, problems). Diffs the archived source against the addendum's transcription."""
    problems, rows = [], []
    only_src = sorted(set(src) - set(TRANSCRIPTION))
    only_add = sorted(set(TRANSCRIPTION) - set(src))
    if only_src or only_add:
        problems.append(f"family sets differ: source-only={only_src} addendum-only={only_add}")

    for fid in sorted(set(src) & set(TRANSCRIPTION)):
        s_raw, a_raw = src[fid], TRANSCRIPTION[fid]
        cross = "all" in s_raw.lower() and not _REF.findall(s_raw)
        if cross:
            if _REF.findall(a_raw):
                problems.append(f"{fid}: source says cross-cutting, addendum lists dimensions")
            rows.extend({"family_id": fid, "dimension_id": f"COH-D{n:02d}",
                         "framework_version": FRAMEWORK_VERSION, "role_note": CROSS_CUTTING}
                        for n in range(1, 23))
            continue
        s_pairs, a_pairs = parse(s_raw), parse(a_raw)
        s_dims = {d for d, _ in s_pairs}
        a_dims = {d for d, _ in a_pairs}
        if s_dims != a_dims:
            problems.append(f"{fid}: TRANSCRIPTION disagrees with the archived source -- "
                            f"addendum-only={sorted(a_dims - s_dims)} "
                            f"source-only={sorted(s_dims - a_dims)}")
            continue
        s_q = {d: q for d, q in s_pairs if q}
        a_q = {d: q for d, q in a_pairs if q}
        if s_q != a_q:
            problems.append(f"{fid}: qualifier annotations differ -- source={s_q} addendum={a_q}")
        for d, q in s_pairs:
            rows.append({"family_id": fid, "dimension_id": d,
                         "framework_version": FRAMEWORK_VERSION, "role_note": q})
    return rows, problems


def verify(rows: list[dict], db: MarketIntelDB) -> list[str]:
    problems = []
    ws = db.wb[TAXONOMY]
    head = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    tax = {}
    for r in range(2, ws.max_row + 1):
        d = {h: ws.cell(r, i + 1).value for i, h in enumerate(head) if h}
        if d.get("id"):
            tax[str(d["id"]).strip()] = (str(d.get("entity_type") or ""),
                                         str(d.get("framework_version") or ""))
    for row in rows:
        for col, want in (("family_id", "failure_family"), ("dimension_id", "dimension")):
            rid = row[col]
            if rid not in tax:
                problems.append(f"{row['family_id']}/{row['dimension_id']}: {rid} absent from "
                                f"{TAXONOMY}")
            elif tax[rid][0] != want:
                problems.append(f"{row['family_id']}/{row['dimension_id']}: {rid} is "
                                f"{tax[rid][0]!r}, must be {want!r}")
            elif tax[rid][1] != row["framework_version"]:
                problems.append(f"{rid}: framework_version {tax[rid][1]!r} does not match the "
                                f"mapping's {row['framework_version']!r}")
    dupes = [k for k, n in collections.Counter(
        (r["family_id"], r["dimension_id"], r["framework_version"]) for r in rows).items() if n > 1]
    if dupes:
        problems.append(f"duplicate (family, dimension, version) triple(s): {dupes[:5]}")
    if len(rows) != EXPECTED_ROWS:
        problems.append(f"expected {EXPECTED_ROWS} rows, built {len(rows)} -- STOP and diff "
                        f"against the addendum table; do not adjust EXPECTED_ROWS to fit")
    return problems


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(SOURCE))
    ap.add_argument("--db", default=str(ROOT / "data" / "market_intel_db.xlsx"))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    src = read_source(Path(args.source))
    rows, problems = build(src)
    db = MarketIntelDB(Path(args.db))
    problems += verify(rows, db)

    per_family = collections.Counter(r["family_id"] for r in rows)
    listed = sum(n for f, n in per_family.items() if f != "COH-F19")
    print(f"source {Path(args.source).name} + addendum transcription, diffed: "
          f"{len(per_family)} families")
    print(f"  {listed} listed dimension reference(s) across F01-F18, "
          f"{per_family.get('COH-F19', 0)} from COH-F19's cross-cutting expansion, "
          f"{len(rows)} rows")
    print(f"  role_note set on: "
          f"{sorted({(r['family_id'], r['dimension_id'], r['role_note']) for r in rows if r['role_note'] and r['role_note'] != CROSS_CUTTING})}")
    if problems:
        print("\nREFUSING TO LOAD:")
        for p in problems:
            print(f"  !!  {p}")
        return 1

    ws = db.wb[SHEET]
    header = [ws.cell(1, c).value for c in range(1, len(COLUMNS) + 1)]
    if header != COLUMNS:
        raise SystemExit(f"ABORT: target header drift.\n  expected {COLUMNS}\n  actual   {header}")
    existing = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            d = {c: ws.cell(r, i + 1).value for i, c in enumerate(COLUMNS)}
            existing[(str(d["family_id"]), str(d["dimension_id"]),
                      str(d["framework_version"]))] = d

    fresh, skipped, conflicts = [], 0, []
    for row in rows:
        key = (row["family_id"], row["dimension_id"], row["framework_version"])
        prior = existing.get(key)
        if prior is None:
            fresh.append(row)
        elif str(prior.get("role_note") or "") == str(row["role_note"] or ""):
            skipped += 1
        else:
            conflicts.append(f"{key}: role_note {prior.get('role_note')!r} -> {row['role_note']!r}")
    if conflicts:
        print("\nREFUSING TO LOAD -- triple(s) present with a different role_note. A mapping "
              "change is new rows at a new framework_version, never a rewrite:")
        for c in conflicts[:10]:
            print(f"  !!  {c}")
        return 1

    print(f"  {len(fresh)} to load, {skipped} already present and identical")
    if not fresh:
        print("\n  nothing to do")
        return 0
    if not args.apply:
        print("\n  DRY RUN -- nothing written. Re-run with --apply.")
        return 0
    b, _ = backup_workbook(Path(args.db))
    for row in fresh:
        ws.append([row[c] for c in COLUMNS])
    db.save()
    print(f"\n  loaded {len(fresh)} row(s); backup {b.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
