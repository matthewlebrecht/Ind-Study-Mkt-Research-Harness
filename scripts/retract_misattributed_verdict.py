#!/usr/bin/env python3
"""
Retract ONE audit verdict that was attributed to a reviewer who did not make it.

THE INCIDENT (2026-09-20). Matthew Lebrecht said "I review all manual review rows and they are all
supported". That was applied to the 30 rows of the H-FIRSTPARTY-01 v1.3 audit census -- correct, its
verdict vocabulary is exactly `supported` -- and ALSO to O00530, which is not an audit-census row. It
is the last live row of the Front Line Power Construction identity question, and the question it
carries is not "is this claim supported at the strength stated" but "is this release about Power
Construction (A073) or about Front Line Power Construction, a different company". Matthew then said
that call is his to make directly. It had not been made, so the row's `review_source = human` was a
claim about a judgment that did not happen.

WHY THIS IS NOT A BREACH OF "NEVER OVERWRITE HUMAN REVIEW". That convention protects a judgment a
person actually made. Here the stored judgment is the thing that is wrong: the row records a reviewer
verdict nobody gave. Leaving it would put a false attribution into the evidence base and into the
count of reviewed rows, which is worse than removing it. The retraction restores the row to exactly
its pre-verdict state and nothing else.

GUARDED BOTH WAYS, as a correction of this kind has to be:
  * it refuses unless the row currently holds exactly the mis-applied values;
  * it refuses unless the restored values equal the row as committed at `BEFORE_COMMIT`;
  * it verifies every OTHER Observations row is unchanged before saving.

    python scripts/retract_misattributed_verdict.py            # dry run on a temp copy
    python scripts/retract_misattributed_verdict.py --apply
"""

from __future__ import annotations

import argparse
import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.workbook_backup import backup_workbook  # noqa: E402

OBSERVATION_ID = "O00530"
BEFORE_COMMIT = "e826241"          # the commit before the verdict was applied
FIELDS = ("review_status", "review_source", "audit_verdict", "reviewer_notes")
EXPECT_NOW = {"review_status": "accepted", "review_source": "human", "audit_verdict": "supported"}


def load(path_or_buf):
    wb = openpyxl.load_workbook(path_or_buf, read_only=True, data_only=True)
    ws = wb["Observations"]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    return headers, [dict(zip(headers, r)) for r in it if r and r[0]]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    live_path = ROOT / "data" / "market_intel_db.xlsx"
    target = live_path if args.apply else Path(tempfile.mkdtemp()) / "dry.xlsx"
    if not args.apply:
        shutil.copy2(live_path, target)

    blob = subprocess.run(["git", "show", f"{BEFORE_COMMIT}:data/market_intel_db.xlsx"],
                          cwd=ROOT, capture_output=True)
    if blob.returncode != 0:
        print(f"cannot read the workbook at {BEFORE_COMMIT}")
        return 2
    _, before_rows = load(io.BytesIO(blob.stdout))
    before = next(r for r in before_rows if r["observation_id"] == OBSERVATION_ID)

    _, now_rows = load(live_path)
    now = next(r for r in now_rows if r["observation_id"] == OBSERVATION_ID)

    for k, v in EXPECT_NOW.items():
        if str(now.get(k) or "") != v:
            print(f"REFUSING: {OBSERVATION_ID}.{k} is {now.get(k)!r}, expected the mis-applied {v!r}. "
                  f"Nothing written.")
            return 2
    print(f"{OBSERVATION_ID}: restoring {len(FIELDS)} field(s) to their state at {BEFORE_COMMIT}")
    for f in FIELDS:
        print(f"   {f}: {str(now.get(f))[:60]!r} -> {str(before.get(f))[:60]!r}")

    wb = openpyxl.load_workbook(target)
    ws = wb["Observations"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col = {h: i + 1 for i, h in enumerate(headers) if h}
    row_no = next(r for r in range(2, ws.max_row + 1)
                  if str(ws.cell(r, 1).value or "") == OBSERVATION_ID)
    for f in FIELDS:
        ws.cell(row_no, col[f]).value = before.get(f)

    out = Path(tempfile.mkdtemp()) / "verify.xlsx"
    wb.save(out)
    _, after_rows = load(out)
    changed = [a["observation_id"] for a, n in zip(after_rows, now_rows)
               if a != n and a["observation_id"] != OBSERVATION_ID]
    if changed:
        print(f"REFUSING: {len(changed)} other row(s) would change: {changed[:5]}")
        return 2
    restored = next(r for r in after_rows if r["observation_id"] == OBSERVATION_ID)
    if any(str(restored.get(f) or "") != str(before.get(f) or "") for f in FIELDS):
        print("REFUSING: the restored row does not equal the committed one")
        return 2
    print("verified: the row equals its pre-verdict state and no other row changed")

    if args.apply:
        backup_workbook(live_path)
        wb.save(live_path)
        print("saved market_intel_db.xlsx")
    else:
        print("DRY RUN -- nothing written.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
