"""
H-LOCALRECORDS-01 -- fragmented public records, one harness with four sub-scopes (Week 4 item 4).

WHY ONE HARNESS
---------------
State WARN lists, city permit portals, council records and website technology traces are each
expected to be thin for this universe, and each is fragmented by jurisdiction. Four standalone
instruments would duplicate the same access and rate-limit handling four times over, so they
share `source.py` and are declared here as sub-scopes, each with its own signal type, family,
instrument class and jurisdictional scope. Measured before building (2026-09-13):

  sub-scope   signal (id)                                source(s)                          class
  ---------   -----------------------------------------  ---------------------------------  -----
  warn        state_warn_notice (ST-0013), family 3      Texas TWC yearly xlsx, Utah DWS    IC4
                                                         HTML table, California EDD current
                                                         fiscal-year xlsx
  permits     municipal_permit_as_contractor (ST-0014),  Chicago, Seattle, Austin Socrata   IC3
              family 8                                   building permits
  council     council_matter_title_mention (ST-0015),    Seattle Legistar web API           IC3
              family 16
  techstack   website_technology_fingerprint (ST-0016),  company homepages as archived by   IC3
              family 5                                   H-EXECID-01

WHAT EACH ONE CAN AND CANNOT SAY
--------------------------------
* WARN: a statutory notice (IC4) that an employer is cutting 50+ jobs at a site. The company is
  the employer, so it is about the company's own operations -- but it records contraction,
  closure or relocation, never modernization, and routes to no theme. Scoped absence for
  companies headquartered in a reached state; a notice is filed where the SITE is, so presence
  is searched for every buyer. California's list is only the current fiscal year in structured
  form (earlier years are PDFs this version does not parse), so a California absence covers
  2026-07-01 onward and the row says so. Out of scope, measured: Missouri (Incapsula challenge),
  Minnesota (perfdrive redirect), Washington (the linked data page is gone), Illinois (not
  located), Tennessee (connection reset, unattributed).
* permits: in all three cities the buyers appear as the CONTRACTOR on someone else's project
  (Chicago: Walsh 665 general-contractor roles to 12 "owner" roles, and the owner rows are GCs
  misfiled in the owner slot or individuals who share a surname). So a row records local
  project volume as a seller -- the procurement direction problem again -- and owner-role hits
  are counted and logged, never written. Scoped absence only for companies headquartered in
  the city (Chicago, Seattle).
* council: Seattle is the only one of the universe's HQ cities on the Legistar web API
  (Chicago and Salt Lake City are not). Matter TITLES only -- the signal is defined as that,
  so an empty result is complete for it and nothing more. Full-phrase search only: searching
  a first word matched "Lease" ordinances for Lease Crutcher Lewis.
* techstack: specific front-end markers (asset hosts, generator tags, platform paths) in the
  company's served homepage, read from H-EXECID-01's dated archive, no new fetch. Front end
  only: it says what the public site is built on, not the company's internal systems.

None is a theme instrument in core/composition.py; the silence of none licenses anything about
a theme.

IDENTITY
--------
H-PROCUREMENT-01's `name_matches` (exact tokens, initials collapsed, JVs and professional
practices refused, common-word names matched in full), after removing what these sources add
to a name without changing the entity: "***MAIN***" / "(MAIN)" markers, state codes and
directionals ("LEASE CRUTCHER LEWIS WA LLC", "SWINERTON BUILDERS NW"), and, on WARN site names,
facility words ("FedEx Corporation Facility (Ft. Worth)").

    python harnesses/h_localrecords_01/harness.py              # dry run
    python harnesses/h_localrecords_01/harness.py --commit
    python harnesses/h_localrecords_01/harness.py --offline
    python harnesses/h_localrecords_01/harness.py --scopes warn,techstack
"""

from __future__ import annotations

import argparse
import datetime as _dt
import glob
import io
import json
import os
import re
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urljoin

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from core.cache import slug  # noqa: E402
from core.db import MarketIntelDB, Observation, today  # noqa: E402
from core.resolution import US_STATES, state_code, strip_legal_suffix  # noqa: E402
from core.windows import STALENESS_YEARS  # noqa: E402
from harnesses.h_localrecords_01.source import AccessError, Source  # noqa: E402
from core import aliases  # noqa: E402
from harnesses.h_firstparty_01.article import COMMON_WORD_NAMES  # noqa: E402
from harnesses.h_procurement_01.harness import (  # noqa: E402
    LOCAL_COMMON_WORDS, SEARCH_GENERIC, distinctive, is_common_word_name, name_matches,
    norm_tokens_raw)

HARNESS_ID = "H-LOCALRECORDS-01"
HARNESS_NAME = "Fragmented Public Records Reader (WARN, permits, council, website tech)"
VERSION = "v1.2"
OUTPUT_DIR = ROOT / "harness_output" / HARNESS_ID

SIG_WARN, FAM_WARN = "state_warn_notice", "3_workforce_org_exhaust"
SIG_PERMIT, FAM_PERMIT = "municipal_permit_as_contractor", "8_physical_footprint_capacity"
SIG_COUNCIL, FAM_COUNCIL = "council_matter_title_mention", "16_local_community_records"
SIG_TECH, FAM_TECH = "website_technology_fingerprint", "5_technology_stack_traces"
SCOPES = ("warn", "permits", "council", "techstack")
SIGNAL_OF = {"warn": SIG_WARN, "permits": SIG_PERMIT, "council": SIG_COUNCIL, "techstack": SIG_TECH}
FAMILY_OF = {SIG_WARN: FAM_WARN, SIG_PERMIT: FAM_PERMIT, SIG_COUNCIL: FAM_COUNCIL, SIG_TECH: FAM_TECH}

# ------------------------------------------------------------------ identity helpers
DIRECTIONALS = {"N", "S", "E", "W", "NW", "NE", "SW", "SE", "MAIN"}
FACILITY_WORDS = {"FACILITY", "FACILITIES", "PLANT", "WAREHOUSE", "SITE", "LOCATION", "TERMINAL",
                  "OFFICE", "OFFICES", "BRANCH", "HUB", "DC", "CAMPUS", "YARD", "DEPOT", "STORE"}


def clean_candidate(name: str, facility: bool = False, drop: set | None = None) -> str:
    """Remove what a record adds to a name without changing the entity.

    Seattle and Austin contractor strings carry "***MAIN***" / "(MAIN)" markers, state codes and
    directionals ("LEASE CRUTCHER LEWIS WA LLC", "SWINERTON BUILDERS NW"); WARN site names carry
    facility words and the site's city. Dropping these is not loosening identity: name_matches
    still requires every distinctive company token, and a name that carries anything else is
    still refused."""
    s = re.sub(r"\*+\s*MAIN\s*\*+|\(\s*MAIN\s*\)", " ", str(name or ""), flags=re.I)
    s = re.sub(r"\(.*?\)", " ", s)
    words = re.split(r"[^A-Za-z0-9&'+.\-]+", s)
    out = []
    for w in words:
        u = re.sub(r"[^A-Z0-9]", "", w.upper())
        if not u:
            continue
        if u in DIRECTIONALS or (len(u) == 2 and u in US_STATES and out):
            continue
        # An entity numeral is a successor registration of the same firm, not another name:
        # WALSH CONSTRUCTION COMPANY II, LLC holds 97 of Walsh's Chicago permits, and without
        # this name_matches read "II" as an extra word and refused it (caught by
        # core/tests/test_localrecords.py before the second dry run was read).
        if u in {"II", "III", "IV"} and out:
            continue
        if facility and u in FACILITY_WORDS:
            continue
        if drop and u in drop:
            continue
        out.append(w)
    return " ".join(out)


# Legal forms dropped when comparing operating-name FORMS. GROUP, COMPANIES, HOLDINGS and the
# line-of-business words are kept on purpose: "THE WALSH GROUP" and a bare "WALSH" must not
# compare equal.
FORM_DROP = {"THE", "INC", "INCORPORATED", "LLC", "LLP", "LP", "LTD", "LIMITED", "CORP",
             "CORPORATION", "CO", "COMPANY", "PLC", "II", "III", "IV", "AND"}

# Verified legal names per company_id, filled in main() from released H-PROCUREMENT-01 rows
# (entities SAM registers under the company, accepted there with HQ-state corroboration) and
# from the human alias registry. Tests set it directly.
VERIFIED: dict[str, list[str]] = {}

# Names a human has confirmed are a DIFFERENT firm, per company_id, from data/company_aliases.json
# `not_variants` (v1.2, 2026-09-15). Loaded on first use; tests set it directly.
NOT_VARIANTS: dict[str, list[str]] | None = None


def not_variants_for(company_id: str) -> list[str]:
    global NOT_VARIANTS
    if NOT_VARIANTS is None:
        NOT_VARIANTS = {cid: e.get("not_variants", []) for cid, e in aliases.load().items()}
    return NOT_VARIANTS.get(company_id, [])


def form_tokens(name: str) -> tuple:
    return tuple(t for t in norm_tokens_raw(clean_candidate(name)) if t not in FORM_DROP)


def identity_ok(company: dict, candidate: str) -> tuple[bool, str]:
    """name_matches, then -- for a name with ONE distinctive token -- an exact operating-name form.

    The first full dry run (2026-09-13) accepted, for The Walsh Group in Chicago, a bare "WALSH"
    (a web-portal applicant), WALSH SERVICES INCORPORATED (a Crestwood plumber) and WALSH
    BROTHERS CONSTRUCTION; for Mortenson a Wilmette masonry firm called MORTENSON CONSTRUCTION;
    for The Yates Companies an Elgin, Texas "Yates Construction Inc."; for McCarthy Holdings a
    "MCCARTHY CONSTRUCTION CO". Each name reduces to one surname, and an HQ-state test cannot
    separate them (the Walsh plumber is in Illinois too). What does separate them, measured on
    every case, is the exact form of the name: the company's canonical name, a legal name SAM
    registers for it (WALSH CONSTRUCTION COMPANY, GILBANE BUILDING COMPANY, MCCARTHY BUILDING
    COMPANIES, SWINERTON BUILDERS), or a human alias. Anything else is refused and logged; the
    stated cost is a real operating name nobody has verified (Clayco Construction, McShane
    Construction Company), which is an alias to seed, not a guess to make."""
    # v1.2 (2026-09-15): a name a human has confirmed is a different firm is refused as such, first,
    # so the log records a settled non-match rather than a pending identity question (Clayco:
    # the Austin/Hutto "CLAYCO CONSTRUCTION", a small residential contractor, is not Clayco).
    confirmed_other = {form_tokens(v) for v in not_variants_for(company["company_id"])} - {()}
    if confirmed_other:
        oparts = [p for p in re.split(r"\s+(?:dba|d/b/a|doing business as)\s+", candidate, flags=re.I) if p.strip()]
        if {form_tokens(p) for p in oparts} & confirmed_other:
            return False, ("confirmed_different_entity: listed under not_variants in data/company_aliases.json "
                           "by a human decision")
    # v1.1 (2026-09-14): an exact form of a VERIFIED legal name or source-record alias is the
    # identity, and is checked before name_matches. Seattle files Mortenson's permits as "MA
    # Mortenson Company", which name_matches refuses for the extra initials token exactly as the
    # procurement matcher refuses W. W. CLYDE & CO.; the alias M. A. MORTENSON COMPANY (the
    # USASpending recipient name) is what separates it from the Wilmette MORTENSON CONSTRUCTION.
    verified = {form_tokens(v) for v in VERIFIED.get(company["company_id"], [])} - {()}
    if verified:
        vparts = [p for p in re.split(r"\s+(?:dba|d/b/a|doing business as)\s+", candidate, flags=re.I) if p.strip()]
        if {form_tokens(p) for p in vparts} & verified:
            return True, "exact form of a SAM-verified legal name or a source-record alias"
    ok, why = name_matches(company["canonical_name"], candidate)
    if not ok:
        return False, why
    if len(distinctive(re.sub(r"\(.*?\)", "", company["canonical_name"]))) >= 2:
        return True, why
    # A dba clause is split here as name_matches splits it: Austin files Layton's permits under
    # "Layton Construction Company, LLC dba Layton Builders of Texas, LLC", which compared as one
    # form never equals "Layton Construction" (second full dry run, 2026-09-13).
    parts = [p for p in re.split(r"\s+(?:dba|d/b/a|doing business as)\s+", candidate, flags=re.I) if p.strip()]
    cands = {form_tokens(p) for p in parts}
    cand = form_tokens(candidate)
    forms = {form_tokens(company["canonical_name"])}
    forms |= {form_tokens(v) for v in VERIFIED.get(company["company_id"], [])}
    if cands & forms:
        return True, "one-token name matched on an exact operating-name form (canonical, SAM-verified or alias)"
    return False, (f"single_token_unverified_form: {' '.join(cand)} is not the canonical name, a SAM-verified "
                   f"legal name or an alias for {company['canonical_name']}")


def load_verified(db) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    ws = db.wb["Observations"]
    headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
    ix = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[ix["harness_id"]] != "H-PROCUREMENT-01":
            continue
        if row[ix["publication_state"]] != "released":
            continue
        m = re.search(r"Recipient: (.*?)\. A federal", str(row[ix["observation_text"]] or ""))
        if not m:
            continue
        for part in m.group(1).split("; "):
            name = re.sub(r"\s*\(UEI.*$", "", part).strip()
            if name:
                out.setdefault(str(row[ix["company_id"]]), []).append(name)
    for cid, entry in aliases.load().items():
        out.setdefault(cid, []).extend(entry.get("name_variants", []))
    return {k: sorted(set(v)) for k, v in out.items()}


def search_token(company_name: str) -> str:
    """The text a LIKE / substring query searches: the longest distinctive, non-common token of
    five or more letters, else the name minus its legal suffix. Identity is decided afterwards
    on every string the query returns."""
    name = re.sub(r"\(.*?\)", "", company_name).strip()
    # Never an ordinary or generic word: the first dry run searched GENERAL for Joeris General
    # Contractors and MANAGEMENT for JRM Construction Management, which can push the company
    # itself past the grouped query's row limit.
    d = [t for t in distinctive(name) if t.isalpha() and len(t) >= 5
         and t.lower() not in COMMON_WORD_NAMES and t.lower() not in LOCAL_COMMON_WORDS
         and t not in SEARCH_GENERIC]
    if d and not is_common_word_name(name):
        return max(d, key=len)
    return strip_legal_suffix(name).upper()


def parse_date(v) -> _dt.date | None:
    """Dates as these sources actually write them: datetimes, ISO strings, M/D/YY and the
    Utah list's '01/07//09'."""
    if v is None or v == "":
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = str(v).strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"(\d{1,2})/+(\d{1,2})/+(\d{2,4})$", s)
    if m:
        y = int(m.group(3))
        y = y + 2000 if y < 100 else y
        try:
            return _dt.date(y, int(m.group(1)), int(m.group(2)))
        except ValueError:
            return None
    return None


def window(stamp: str) -> tuple[_dt.date, _dt.date]:
    end = _dt.date.fromisoformat(stamp)
    try:
        return end.replace(year=end.year - STALENESS_YEARS), end
    except ValueError:
        return end.replace(year=end.year - STALENESS_YEARS, day=28), end


def hq_city(company: dict) -> str:
    hq = str(company.get("hq_state") or "")
    return hq.split(",")[0].strip().upper() if "," in hq else ""


# ------------------------------------------------------------------ WARN
WARN_PAGES = {
    "TX": ("Texas Workforce Commission", "https://www.twc.texas.gov/data-reports/warn-notice"),
    "UT": ("Utah Department of Workforce Services", "https://jobs.utah.gov/employer/business/warnnotices.html"),
    "CA": ("California EDD", "https://edd.ca.gov/en/jobs_and_training/Layoff_Services_WARN"),
}


def _xlsx_rows(content: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]


def warn_notices(src: Source, state: str, start: _dt.date, log: dict) -> list[dict]:
    """Every notice in the state's list dated inside the window, as {state, date, name, count,
    city, url}. Files are discovered from the list page's own links, never constructed."""
    agency, page = WARN_PAGES[state]
    final, html = src.get(page, key=f"warn_page_{state}")
    out: list[dict] = []
    if state == "UT":
        for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", html, re.S):
            cells = [re.sub(r"<[^>]+>|\s+", " ", c).strip() for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, re.S)]
            if len(cells) >= 4 and cells[1] and cells[0].lower() != "date of notice":
                d = parse_date(cells[0])
                if d and d >= start:
                    out.append({"state": state, "date": d, "name": cells[1], "city": cells[2],
                                "count": cells[3], "url": final})
        log["UT"] = {"page": final, "notices_in_window": len(out)}
        return out
    links = sorted({urljoin(final, h) for h in re.findall(r'href="([^"]+\.xlsx)"', html, re.I)})
    if state == "TX":
        files = [l for l in links if re.search(r"warn-act-listings-(\d{4})", l)
                 and int(re.search(r"warn-act-listings-(\d{4})", l).group(1)) >= start.year]
    else:
        files = [l for l in links if re.search(r"warn_report", l, re.I)]
    log[state] = {"page": final, "files": files}
    if not files:
        raise AccessError(f"{agency}: no structured WARN file linked from {final}",
                          "source_drift_detected", "code_change")
    for f in files:
        _, content = src.get(f, key=f"warn_{state}_{os.path.basename(f)}", binary=True)
        for title, rows in _xlsx_rows(content):
            header_i = next((i for i, r in enumerate(rows[:6])
                             if any(str(v or "").strip().upper() in ("JOB_SITE_NAME", "COMPANY") for v in r)), None)
            if header_i is None:
                continue
            header = [re.sub(r"\s+", " ", str(v or "")).strip().upper() for v in rows[header_i]]
            ix = {h: i for i, h in enumerate(header)}
            name_i = ix.get("JOB_SITE_NAME", ix.get("COMPANY"))
            date_i = ix.get("NOTICE_DATE", ix.get("NOTICE DATE"))
            count_i = ix.get("TOTAL_LAYOFF_NUMBER", ix.get("NO. OF EMPLOYEES"))
            city_i = ix.get("CITY_NAME", ix.get("ADDRESS"))
            for r in rows[header_i + 1:]:
                d = parse_date(r[date_i]) if date_i is not None and date_i < len(r) else None
                if not d or d < start or name_i is None or not r[name_i]:
                    continue
                out.append({"state": state, "date": d, "name": str(r[name_i]).strip(),
                            "city": str(r[city_i] or "") if city_i is not None else "",
                            "count": r[count_i] if count_i is not None else "", "url": f,
                            "sheet": title})
    log[state]["notices_in_window"] = len(out)
    return out


def match_warn(company: dict, notices: list[dict]) -> tuple[list[dict], list[dict]]:
    """(accepted notices, refused candidates sharing a distinctive token)."""
    comp = set(distinctive(company["canonical_name"]))
    acc, ref = [], []
    for n in notices:
        city_words = {re.sub(r"[^A-Z0-9]", "", w.upper()) for w in re.split(r"\W+", str(n.get("city") or ""))}
        cand = clean_candidate(n["name"], facility=True, drop={w for w in city_words if len(w) > 2})
        if not comp & set(distinctive(cand)):
            continue
        ok, why = identity_ok(company, cand)
        (acc if ok else ref).append({**n, "cleaned": cand, "reason": why})
    return acc, ref


# ------------------------------------------------------------------ permits
CITIES = {
    "CHICAGO": {"label": "City of Chicago", "base": "https://data.cityofchicago.org", "ds": "ydr8-5enu",
                "date": "issue_date", "names": [f"contact_{i}_name" for i in range(1, 16)],
                "types": [f"contact_{i}_type" for i in range(1, 16)]},
    "SEATTLE": {"label": "City of Seattle", "base": "https://data.seattle.gov", "ds": "76t5-zqzr",
                "date": "issueddate", "names": ["contractorcompanyname"], "types": None},
    "AUSTIN": {"label": "City of Austin", "base": "https://data.austintexas.gov", "ds": "3syk-w9eu",
               "date": "issue_date", "names": ["contractor_company_name"], "types": None},
}


def _soql_like(field: str, token: str) -> str:
    return f"upper({field}) like '%{token.upper().replace(chr(39), chr(39) * 2)}%'"


def permits_for(src: Source, city: str, company: dict, start: _dt.date) -> dict:
    cfg = CITIES[city]
    token = search_token(company["canonical_name"])
    base, ds, datef = cfg["base"], cfg["ds"], cfg["date"]
    url = f"{base}/resource/{ds}.json"
    since = f"{datef} >= '{start.isoformat()}T00:00:00'"
    name = company["canonical_name"]
    if cfg["types"] is None:
        field = cfg["names"][0]
        rows = src.get_json(url, {"$select": f"{field}, count(*) as n, max({datef}) as latest",
                                  "$where": f"{_soql_like(field, token)} and {since}",
                                  "$group": field, "$order": "n DESC", "$limit": 200},
                            key=f"permits_{city}_{token}")
        accepted, refused = [], []
        for r in rows:
            raw = str(r.get(field) or "")
            ok, why = identity_ok(company, clean_candidate(raw))
            (accepted if ok else refused).append({"name": raw, "permits": int(r.get("n") or 0),
                                                  "latest": str(r.get("latest") or "")[:10], "reason": why})
        return {"city": city, "token": token, "accepted": accepted, "refused": refused[:12],
                "contractor_permits": sum(a["permits"] for a in accepted),
                "latest": max((a["latest"] for a in accepted), default=""), "roles": {"CONTRACTOR": sum(a["permits"] for a in accepted)},
                "owner_role_permits": 0, "url": f"{base}/d/{ds}"}
    # Chicago: up to 15 typed contacts per permit
    where = "(" + " or ".join(_soql_like(f, token) for f in cfg["names"]) + f") and {since}"
    select = ", ".join(["permit_", datef, "permit_type"] + cfg["names"] + cfg["types"])
    rows, offset = [], 0
    while True:
        page = src.get_json(url, {"$select": select, "$where": where, "$limit": 1000, "$offset": offset,
                                  "$order": f"{datef} DESC"}, key=f"permits_{city}_{token}_{offset}")
        rows.extend(page)
        if len(page) < 1000 or offset >= 9000:
            break
        offset += 1000
    verdicts: dict[str, tuple[bool, str]] = {}
    roles, contractor, owner, latest = Counter(), 0, 0, ""
    accepted_names, refused = Counter(), Counter()
    for r in rows:
        is_contractor = is_owner = False
        for nf, tf in zip(cfg["names"], cfg["types"]):
            raw = str(r.get(nf) or "")
            if token.upper() not in raw.upper():
                continue
            if raw not in verdicts:
                verdicts[raw] = identity_ok(company, clean_candidate(raw))
            if not verdicts[raw][0]:
                refused[raw] += 1
                continue
            accepted_names[raw] += 1
            t = str(r.get(tf) or "").upper()
            roles[t or "(no type)"] += 1
            if "CONTRACTOR" in t:
                is_contractor = True
            elif "OWNER" in t:
                is_owner = True
        if is_contractor:
            contractor += 1
            latest = max(latest, str(r.get(datef) or "")[:10])
        elif is_owner:
            owner += 1
    return {"city": city, "token": token, "rows_read": len(rows), "capped": len(rows) >= 10000,
            "accepted": [{"name": k, "mentions": v} for k, v in accepted_names.most_common(8)],
            "refused": [{"name": k, "mentions": v, "reason": verdicts[k][1]} for k, v in refused.most_common(8)],
            "contractor_permits": contractor, "owner_role_permits": owner, "roles": dict(roles.most_common(8)),
            "latest": latest, "url": f"{base}/d/{ds}"}


# ------------------------------------------------------------------ council
LEGISTAR = "https://webapi.legistar.com/v1/seattle/matters"


def council_for(src: Source, company: dict) -> dict:
    """Seattle Legistar matters whose TITLE carries the company phrase. Full phrase only."""
    phrase = strip_legal_suffix(re.sub(r"\(.*?\)", "", company["canonical_name"])).strip()
    q = phrase.replace("'", "''")
    rows = src.get_json(LEGISTAR, {"$filter": f"substringof('{q}', MatterTitle)", "$top": 200},
                        key=f"council_seattle_{phrase}")
    if not isinstance(rows, list):
        raise AccessError(f"Legistar answered with {type(rows).__name__}, not a list", "source_drift_detected", "code_change")
    one_word = len(phrase.split()) == 1
    verified = [v for v in VERIFIED.get(company["company_id"], []) if len(v.split()) >= 2]
    context = (r"\s*,?\s*(?:Inc|LLC|L\.L\.C|Co|Company|Corp|Corporation|Construction|Builders|"
               r"Contractors|Contracting|Group|Companies|Holdings)\b")
    hits = []
    for m in rows:
        title = str(m.get("MatterTitle") or "")
        # A one-word name in council prose is usually a person: the first dry run matched
        # "Appointment of Cali Mortenson Ellis" for Mortenson. It counts only when a legal-form
        # or line-of-business word follows it, or a verified legal name appears.
        if one_word:
            found = bool(re.search(rf"\b{re.escape(phrase)}{context}", title, re.I)) or any(
                re.search(rf"\b{re.escape(v)}\b", title, re.I) for v in verified)
        else:
            found = bool(re.search(rf"\b{re.escape(phrase)}\b", title, re.I))
        if found:
            hits.append({"id": m.get("MatterId"), "file": m.get("MatterFile"), "type": m.get("MatterTypeName"),
                         "intro": str(m.get("MatterIntroDate") or "")[:10], "title": title[:300]})
    return {"phrase": phrase, "returned": len(rows), "hits": hits}


# ------------------------------------------------------------------ techstack
# Specific markers only (convention 16): an asset host, a generator tag or a platform path --
# never a bare word that page prose could carry. Measured on the 108 archived homepages.
FINGERPRINTS = {
    "WordPress": r"/wp-content/|/wp-includes/|<meta name=\"generator\" content=\"WordPress",
    "Drupal": r"<meta name=\"Generator\" content=\"Drupal|/sites/default/files/|drupal\.js",
    "Sitecore": r"/-/media/[^\"'\s]+\.(?:ashx|jpg|jpeg|png|svg|pdf)|/sitecore/shell",
    "Adobe Experience Manager": r"/etc\.clientlibs/|/content/dam/",
    "Squarespace": r"static1\.squarespace\.com",
    "Wix": r"static\.wixstatic\.com",
    "Webflow": r"data-wf-page=|assets\.website-files\.com|cdn\.prod\.website-files\.com",
    "HubSpot": r"js\.hs-scripts\.com|js\.hsforms\.net|js\.hs-analytics\.net",
    "Salesforce Pardot": r"pi\.pardot\.com|go\.pardot\.com",
    "Marketo": r"munchkin\.marketo\.net",
    "Google Tag Manager": r"googletagmanager\.com/gtm\.js",
    "Google Analytics 4": r"googletagmanager\.com/gtag/js\?id=G-",
    "Shopify": r"cdn\.shopify\.com",
    "Cloudflare": r"/cdn-cgi/",
    "OneTrust": r"cdn\.cookielaw\.org|otSDKStub",
    "Microsoft Clarity": r"clarity\.ms/tag",
}
EXECID_PAGES = ROOT / "harness_output" / "H-EXECID-01" / "raw" / "pages"


def archived_homepage(company: dict) -> tuple[str, str, int, str] | None:
    """(archive date, url, status, html) for the company's homepage as H-EXECID-01 archived it,
    latest partition first; None if never archived."""
    site = re.sub(r"^https?://", "", str(company.get("website") or "")).strip("/")
    if not site:
        return None
    names = {"p_" + slug(site, maxlen=110) + ".txt", "p_" + slug(re.sub(r"^www\.", "", site), maxlen=110) + ".txt"}
    found = sorted((p for n in names for p in glob.glob(str(EXECID_PAGES / "*" / n))),
                   key=lambda p: Path(p).parent.name, reverse=True)
    if not found:
        return None
    text = Path(found[0]).read_text(encoding="utf-8", errors="replace")
    head, _, body = text.partition("\n")
    try:
        status = int(head.strip())
    except ValueError:
        status, body = 0, text
    return Path(found[0]).parent.name, str(company.get("website")), status, body


def fingerprint(html: str) -> dict[str, str]:
    out = {}
    for tech, pattern in FINGERPRINTS.items():
        m = re.search(pattern, html, re.I)
        if m:
            out[tech] = m.group(0)[:60]
    return out


# ------------------------------------------------------------------ rows
def obs(company, signal, topic, text, excerpt, url, pub, stamp, grade, conf, strength="measured_result"):
    return Observation(company_id=company["company_id"], evidence_family=FAMILY_OF[signal],
                       evidence_role="buyer_acts", topic=topic, organizational_state="unknown",
                       signal_strength=strength, observation_text=text, evidence_excerpt=excerpt[:2000],
                       source_url=url, publication_date=pub, retrieval_date=stamp, source_grade=grade,
                       harness_id=HARNESS_ID, harness_version=VERSION, confidence_0_1=conf)


def conf_for(company) -> float:
    return 0.85 if len(distinctive(company["canonical_name"])) >= 2 else 0.75


# ------------------------------------------------------------------ run
def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=f"{HARNESS_ID} -- {HARNESS_NAME}")
    ap.add_argument("--commit", action="store_true")
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--companies")
    ap.add_argument("--scopes", default=",".join(SCOPES))
    ap.add_argument("--pause", type=float, default=0.3)
    ap.add_argument("--show-rows", action="store_true")
    args = ap.parse_args()
    scopes = [s.strip() for s in args.scopes.split(",") if s.strip() in SCOPES]

    db = MarketIntelDB()
    buyers = [c for c in db.companies() if c.get("qualification_status") != "provider_benchmark"]
    if args.companies:
        want = {x.strip() for x in args.companies.split(",")}
        buyers = [c for c in buyers if c["company_id"] in want]
    for c in buyers:
        c["_state"], c["_city"] = state_code(str(c.get("hq_state") or "")), hq_city(c)
    stamp = today()
    start, end = window(stamp)
    VERIFIED.update(load_verified(db))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    src = Source(OUTPUT_DIR / "raw", offline=args.offline, stamp=stamp, pause=args.pause)

    # Declared scope: the (company, signal) pairs whose EMPTY result is a claim.
    declared = []
    if "warn" in scopes:
        declared += [(c["company_id"], SIG_WARN) for c in buyers if c["_state"] in WARN_PAGES]
    if "permits" in scopes:
        declared += [(c["company_id"], SIG_PERMIT) for c in buyers if c["_city"] in CITIES]
    if "council" in scopes:
        declared += [(c["company_id"], SIG_COUNCIL) for c in buyers if c["_city"] == "SEATTLE"]
    if "techstack" in scopes:
        declared += [(c["company_id"], SIG_TECH) for c in buyers]
    run = db.open_run(harness_id=HARNESS_ID, harness_name=HARNESS_NAME, version=VERSION,
                      primary_family=FAM_WARN, scope=declared, signal_families=FAMILY_OF, commit=args.commit)
    log = {"harness_id": HARNESS_ID, "version": VERSION, "date": stamp, "offline": args.offline,
           "window": [start.isoformat(), end.isoformat()], "scopes": scopes, "sources": {}, "companies": [],
           "verified_names": VERIFIED}
    print(f"{HARNESS_ID} {VERSION} -- {len(buyers)} buyers, scopes {scopes}, window {start}..{end} "
          f"({'offline' if args.offline else 'live'})\n")

    # ---- load the list sources once
    notices, warn_state_error = [], {}
    if "warn" in scopes:
        for st in WARN_PAGES:
            try:
                got = warn_notices(src, st, start, log["sources"])
                notices.extend(got)
                print(f"  WARN {st}: {len(got)} notice(s) in window")
            except AccessError as e:
                warn_state_error[st] = e
                print(f"  WARN {st}: UNREAD -- {str(e)[:120]}")

    proposed: list[Observation] = []
    for c in sorted(buyers, key=lambda x: x["company_id"]):
        cid, name = c["company_id"], c["canonical_name"]
        entry = {"company_id": cid, "name": name, "hq_state": c["_state"], "hq_city": c["_city"]}
        marks = []

        # ---- WARN
        if "warn" in scopes:
            acc, ref = match_warn(c, notices)
            entry["warn"] = {"accepted": [{k: str(v) for k, v in n.items()} for n in acc],
                             "refused": [{"name": n["name"], "state": n["state"], "reason": n["reason"]} for n in ref[:8]]}
            by_state: dict[str, list] = {}
            for n in acc:
                by_state.setdefault(n["state"], []).append(n)
            rows = []
            for st, ns in sorted(by_state.items()):
                ns.sort(key=lambda n: n["date"], reverse=True)
                agency, page = WARN_PAGES[st]
                total = sum(int(re.sub(r"\D", "", str(n["count"])) or 0) for n in ns)
                ca_note = (" California's structured list covers the current fiscal year only (from "
                           "2026-07-01); earlier notices are not read." if st == "CA" else "")
                text = (f"{name} appears in the {agency} WARN notice list {len(ns)} time(s) in the five years to "
                        f"{end} (latest {ns[0]['date']}), {total} affected worker(s) in total, as "
                        f"\"{ns[0]['name']}\". A statutory layoff or closure notice (IC4) about the company's own "
                        f"workforce at a site: it records a reduction, closure or relocation, not modernization, "
                        f"and routes to no theme.{ca_note}")
                excerpt = " | ".join(f"{n['date']}: {n['name']} ({n.get('city') or 'n/a'}), {n['count']} workers"
                                     for n in ns[:6])
                rows.append(obs(c, SIG_WARN, "warn_layoff_notice", text, excerpt, page, ns[0]["date"].isoformat(),
                                stamp, "A", conf_for(c)))
            proposed.extend(rows)
            own = [r for r in rows if r.source_url == WARN_PAGES.get(c["_state"], ("", ""))[1]]
            other = [r for r in rows if r not in own]
            if c["_state"] in WARN_PAGES:
                if c["_state"] in warn_state_error:
                    e = warn_state_error[c["_state"]]
                    run.attempt(cid, SIG_WARN, outcome="not_covered", failure_stage="fetch",
                                failure_category=e.failure_category, fix_class=e.fix_class,
                                failure_detail=str(e)[:500], source_url_attempted=WARN_PAGES[c["_state"]][1])
                elif own:
                    run.attempt(cid, SIG_WARN, outcome="covered", records_written=len(own),
                                source_url_attempted=own[0].source_url, candidates_evaluated=len(acc) + len(ref))
                else:
                    run.attempt(cid, SIG_WARN, outcome="absent_confirmed", source_url_attempted=WARN_PAGES[c["_state"]][1],
                                candidates_evaluated=len(ref), candidates_discarded=len(ref))
            if other:
                run.attempt(cid, SIG_WARN, outcome="covered", scope="incidental", records_written=len(other),
                            source_url_attempted=other[0].source_url)
            if rows:
                marks.append(f"WARN {sorted(by_state)}")

        # ---- permits
        if "permits" in scopes:
            results, errors = {}, {}
            for city in CITIES:
                try:
                    results[city] = permits_for(src, city, c, start)
                except AccessError as e:
                    errors[city] = e
            entry["permits"] = {"results": results, "errors": {k: str(v) for k, v in errors.items()}}
            rows = []
            for city, r in results.items():
                if not r["contractor_permits"]:
                    continue
                cfg = CITIES[city]
                owner_note = (f" {r['owner_role_permits']} permit(s) name the company in an owner role; not counted, "
                              f"because the owner slot in this dataset carries misfiled contractors and namesakes."
                              if r.get("owner_role_permits") else "")
                cap = " The listing hit its row cap, so the count is a floor." if r.get("capped") else ""
                text = (f"{name} is named as contractor on {r['contractor_permits']} building permit(s) issued by the "
                        f"{cfg['label']} in the five years to {end} (latest issue {r['latest'] or 'n/a'}). A municipal "
                        f"record in which the company is the CONTRACTOR on someone else's project: local project "
                        f"volume as a seller, not its own facilities or modernization; routes to no theme."
                        f"{owner_note}{cap}")
                excerpt = (f"search token {r['token']} | accepted names: "
                           + "; ".join(f"{a['name']} ({a.get('permits', a.get('mentions'))})" for a in r["accepted"][:6])
                           + f" | roles: {r.get('roles')}")
                if r["refused"]:
                    excerpt += " | refused: " + "; ".join(f"{x['name']} ({x['reason'][:40]})" for x in r["refused"][:4])
                rows.append(obs(c, SIG_PERMIT, "municipal_permits_as_contractor", text, excerpt, r["url"],
                                r["latest"] or "", stamp, "A", conf_for(c)))
            proposed.extend(rows)
            own_city = c["_city"] if c["_city"] in CITIES else ""
            own = [r for r in rows if own_city and r.source_url == f"{CITIES[own_city]['base']}/d/{CITIES[own_city]['ds']}"]
            other = [r for r in rows if r not in own]
            if own_city:
                if own_city in errors:
                    e = errors[own_city]
                    run.attempt(cid, SIG_PERMIT, outcome="not_covered", failure_stage="fetch",
                                failure_category=e.failure_category, fix_class=e.fix_class,
                                failure_detail=str(e)[:500], source_url_attempted=CITIES[own_city]["base"])
                elif own:
                    run.attempt(cid, SIG_PERMIT, outcome="covered", records_written=len(own),
                                source_url_attempted=own[0].source_url)
                else:
                    run.attempt(cid, SIG_PERMIT, outcome="absent_confirmed",
                                source_url_attempted=f"{CITIES[own_city]['base']}/d/{CITIES[own_city]['ds']}",
                                candidates_evaluated=len(results[own_city]["refused"]),
                                candidates_discarded=len(results[own_city]["refused"]))
            if other:
                run.attempt(cid, SIG_PERMIT, outcome="covered", scope="incidental", records_written=len(other),
                            source_url_attempted=other[0].source_url)
            if rows:
                marks.append(f"permits {[r.source_url.split('/')[2] for r in rows]}")

        # ---- council
        if "council" in scopes:
            try:
                res = council_for(src, c)
                err = None
            except AccessError as e:
                res, err = {"hits": []}, e
            entry["council"] = {**res, "error": str(err) if err else ""}
            rows = []
            if res["hits"]:
                hits = sorted(res["hits"], key=lambda h: h["intro"], reverse=True)
                text = (f"{name} is named in {len(hits)} Seattle City Council matter title(s) (latest introduced "
                        f"{hits[0]['intro'] or 'n/a'}). A council record naming the company, usually as a party to "
                        f"a city agreement or land-use action; it records a civic transaction, not modernization.")
                excerpt = " | ".join(f"{h['file']} ({h['type']}, {h['intro']}): {h['title'][:200]}" for h in hits[:4])
                rows.append(obs(c, SIG_COUNCIL, "city_council_matter_mention", text, excerpt,
                                "https://seattle.legistar.com/Legislation.aspx", hits[0]["intro"], stamp, "A", conf_for(c)))
            proposed.extend(rows)
            if c["_city"] == "SEATTLE":
                if err:
                    run.attempt(cid, SIG_COUNCIL, outcome="not_covered", failure_stage="fetch",
                                failure_category=err.failure_category, fix_class=err.fix_class,
                                failure_detail=str(err)[:500], source_url_attempted=LEGISTAR)
                elif rows:
                    run.attempt(cid, SIG_COUNCIL, outcome="covered", records_written=1, source_url_attempted=LEGISTAR)
                else:
                    run.attempt(cid, SIG_COUNCIL, outcome="absent_confirmed", source_url_attempted=LEGISTAR,
                                candidates_evaluated=res.get("returned", 0), candidates_discarded=res.get("returned", 0))
            elif rows:
                run.attempt(cid, SIG_COUNCIL, outcome="covered", scope="incidental", records_written=1,
                            source_url_attempted=LEGISTAR)
            if rows:
                marks.append("council")

        # ---- techstack
        if "techstack" in scopes:
            arch = archived_homepage(c)
            if arch is None:
                run.attempt(cid, SIG_TECH, outcome="not_covered", failure_stage="discovery",
                            failure_category="source_not_found", fix_class="source_limitation",
                            failure_detail="no homepage in H-EXECID-01's archive", source_url_attempted=str(c.get("website") or ""))
                entry["techstack"] = {"archived": False}
            else:
                adate, url, status, html = arch
                fps = fingerprint(html) if status == 200 else {}
                entry["techstack"] = {"archive_date": adate, "status": status, "fingerprints": fps}
                if status != 200:
                    run.attempt(cid, SIG_TECH, outcome="not_covered", failure_stage="fetch",
                                failure_category="access_blocked" if status in (401, 403) else "source_unavailable",
                                fix_class="source_limitation" if status in (401, 403) else "transient",
                                failure_detail=f"archived homepage response was HTTP {status} ({adate})",
                                source_url_attempted=url)
                elif fps:
                    text = (f"{name}'s homepage, as served on {adate}, carries markers of: {', '.join(sorted(fps))}. "
                            f"Front end only: what the public website is built and measured with, not the "
                            f"company's internal systems; routes to no theme.")
                    excerpt = " | ".join(f"{k}: {v}" for k, v in sorted(fps.items()))
                    o = obs(c, SIG_TECH, "website_technology_stack", text, excerpt, url, adate, adate, "B", 0.8)
                    proposed.append(o)
                    run.attempt(cid, SIG_TECH, outcome="covered", records_written=1, source_url_attempted=url)
                    marks.append(f"tech {len(fps)}")
                else:
                    run.attempt(cid, SIG_TECH, outcome="absent_confirmed", source_url_attempted=url)

        log["companies"].append(entry)
        if marks:
            print(f"  [ok] {cid} {name} ({c['_city'] or '?'}, {c['_state'] or '?'}): {'; '.join(marks)}")

    if args.show_rows:
        for o in proposed:
            print()
            print(f"  ---- {o.company_id} {o.topic} grade={o.source_grade} conf={o.confidence_0_1} pub={o.publication_date}")
            print(f"  TEXT: {o.observation_text}")
            print(f"  EXCERPT: {o.evidence_excerpt}")
            print(f"  URL: {o.source_url}")
    report = db.sync_observations(proposed)
    run.observations_written = report.written
    summary = run.close()
    by_topic = Counter(o.topic for o in proposed)
    print(f"\n  {len(buyers)} buyers - {len(proposed)} observations {dict(by_topic)} - {src.cache.fetch_count} fetches")
    print(f"  attempts: {summary['attempts_covered']} covered, {summary['attempts_absent_confirmed']} absent_confirmed, "
          f"{summary['attempts_not_covered']} not_covered (scoped)")
    print(f"  dedupe: {report.summary()}")
    if run.derived_known_issues():
        print(f"  issues: {run.derived_known_issues()}")
    log["summary"] = summary
    log["dedupe"] = report.summary()
    log["held"] = report.conflicts
    path = OUTPUT_DIR / f"run-{stamp}{'' if args.commit else '-dryrun'}.json"
    path.write_text(json.dumps(log, indent=2, default=str), encoding="utf-8")
    if args.commit:
        db.save()
        print(f"  committed to market_intel_db.xlsx ({summary['run_id']})")
    else:
        print("  DRY RUN -- nothing written. Re-run with --commit to write.")
    print(f"  run log: {path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
